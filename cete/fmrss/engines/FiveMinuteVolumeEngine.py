import gc
import os
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd


try:
    from engines._volume_rolling import (
        rolling_volume_baselines
    )
    CYTHON_VOLUME_AVAILABLE = True
except ImportError:
    try:
        from _volume_rolling import (
            rolling_volume_baselines
        )
        CYTHON_VOLUME_AVAILABLE = True
    except ImportError:
        rolling_volume_baselines = None
        CYTHON_VOLUME_AVAILABLE = False


class FiveMinuteVolumeEngine:

    # =====================================================
    # Required Input Columns
    # =====================================================

    REQUIRED_COLUMNS = [
        "Symbol",
        "DateTime",
        "Volume",
        "Stable_Regime",
        "Regime_Stable"
    ]

    # =====================================================
    # Constructor
    # =====================================================

    def __init__(
        self,
        rolling_volume_window=20,
        rolling_volume_min_periods=15,
        time_slot_lookback_days=10,
        time_slot_min_periods=5,
        momentum_minimum_rvol=1.25,
        momentum_minimum_time_rvol=1.20,
        reversion_minimum_rvol=1.00,
        maximum_volume_factor=3.0,
        volume_floor=1e-12,
        use_cython=True,
        memory_optimized=True,
        retain_results_in_memory=False,
        save_full_history_to_excel=False,
        report_output_file=(
            "/home/devinderjeet/fmrss/report/"
            "ResultVolumeEngine.xlsx"
        ),
        save_excel=True,
    ):

        if rolling_volume_window < 2:
            raise ValueError(
                "rolling_volume_window must be at least 2"
            )

        if not 1 <= rolling_volume_min_periods <= rolling_volume_window:
            raise ValueError(
                "rolling_volume_min_periods must be between "
                "1 and rolling_volume_window"
            )

        if time_slot_lookback_days < 2:
            raise ValueError(
                "time_slot_lookback_days must be at least 2"
            )

        if not 1 <= time_slot_min_periods <= time_slot_lookback_days:
            raise ValueError(
                "time_slot_min_periods must be between "
                "1 and time_slot_lookback_days"
            )

        if momentum_minimum_rvol <= 0:
            raise ValueError(
                "momentum_minimum_rvol must be positive"
            )

        if momentum_minimum_time_rvol <= 0:
            raise ValueError(
                "momentum_minimum_time_rvol must be positive"
            )

        if reversion_minimum_rvol <= 0:
            raise ValueError(
                "reversion_minimum_rvol must be positive"
            )

        if maximum_volume_factor <= 0:
            raise ValueError(
                "maximum_volume_factor must be positive"
            )

        if volume_floor <= 0:
            raise ValueError(
                "volume_floor must be positive"
            )

        self.rolling_volume_window = int(
            rolling_volume_window
        )

        self.rolling_volume_min_periods = int(
            rolling_volume_min_periods
        )

        self.time_slot_lookback_days = int(
            time_slot_lookback_days
        )

        self.time_slot_min_periods = int(
            time_slot_min_periods
        )

        self.momentum_minimum_rvol = float(
            momentum_minimum_rvol
        )

        self.momentum_minimum_time_rvol = float(
            momentum_minimum_time_rvol
        )

        self.reversion_minimum_rvol = float(
            reversion_minimum_rvol
        )

        self.maximum_volume_factor = float(
            maximum_volume_factor
        )

        self.volume_floor = float(
            volume_floor
        )

        self.use_cython = bool(use_cython)
        self.memory_optimized = bool(memory_optimized)
        self.retain_results_in_memory = bool(
            retain_results_in_memory
        )
        self.save_full_history_to_excel = bool(
            save_full_history_to_excel
        )
        self.report_output_file = (
            str(report_output_file)
            if report_output_file
            else None
        )
        self.save_excel_enabled = bool(save_excel)

        self.volume_backend = (
            "CYTHON"
            if self.use_cython
            and CYTHON_VOLUME_AVAILABLE
            else "PYTHON"
        )

    # =====================================================
    # Validate Input
    # =====================================================

    def _validate_input(self, symbol, df):

        if not isinstance(df, pd.DataFrame):
            raise TypeError(
                f"{symbol} data must be a pandas DataFrame"
            )

        if df.empty:
            raise ValueError(
                f"{symbol} regime data is empty"
            )

        missing_columns = [
            column
            for column in self.REQUIRED_COLUMNS
            if column not in df.columns
        ]

        if missing_columns:
            raise ValueError(
                f"{symbol} is missing required columns: "
                f"{missing_columns}"
            )

    # =====================================================
    # Prepare Input
    # =====================================================

    def _prepare_input(self, symbol, df):

        self._validate_input(
            symbol=symbol,
            df=df
        )

        # Share immutable upstream blocks while allocating new volume-feature
        # arrays separately.  The caller can release regime_data after this
        # stage without duplicating the full residual/regime history.
        result = df.copy(deep=not self.memory_optimized)

        result["DateTime"] = pd.to_datetime(
            result["DateTime"],
            errors="coerce"
        )

        result["Volume"] = pd.to_numeric(
            result["Volume"],
            errors="coerce"
        )

        result["Volume"] = (
            result["Volume"]
            .replace(
                [np.inf, -np.inf],
                np.nan
            )
        )
        if self.memory_optimized:
            result["Volume"] = result["Volume"].astype(
                np.float32,
                copy=False,
            )

        result["Stable_Regime"] = (
            result["Stable_Regime"]
            .fillna("NOT_READY")
            .astype(str)
            .str.strip()
            .str.upper()
            .astype("category")
        )

        result["Regime_Stable"] = (
            result["Regime_Stable"]
            .fillna(False)
            .astype(bool)
        )

        result.dropna(
            subset=["DateTime"],
            inplace=True
        )

        invalid_volume = (
            result["Volume"].notna()
            & (result["Volume"] < 0)
        )

        if invalid_volume.any():
            result.drop(
                index=result.index[invalid_volume],
                inplace=True
            )

        result["Symbol"] = (
            str(symbol)
            .strip()
            .upper()
        )

        result.sort_values(
            "DateTime",
            inplace=True
        )

        result.drop_duplicates(
            subset=["DateTime"],
            keep="last",
            inplace=True
        )

        result.reset_index(
            drop=True,
            inplace=True
        )

        if result.empty:
            raise ValueError(
                f"{symbol} has no valid timestamped rows"
            )

        # Integer time-slot keys use less memory and compare faster than
        # Python datetime.time objects.
        result["TimeSlot_Minute"] = (
            result["DateTime"].dt.hour * 60
            + result["DateTime"].dt.minute
        ).astype(np.int16)

        return result

    # =====================================================
    # Cython Rolling Volume Baselines
    # =====================================================

    def _calculate_cython_volume_baselines(self, df):

        volume_values = np.require(
            df["Volume"].to_numpy(dtype=np.float64, copy=False),
            dtype=np.float64,
            requirements=["C", "A", "W"],
        )

        time_slot_values = np.require(
            df["TimeSlot_Minute"].to_numpy(
                dtype=np.int16,
                copy=False,
            ),
            dtype=np.int16,
            requirements=["C", "A", "W"],
        )

        (
            rolling_median,
            rolling_observations,
            time_slot_median,
            time_slot_observations
        ) = rolling_volume_baselines(
            volume_values,
            time_slot_values,
            self.rolling_volume_window,
            self.rolling_volume_min_periods,
            self.time_slot_lookback_days,
            self.time_slot_min_periods
        )

        valid_rolling = (
            np.isfinite(rolling_median)
            & (rolling_median > self.volume_floor)
            & (
                rolling_observations
                >= self.rolling_volume_min_periods
            )
        )

        valid_time_slot = (
            np.isfinite(time_slot_median)
            & (time_slot_median > self.volume_floor)
            & (
                time_slot_observations
                >= self.time_slot_min_periods
            )
        )

        relative_volume = np.full(
            len(df),
            np.nan,
            dtype=np.float64
        )

        time_adjusted_rvol = np.full(
            len(df),
            np.nan,
            dtype=np.float64
        )

        finite_volume = np.isfinite(volume_values)

        np.divide(
            volume_values,
            rolling_median,
            out=relative_volume,
            where=valid_rolling & finite_volume
        )

        np.divide(
            volume_values,
            time_slot_median,
            out=time_adjusted_rvol,
            where=valid_time_slot & finite_volume
        )

        df["Rolling_Volume_Median"] = rolling_median
        df["Rolling_Volume_Observations"] = (
            rolling_observations
        )
        df["Rolling_Volume_Ready"] = valid_rolling
        df["Relative_Volume"] = relative_volume

        df["TimeSlot_Volume_Median"] = time_slot_median
        df["TimeSlot_Volume_Observations"] = (
            time_slot_observations
        )
        df["TimeSlot_Volume_Ready"] = valid_time_slot
        df["Time_Adjusted_RVOL"] = time_adjusted_rvol

        return df

    # =====================================================
    # Rolling Volume Backend Dispatcher
    # =====================================================

    def _calculate_volume_baselines(self, df):

        if self.volume_backend == "CYTHON":
            return self._calculate_cython_volume_baselines(df)

        df = self._calculate_rolling_relative_volume(df)
        return self._calculate_time_adjusted_volume(df)

    # =====================================================
    # Rolling Relative Volume
    # =====================================================

    def _calculate_rolling_relative_volume(self, df):

        volume = df["Volume"]

        # Exclude the current candle from both the baseline and count.
        rolling_median = (
            volume
            .rolling(
                window=self.rolling_volume_window,
                min_periods=self.rolling_volume_min_periods
            )
            .median()
            .shift(1)
        )

        rolling_observations = (
            volume
            .rolling(
                window=self.rolling_volume_window,
                min_periods=1
            )
            .count()
            .shift(1)
            .fillna(0)
            .astype(np.int16)
        )

        valid_baseline = (
            rolling_median.notna()
            & (rolling_median > self.volume_floor)
            & (
                rolling_observations
                >= self.rolling_volume_min_periods
            )
        )

        df["Rolling_Volume_Median"] = rolling_median

        df["Rolling_Volume_Observations"] = (
            rolling_observations
        )

        df["Rolling_Volume_Ready"] = valid_baseline

        df["Relative_Volume"] = np.where(
            valid_baseline & volume.notna(),
            volume / rolling_median,
            np.nan
        )

        return df

    # =====================================================
    # Same-Time Previous-Day Volume
    # =====================================================

    def _calculate_time_adjusted_volume(self, df):

        grouped_volume = df.groupby(
            "TimeSlot_Minute",
            sort=False
        )["Volume"]

        # Within each time slot, shift(1) excludes today's current slot.
        # The resulting window contains only earlier trading days.
        time_slot_median = grouped_volume.transform(
            lambda series: (
                series
                .rolling(
                    window=self.time_slot_lookback_days,
                    min_periods=self.time_slot_min_periods
                )
                .median()
                .shift(1)
            )
        )

        time_slot_observations = grouped_volume.transform(
            lambda series: (
                series
                .rolling(
                    window=self.time_slot_lookback_days,
                    min_periods=1
                )
                .count()
                .shift(1)
                .fillna(0)
            )
        ).astype(np.int16)

        valid_time_baseline = (
            time_slot_median.notna()
            & (time_slot_median > self.volume_floor)
            & (
                time_slot_observations
                >= self.time_slot_min_periods
            )
        )

        df["TimeSlot_Volume_Median"] = (
            time_slot_median
        )

        df["TimeSlot_Volume_Observations"] = (
            time_slot_observations
        )

        df["TimeSlot_Volume_Ready"] = (
            valid_time_baseline
        )

        df["Time_Adjusted_RVOL"] = np.where(
            valid_time_baseline & df["Volume"].notna(),
            df["Volume"] / time_slot_median,
            np.nan
        )

        return df

    # =====================================================
    # Combined Volume Factor
    # =====================================================

    def _calculate_volume_factor(self, df):

        both_ready = (
            df["Rolling_Volume_Ready"]
            & df["TimeSlot_Volume_Ready"]
        )

        volume_product = (
            df["Relative_Volume"]
            * df["Time_Adjusted_RVOL"]
        )

        volume_product = volume_product.where(
            volume_product >= 0,
            np.nan
        )

        df["Volume_Features_Ready"] = both_ready

        df["Volume_Factor_Raw"] = np.where(
            both_ready,
            np.sqrt(volume_product),
            np.nan
        )

        df["Volume_Factor"] = (
            df["Volume_Factor_Raw"]
            .clip(
                lower=0.0,
                upper=self.maximum_volume_factor
            )
        )

        return df

    # =====================================================
    # Volume State
    # =====================================================

    @staticmethod
    def _classify_volume_state(df):

        factor = df["Volume_Factor"]

        df["Volume_State"] = pd.Categorical(np.select(
            condlist=[
                ~df["Volume_Features_Ready"],
                factor >= 2.0,
                factor >= 1.25,
                factor >= 0.75
            ],
            choicelist=[
                "NOT_READY",
                "EXTREME",
                "HIGH",
                "NORMAL"
            ],
            default="LOW"
        ))

        return df

    # =====================================================
    # Regime-Specific Volume Confirmation
    # =====================================================

    def _calculate_volume_confirmation(self, df):

        momentum_ready = (
            df["Regime_Stable"]
            & df["Stable_Regime"].eq("MOMENTUM")
            & df["Rolling_Volume_Ready"]
            & df["TimeSlot_Volume_Ready"]
        )

        reversion_ready = (
            df["Regime_Stable"]
            & df["Stable_Regime"].eq("MEAN_REVERSION")
            & df["Rolling_Volume_Ready"]
        )

        df["Momentum_Volume_Confirmed"] = (
            momentum_ready
            & (
                df["Relative_Volume"]
                >= self.momentum_minimum_rvol
            )
            & (
                df["Time_Adjusted_RVOL"]
                >= self.momentum_minimum_time_rvol
            )
        )

        df["Reversion_Volume_Confirmed"] = (
            reversion_ready
            & (
                df["Relative_Volume"]
                >= self.reversion_minimum_rvol
            )
        )

        df["Volume_Confirmation_Ready"] = (
            momentum_ready
            | reversion_ready
        )

        df["Volume_Confirmed"] = np.select(
            condlist=[
                df["Stable_Regime"].eq("MOMENTUM"),
                df["Stable_Regime"].eq("MEAN_REVERSION")
            ],
            choicelist=[
                df["Momentum_Volume_Confirmed"],
                df["Reversion_Volume_Confirmed"]
            ],
            default=False
        ).astype(bool)

        return df

    def _compact_feature_dtypes(self, df):
        if not self.memory_optimized:
            return df

        for column in [
            "Rolling_Volume_Median",
            "Relative_Volume",
            "TimeSlot_Volume_Median",
            "Time_Adjusted_RVOL",
            "Volume_Factor_Raw",
            "Volume_Factor",
        ]:
            if column in df:
                df[column] = df[column].astype(np.float32, copy=False)

        for column in [
            "Rolling_Volume_Observations",
            "TimeSlot_Volume_Observations",
        ]:
            if column in df:
                df[column] = df[column].astype(np.int16, copy=False)

        if "Volume_State" in df:
            df["Volume_State"] = df["Volume_State"].astype("category")

        return df

    # =====================================================
    # Calculate One Symbol
    # =====================================================

    def calculate_symbol(
        self,
        symbol,
        regime_df
    ):

        symbol = (
            str(symbol)
            .strip()
            .upper()
        )

        df = self._prepare_input(
            symbol=symbol,
            df=regime_df
        )

        df = self._calculate_volume_baselines(df)

        df = self._calculate_volume_factor(df)

        df = self._classify_volume_state(df)

        df = self._calculate_volume_confirmation(df)

        df = self._compact_feature_dtypes(df)

        feature_ready_rows = int(
            df["Volume_Features_Ready"].sum()
        )

        latest = df.iloc[-1]

        report = {
            "Symbol": symbol,
            "Status": (
                "CALCULATED"
                if feature_ready_rows > 0
                else "NOT_READY"
            ),
            "Input_Rows": len(df),
            "Volume_Ready_Rows": feature_ready_rows,
            "Latest_DateTime": latest["DateTime"],
            "Latest_Relative_Volume": (
                latest["Relative_Volume"]
            ),
            "Latest_Time_Adjusted_RVOL": (
                latest["Time_Adjusted_RVOL"]
            ),
            "Latest_Volume_Factor": (
                latest["Volume_Factor"]
            ),
            "Latest_Volume_State": (
                latest["Volume_State"]
            ),
            "Latest_Stable_Regime": (
                latest["Stable_Regime"]
            ),
            "Latest_Volume_Confirmed": bool(
                latest["Volume_Confirmed"]
            ),
            "Volume_Backend": self.volume_backend,
            "Error": "",
        }

        return df, report

    # =====================================================
    # Excel Report
    # =====================================================

    @staticmethod
    def _excel_safe(frame):
        safe = frame.copy(deep=False).replace([np.inf, -np.inf], np.nan)
        for column in safe.columns:
            if isinstance(safe[column].dtype, pd.CategoricalDtype):
                safe[column] = safe[column].astype("string")
            if isinstance(safe[column].dtype, pd.DatetimeTZDtype):
                safe[column] = safe[column].dt.tz_localize(None)
        return safe

    def _build_summary(self, report_df):
        status = report_df["Status"].astype("string")
        state = report_df["Latest_Volume_State"].astype("string")
        confirmed = report_df["Latest_Volume_Confirmed"].fillna(False)
        return pd.DataFrame(
            [
                ["RUN", "Report generated", pd.Timestamp.now()],
                ["RUN", "Volume backend", self.volume_backend],
                ["RUN", "Memory optimized", self.memory_optimized],
                [
                    "RUN",
                    "Results retained in memory",
                    self.retain_results_in_memory,
                ],
                [
                    "RUN",
                    "Full history saved to Excel",
                    self.save_full_history_to_excel,
                ],
                ["COUNTS", "Total symbols", int(len(report_df))],
                ["COUNTS", "Calculated symbols", int(status.eq("CALCULATED").sum())],
                ["COUNTS", "Not-ready symbols", int(status.eq("NOT_READY").sum())],
                ["COUNTS", "Failed symbols", int(status.eq("FAILED").sum())],
                ["COUNTS", "Volume-confirmed symbols", int(confirmed.sum())],
                ["STATE", "EXTREME symbols", int(state.eq("EXTREME").sum())],
                ["STATE", "HIGH symbols", int(state.eq("HIGH").sum())],
                ["STATE", "NORMAL symbols", int(state.eq("NORMAL").sum())],
                ["STATE", "LOW symbols", int(state.eq("LOW").sum())],
                ["STATE", "NOT_READY symbols", int(state.eq("NOT_READY").sum())],
                ["CONFIGURATION", "Rolling volume window", self.rolling_volume_window],
                ["CONFIGURATION", "Rolling minimum periods", self.rolling_volume_min_periods],
                ["CONFIGURATION", "Time-slot lookback days", self.time_slot_lookback_days],
                ["CONFIGURATION", "Time-slot minimum periods", self.time_slot_min_periods],
                ["CONFIGURATION", "Momentum minimum RVOL", self.momentum_minimum_rvol],
                ["CONFIGURATION", "Momentum minimum time RVOL", self.momentum_minimum_time_rvol],
                ["CONFIGURATION", "Reversion minimum RVOL", self.reversion_minimum_rvol],
                ["CONFIGURATION", "Maximum volume factor", self.maximum_volume_factor],
            ],
            columns=["Section", "Metric", "Value"],
        )

    @staticmethod
    def _style_workbook(writer):
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(color="FFFFFF", bold=True)
        green_fill = PatternFill("solid", fgColor="C6EFCE")
        red_fill = PatternFill("solid", fgColor="FFC7CE")
        amber_fill = PatternFill("solid", fgColor="FFEB9C")

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            if worksheet.title == "Volume_History":
                from openpyxl.utils import get_column_letter

                for index in range(1, worksheet.max_column + 1):
                    worksheet.column_dimensions[
                        get_column_letter(index)
                    ].width = 21
                continue

            for cells in worksheet.columns:
                values = [
                    "" if cell.value is None else str(cell.value)
                    for cell in cells[:200]
                ]
                width = min(
                    max(max(map(len, values), default=0) + 2, 11),
                    34,
                )
                worksheet.column_dimensions[cells[0].column_letter].width = width

            headers = {
                cell.value: cell.column_letter
                for cell in worksheet[1]
                if cell.value is not None
            }
            datetime_column = headers.get("Latest_DateTime")
            if datetime_column:
                for row in range(2, worksheet.max_row + 1):
                    worksheet[f"{datetime_column}{row}"].number_format = (
                        "yyyy-mm-dd hh:mm:ss"
                    )

            # Avoid invalid conditional ranges on header-only worksheets.
            if worksheet.max_row < 2:
                continue

            confirmed_column = headers.get("Latest_Volume_Confirmed")
            if confirmed_column:
                target = (
                    f"{confirmed_column}2:"
                    f"{confirmed_column}{worksheet.max_row}"
                )
                worksheet.conditional_formatting.add(
                    target,
                    FormulaRule(
                        formula=[f"{confirmed_column}2=TRUE"],
                        fill=green_fill,
                    ),
                )

            status_column = headers.get("Status")
            if status_column:
                target = (
                    f"{status_column}2:"
                    f"{status_column}{worksheet.max_row}"
                )
                worksheet.conditional_formatting.add(
                    target,
                    FormulaRule(
                        formula=[f'{status_column}2="FAILED"'],
                        fill=red_fill,
                    ),
                )
                worksheet.conditional_formatting.add(
                    target,
                    FormulaRule(
                        formula=[f'{status_column}2="NOT_READY"'],
                        fill=amber_fill,
                    ),
                )

    def _write_volume_history(self, writer, volume_data):
        maximum_data_rows = 1_048_575
        total_rows = sum(
            len(frame)
            for frame in volume_data.values()
            if isinstance(frame, pd.DataFrame) and not frame.empty
        )
        if total_rows > maximum_data_rows:
            raise ValueError(
                "Volume_History exceeds Excel's worksheet limit: "
                f"{total_rows:,} rows > {maximum_data_rows:,} rows"
            )

        if total_rows == 0:
            pd.DataFrame(columns=self.REQUIRED_COLUMNS).to_excel(
                writer, sheet_name="Volume_History", index=False
            )
            return 0

        start_row = 0
        write_header = True
        written_rows = 0
        expected_columns = None
        for symbol, frame in volume_data.items():
            if not isinstance(frame, pd.DataFrame) or frame.empty:
                continue
            if expected_columns is None:
                expected_columns = list(frame.columns)
            elif list(frame.columns) != expected_columns:
                frame = frame.reindex(columns=expected_columns)
            chunk = self._excel_safe(frame)
            chunk.to_excel(
                writer,
                sheet_name="Volume_History",
                index=False,
                header=write_header,
                startrow=start_row,
            )
            count = len(chunk)
            written_rows += count
            start_row += count + (1 if write_header else 0)
            write_header = False
            del chunk
        return written_rows

    def _save_excel_report(self, report_df, output_file, volume_data=None):
        output_path = Path(output_file).expanduser().resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        confirmed = report_df.loc[
            report_df["Latest_Volume_Confirmed"].fillna(False)
        ].copy()
        unconfirmed = report_df.loc[
            report_df["Status"].eq("CALCULATED")
            & ~report_df["Latest_Volume_Confirmed"].fillna(False)
        ].copy()
        not_ready_failed = report_df.loc[
            report_df["Status"].isin(["NOT_READY", "FAILED"])
        ].copy()
        summary = self._build_summary(report_df)

        descriptor, temporary_name = tempfile.mkstemp(
            prefix="VolumeEngine_",
            suffix=".xlsx",
            dir=output_path.parent,
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(
                temporary_name,
                engine="openpyxl",
                datetime_format="yyyy-mm-dd hh:mm:ss",
            ) as writer:
                self._excel_safe(summary).to_excel(
                    writer, sheet_name="Summary", index=False
                )
                self._excel_safe(confirmed).to_excel(
                    writer, sheet_name="Confirmed_Volume", index=False
                )
                self._excel_safe(unconfirmed).to_excel(
                    writer, sheet_name="Unconfirmed_Volume", index=False
                )
                self._excel_safe(not_ready_failed).to_excel(
                    writer, sheet_name="Not_Ready_Failed", index=False
                )
                self._excel_safe(report_df).to_excel(
                    writer, sheet_name="All_Shares", index=False
                )
                if self.save_full_history_to_excel:
                    history_rows = self._write_volume_history(
                        writer,
                        volume_data or {},
                    )
                else:
                    history_rows = 0
                self._style_workbook(writer)

            # Existing report is replaced only after the new workbook closes.
            os.replace(temporary_name, output_path)
        except Exception:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)
            raise

        report_df.attrs["Excel_File"] = str(output_path)
        report_df.attrs["Volume_History_Rows"] = history_rows
        return str(output_path)

    # =====================================================
    # Calculate from RegimeEngine.xlsx
    # =====================================================

    def calculate_all_from_excel(
        self,
        regime_file,
        sheet_name="Regime_History",
        output_file=None,
    ):
        """Stream regime history and calculate one symbol block at a time."""
        from openpyxl import load_workbook

        input_path = Path(regime_file).expanduser().resolve()
        if not input_path.is_file():
            raise FileNotFoundError(
                f"Regime workbook not found: {input_path}"
            )

        workbook = load_workbook(
            filename=input_path,
            read_only=True,
            data_only=True,
        )
        available_sheets = list(workbook.sheetnames)
        if sheet_name not in available_sheets:
            workbook.close()
            raise ValueError(
                f"Worksheet {sheet_name!r} was not found in "
                f"{input_path}. Available sheets: {available_sheets}"
            )

        volume_data = {}
        reports = []
        processed_symbols = set()
        calculated_symbols = 0

        def failed_report(symbol, row_count, error):
            return {
                "Symbol": symbol,
                "Status": "FAILED",
                "Input_Rows": int(row_count),
                "Volume_Ready_Rows": 0,
                "Latest_DateTime": pd.NaT,
                "Latest_Relative_Volume": np.nan,
                "Latest_Time_Adjusted_RVOL": np.nan,
                "Latest_Volume_Factor": np.nan,
                "Latest_Volume_State": "NOT_READY",
                "Latest_Stable_Regime": "NOT_READY",
                "Latest_Volume_Confirmed": False,
                "Volume_Backend": self.volume_backend,
                "Error": str(error),
            }

        def process_symbol(symbol, rows, columns):
            nonlocal calculated_symbols
            if not symbol or not rows:
                return
            try:
                regime_df = pd.DataFrame.from_records(
                    rows,
                    columns=columns,
                )
                result_df, report = self.calculate_symbol(
                    symbol=symbol,
                    regime_df=regime_df,
                )
                reports.append(report)
                if report["Status"] != "NOT_READY":
                    calculated_symbols += 1
                    if (
                        self.retain_results_in_memory
                        or self.save_full_history_to_excel
                    ):
                        volume_data[symbol] = result_df
                    else:
                        del result_df
                else:
                    del result_df
                del regime_df
            except Exception as error:
                reports.append(
                    failed_report(symbol, len(rows), error)
                )

        try:
            worksheet = workbook[sheet_name]
            row_iterator = worksheet.iter_rows(values_only=True)
            try:
                raw_header = next(row_iterator)
            except StopIteration as error:
                raise ValueError(
                    f"Worksheet {sheet_name!r} is empty"
                ) from error

            columns = [
                str(value).strip() if value is not None else ""
                for value in raw_header
            ]
            missing_columns = [
                column
                for column in self.REQUIRED_COLUMNS
                if column not in columns
            ]
            if missing_columns:
                raise ValueError(
                    f"Worksheet {sheet_name!r} is missing columns: "
                    f"{missing_columns}. Recreate ResidualEngine.xlsx "
                    "and RegimeEngine.xlsx with their updated engines."
                )

            symbol_index = columns.index("Symbol")
            current_symbol = None
            current_rows = []

            for values in row_iterator:
                if not values or all(value is None for value in values):
                    continue
                raw_symbol = values[symbol_index]
                symbol = str(raw_symbol).strip().upper()
                if not symbol or symbol == "NONE":
                    continue
                if current_symbol is None:
                    current_symbol = symbol
                if symbol != current_symbol:
                    process_symbol(current_symbol, current_rows, columns)
                    processed_symbols.add(current_symbol)
                    current_rows.clear()
                    if symbol in processed_symbols:
                        raise ValueError(
                            "Regime_History must be grouped by Symbol; "
                            f"{symbol} appears in multiple blocks"
                        )
                    current_symbol = symbol
                current_rows.append(tuple(values))

            process_symbol(current_symbol, current_rows, columns)
            current_rows.clear()
        finally:
            workbook.close()

        report_df = pd.DataFrame(reports)
        if report_df.empty:
            raise RuntimeError(
                "Regime_History contained no usable symbol rows"
            )
        report_df.sort_values(
            by=["Status", "Symbol"],
            inplace=True,
        )
        report_df.reset_index(drop=True, inplace=True)

        selected_output_file = (
            output_file
            if output_file is not None
            else self.report_output_file
        )
        saved_file = None
        if self.save_excel_enabled and selected_output_file:
            saved_file = self._save_excel_report(
                report_df=report_df,
                output_file=selected_output_file,
                volume_data=volume_data,
            )
            print("ResultVolumeEngine.xlsx file is created")

        report_df.attrs["Excel_File"] = saved_file
        report_df.attrs["Volume_Backend"] = self.volume_backend
        report_df.attrs["Regime_Source_File"] = str(input_path)
        report_df.attrs["Regime_Source_Sheet"] = sheet_name
        report_df.attrs["Results_Retained_In_Memory"] = (
            self.retain_results_in_memory
        )

        if calculated_symbols == 0:
            raise RuntimeError(
                "Volume features could not be calculated for any symbol. "
                f"Diagnostic workbook: {saved_file}"
            )

        if not self.retain_results_in_memory:
            volume_data.clear()
            gc.collect()

        return volume_data, report_df

    # =====================================================
    # Calculate All Symbols
    # =====================================================

    def calculate_all(self, regime_data, output_file=None):

        if not isinstance(regime_data, dict):
            raise TypeError(
                "regime_data must be a dictionary "
                "containing {symbol: DataFrame}"
            )

        if not regime_data:
            raise ValueError(
                "regime_data is empty"
            )

        volume_data = {}
        reports = []

        for symbol, regime_df in regime_data.items():

            symbol = (
                str(symbol)
                .strip()
                .upper()
            )

            try:

                result_df, report = self.calculate_symbol(
                    symbol=symbol,
                    regime_df=regime_df
                )

                reports.append(report)

                if report["Status"] == "NOT_READY":
                    continue

                volume_data[symbol] = result_df

            except Exception as error:

                reports.append({
                    "Symbol": symbol,
                    "Status": "FAILED",
                    "Input_Rows": (
                        len(regime_df)
                        if isinstance(
                            regime_df,
                            pd.DataFrame
                        )
                        else 0
                    ),
                    "Volume_Ready_Rows": 0,
                    "Latest_DateTime": pd.NaT,
                    "Latest_Relative_Volume": np.nan,
                    "Latest_Time_Adjusted_RVOL": np.nan,
                    "Latest_Volume_Factor": np.nan,
                    "Latest_Volume_State": "NOT_READY",
                    "Latest_Stable_Regime": "NOT_READY",
                    "Latest_Volume_Confirmed": False,
                    "Volume_Backend": self.volume_backend,
                    "Error": str(error)
                })


        report_df = pd.DataFrame(reports)

        if not report_df.empty:

            report_df.sort_values(
                by=["Status", "Symbol"],
                inplace=True
            )

            report_df.reset_index(
                drop=True,
                inplace=True
            )

        selected_output_file = (
            output_file
            if output_file is not None
            else self.report_output_file
        )
        saved_file = None
        if self.save_excel_enabled and selected_output_file:
            saved_file = self._save_excel_report(
                report_df=report_df,
                output_file=selected_output_file,
                volume_data=volume_data,
            )
            print("ResultVolumeEngine.xlsx file is created")

        report_df.attrs["Excel_File"] = saved_file
        report_df.attrs["Volume_Backend"] = self.volume_backend

        if not volume_data:
            raise RuntimeError(
                "Volume features could not be calculated for any symbol. "
                f"Diagnostic workbook: {saved_file}"
            )

        report_df.attrs["Results_Retained_In_Memory"] = (
            self.retain_results_in_memory
        )
        if not self.retain_results_in_memory:
            volume_data.clear()
            gc.collect()

        return volume_data, report_df

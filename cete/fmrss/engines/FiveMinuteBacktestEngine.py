import os
import tempfile
from datetime import time
from pathlib import Path

import numpy as np
import pandas as pd


class FiveMinuteBacktestEngine:
    """Portfolio-aware intraday backtester for FMRSS risk plans."""

    REQUIRED_COLUMNS = [
        "Symbol", "DateTime", "Open", "High", "Low", "Close",
        "Stable_Regime", "Mathematical_Signal", "Signal_Emitted",
        "Signal_Strength", "Cross_Sectional_Rank", "Raw_Stop_Distance",
        "Risk_Trade_Accepted",
    ]

    LEDGER_COLUMNS = [
        "Trade_ID", "Symbol", "Regime", "Direction", "Signal_DateTime",
        "Entry_DateTime", "Exit_DateTime", "Raw_Entry_Open", "Entry_Price",
        "Raw_Exit_Price", "Exit_Price",
        "Stop_Loss", "Target_Price", "Quantity", "Position_Value",
        "Risk_Per_Share", "Initial_Risk", "Gross_PnL", "Brokerage",
        "STT", "Exchange_Charges", "SEBI_Charges", "GST", "Stamp_Duty",
        "Trading_Cost",
        "Net_PnL", "Return_Percent", "R_Multiple", "Exit_Reason",
        "Holding_Candles", "Signal_Strength", "Cross_Sectional_Rank",
        "Capital_Before", "Capital_After",
    ]

    REJECTION_COLUMNS = [
        "Symbol", "Signal_DateTime", "Planned_Entry_DateTime", "Reason"
    ]

    def __init__(
        self,
        initial_capital=100000.0,
        risk_per_trade_fraction=0.005,
        maximum_position_fraction=0.20,
        maximum_open_positions=5,
        reward_to_risk_ratio=2.0,
        tick_size=0.05,
        slippage_rate=0.0005,
        brokerage_rate=0.0003,
        brokerage_cap_per_order=20.0,
        charge_brokerage_on_buy=True,
        charge_brokerage_on_sell=True,
        stt_sell_rate=0.00025,
        exchange_transaction_rate=0.0000307,
        sebi_turnover_rate=0.000001,
        gst_rate=0.18,
        stamp_duty_buy_rate=0.00003,
        last_entry_time="14:45",
        forced_exit_time="15:15",
        maximum_holding_candles=72,
        conservative_same_candle=True,
        exit_on_opposite_signal=True,
        output_folder="/home/devinderjeet/fmrss/backtest",
        memory_optimized=True,
    ):
        positive = {
            "initial_capital": initial_capital,
            "risk_per_trade_fraction": risk_per_trade_fraction,
            "maximum_position_fraction": maximum_position_fraction,
            "reward_to_risk_ratio": reward_to_risk_ratio,
            "tick_size": tick_size,
        }
        for name, value in positive.items():
            if not np.isfinite(value) or value <= 0:
                raise ValueError(f"{name} must be finite and positive")

        rates = {
            "slippage_rate": slippage_rate,
            "brokerage_rate": brokerage_rate,
            "brokerage_cap_per_order": brokerage_cap_per_order,
            "stt_sell_rate": stt_sell_rate,
            "exchange_transaction_rate": exchange_transaction_rate,
            "sebi_turnover_rate": sebi_turnover_rate,
            "gst_rate": gst_rate,
            "stamp_duty_buy_rate": stamp_duty_buy_rate,
        }
        for name, value in rates.items():
            if not np.isfinite(value) or value < 0:
                raise ValueError(f"{name} must be finite and non-negative")

        if not isinstance(maximum_open_positions, (int, np.integer)):
            raise TypeError("maximum_open_positions must be an integer")
        if maximum_open_positions < 1:
            raise ValueError("maximum_open_positions must be at least 1")
        if not isinstance(maximum_holding_candles, (int, np.integer)):
            raise TypeError("maximum_holding_candles must be an integer")
        if maximum_holding_candles < 1:
            raise ValueError("maximum_holding_candles must be at least 1")
        if maximum_position_fraction * maximum_open_positions > 1.0 + 1e-12:
            raise ValueError(
                "maximum_position_fraction × maximum_open_positions "
                "cannot exceed 1.0"
            )

        self.initial_capital = float(initial_capital)
        self.risk_per_trade_fraction = float(risk_per_trade_fraction)
        self.maximum_position_fraction = float(maximum_position_fraction)
        self.maximum_open_positions = int(maximum_open_positions)
        self.reward_to_risk_ratio = float(reward_to_risk_ratio)
        self.tick_size = float(tick_size)
        self.slippage_rate = float(slippage_rate)
        self.brokerage_rate = float(brokerage_rate)
        self.brokerage_cap_per_order = float(brokerage_cap_per_order)
        self.charge_brokerage_on_buy = bool(charge_brokerage_on_buy)
        self.charge_brokerage_on_sell = bool(charge_brokerage_on_sell)
        self.stt_sell_rate = float(stt_sell_rate)
        self.exchange_transaction_rate = float(exchange_transaction_rate)
        self.sebi_turnover_rate = float(sebi_turnover_rate)
        self.gst_rate = float(gst_rate)
        self.stamp_duty_buy_rate = float(stamp_duty_buy_rate)
        self.last_entry_time = self._parse_time(last_entry_time)
        self.forced_exit_time = self._parse_time(forced_exit_time)
        if self.last_entry_time >= self.forced_exit_time:
            raise ValueError("last_entry_time must be before forced_exit_time")
        self.maximum_holding_candles = int(maximum_holding_candles)
        self.conservative_same_candle = bool(conservative_same_candle)
        self.exit_on_opposite_signal = bool(exit_on_opposite_signal)
        self.output_folder = Path(output_folder).expanduser()
        self.memory_optimized = bool(memory_optimized)
        self.calculation_backend = "NUMPY_PANDAS_EVENT_DRIVEN"

    @staticmethod
    def _parse_time(value):
        if isinstance(value, time):
            return value
        try:
            return pd.Timestamp(str(value)).time()
        except Exception as error:
            raise ValueError(f"Invalid time value: {value}") from error

    def _round_down(self, value):
        return np.floor((value + 1e-12) / self.tick_size) * self.tick_size

    def _round_up(self, value):
        return np.ceil((value - 1e-12) / self.tick_size) * self.tick_size

    def _validate_and_prepare(self, symbol, df):
        if not isinstance(df, pd.DataFrame) or df.empty:
            raise ValueError(f"{symbol} risk data must be a non-empty DataFrame")
        missing = [column for column in self.REQUIRED_COLUMNS if column not in df]
        if missing:
            raise ValueError(f"{symbol} is missing required columns: {missing}")

        # Only these columns are needed during simulation.
        result = df[self.REQUIRED_COLUMNS].copy()
        result["DateTime"] = pd.to_datetime(result["DateTime"], errors="coerce")
        for column in [
            "Open", "High", "Low", "Close", "Signal_Strength",
            "Cross_Sectional_Rank", "Raw_Stop_Distance",
        ]:
            result[column] = pd.to_numeric(result[column], errors="coerce")
        result.replace([np.inf, -np.inf], np.nan, inplace=True)
        result.dropna(subset=["DateTime"], inplace=True)
        result.sort_values("DateTime", inplace=True)
        result.drop_duplicates("DateTime", keep="last", inplace=True)
        result.reset_index(drop=True, inplace=True)
        result["Symbol"] = symbol
        result["Mathematical_Signal"] = (
            result["Mathematical_Signal"].astype("string").fillna("NO_TRADE")
            .str.strip().str.upper()
        )
        result["Stable_Regime"] = (
            result["Stable_Regime"].astype("string").fillna("NOT_READY")
            .str.strip().str.upper()
        )
        result["Signal_Emitted"] = result["Signal_Emitted"].eq(True)
        result["Risk_Trade_Accepted"] = result["Risk_Trade_Accepted"].eq(True)
        if self.memory_optimized:
            result["Symbol"] = result["Symbol"].astype("category")
            result["Mathematical_Signal"] = result["Mathematical_Signal"].astype("category")
            result["Stable_Regime"] = result["Stable_Regime"].astype("category")
            for column in ["Signal_Strength", "Cross_Sectional_Rank", "Raw_Stop_Distance"]:
                result[column] = result[column].astype(np.float32)
        if result.empty:
            raise ValueError(f"{symbol} has no usable rows")
        return result

    def _build_candidates(self, prepared_data):
        candidates, rejections = [], []
        for symbol, df in prepared_data.items():
            positions = np.flatnonzero(
                (df["Risk_Trade_Accepted"] & df["Signal_Emitted"]).to_numpy()
            )
            for position in positions:
                signal_row = df.iloc[position]
                signal_time = signal_row["DateTime"]
                if position + 1 >= len(df):
                    rejections.append({
                        "Symbol": symbol, "Signal_DateTime": signal_time,
                        "Planned_Entry_DateTime": pd.NaT,
                        "Reason": "NO_NEXT_CANDLE",
                    })
                    continue
                entry_row = df.iloc[position + 1]
                entry_time = entry_row["DateTime"]
                if entry_time.date() != signal_time.date():
                    rejections.append({
                        "Symbol": symbol, "Signal_DateTime": signal_time,
                        "Planned_Entry_DateTime": entry_time,
                        "Reason": "NEXT_CANDLE_IS_NEXT_SESSION",
                    })
                    continue
                if entry_time.time() > self.last_entry_time:
                    rejections.append({
                        "Symbol": symbol, "Signal_DateTime": signal_time,
                        "Planned_Entry_DateTime": entry_time,
                        "Reason": "ENTRY_AFTER_CUTOFF",
                    })
                    continue
                candidates.append({
                    "Symbol": symbol,
                    "Signal_Position": int(position),
                    "Entry_Position": int(position + 1),
                    "Signal_DateTime": signal_time,
                    "Entry_DateTime": entry_time,
                    "Direction": str(signal_row["Mathematical_Signal"]),
                    "Regime": str(signal_row["Stable_Regime"]),
                    "Signal_Strength": float(signal_row["Signal_Strength"]),
                    "Cross_Sectional_Rank": float(
                        signal_row["Cross_Sectional_Rank"]
                    ),
                    "Stop_Distance": float(signal_row["Raw_Stop_Distance"]),
                })
        candidates.sort(
            key=lambda row: (
                row["Entry_DateTime"], row["Cross_Sectional_Rank"],
                -row["Signal_Strength"], row["Symbol"],
            )
        )
        return candidates, rejections

    def _apply_entry_slippage(self, price, direction):
        multiplier = 1.0 + self.slippage_rate if direction == "BUY" else 1.0 - self.slippage_rate
        adjusted = price * multiplier
        return self._round_up(adjusted) if direction == "BUY" else self._round_down(adjusted)

    def _apply_exit_slippage(self, price, direction):
        multiplier = 1.0 - self.slippage_rate if direction == "BUY" else 1.0 + self.slippage_rate
        adjusted = price * multiplier
        return self._round_down(adjusted) if direction == "BUY" else self._round_up(adjusted)

    def _trading_cost(self, entry, exit_price, quantity, direction):
        entry_turnover = entry * quantity
        exit_turnover = exit_price * quantity
        buy_turnover = entry_turnover if direction == "BUY" else exit_turnover
        sell_turnover = exit_turnover if direction == "BUY" else entry_turnover
        buy_brokerage = min(
            self.brokerage_rate * buy_turnover, self.brokerage_cap_per_order
        ) if self.charge_brokerage_on_buy else 0.0
        sell_brokerage = min(
            self.brokerage_rate * sell_turnover, self.brokerage_cap_per_order
        ) if self.charge_brokerage_on_sell else 0.0
        brokerage = buy_brokerage + sell_brokerage
        turnover = entry_turnover + exit_turnover
        exchange = turnover * self.exchange_transaction_rate
        sebi = turnover * self.sebi_turnover_rate
        gst = (brokerage + exchange + sebi) * self.gst_rate
        stt = sell_turnover * self.stt_sell_rate
        stamp = buy_turnover * self.stamp_duty_buy_rate
        total = brokerage + exchange + sebi + gst + stt + stamp
        return {
            "Brokerage": brokerage,
            "STT": stt,
            "Exchange_Charges": exchange,
            "SEBI_Charges": sebi,
            "GST": gst,
            "Stamp_Duty": stamp,
            "Trading_Cost": total,
        }

    def _simulate_exit(self, df, candidate, entry_price, stop, target):
        start = candidate["Entry_Position"]
        direction = candidate["Direction"]
        session_date = candidate["Entry_DateTime"].date()
        last_position = min(len(df) - 1, start + self.maximum_holding_candles - 1)
        for position in range(start, last_position + 1):
            row = df.iloc[position]
            current_time = row["DateTime"]
            if current_time.date() != session_date:
                previous = df.iloc[position - 1]
                return position - 1, float(previous["Close"]), "SESSION_END"

            candle_open = float(row["Open"])
            high, low = float(row["High"]), float(row["Low"])

            # Gap logic applies from the candle after entry. On the entry
            # candle, stop and target are derived from that candle's open.
            if position > start:
                if direction == "BUY":
                    if candle_open <= stop:
                        return position, candle_open, "STOP_GAP"
                    if candle_open >= target:
                        return position, target, "TARGET_GAP"
                else:
                    if candle_open >= stop:
                        return position, candle_open, "STOP_GAP"
                    if candle_open <= target:
                        return position, target, "TARGET_GAP"
            if direction == "BUY":
                stop_hit, target_hit = low <= stop, high >= target
            else:
                stop_hit, target_hit = high >= stop, low <= target

            if stop_hit and target_hit:
                chosen = stop if self.conservative_same_candle else target
                reason = "STOP_AND_TARGET_STOP_FIRST" if self.conservative_same_candle else "STOP_AND_TARGET_TARGET_FIRST"
                return position, chosen, reason
            if stop_hit:
                return position, stop, "STOP_LOSS"
            if target_hit:
                return position, target, "TARGET"

            if self.exit_on_opposite_signal and position > start:
                row_signal = str(row["Mathematical_Signal"])
                opposite = (direction == "BUY" and row_signal == "SELL") or (
                    direction == "SELL" and row_signal == "BUY"
                )
                if bool(row["Signal_Emitted"]) and opposite:
                    return position, float(row["Close"]), "OPPOSITE_SIGNAL"

            if current_time.time() >= self.forced_exit_time:
                return position, float(row["Close"]), "FORCED_INTRADAY_EXIT"

        row = df.iloc[last_position]
        reason = "MAX_HOLDING_CANDLES"
        if last_position == len(df) - 1:
            reason = "END_OF_DATA"
        return last_position, float(row["Close"]), reason

    def _run_portfolio(self, prepared_data, candidates, rejections):
        trades, active, unsettled = [], [], []
        realised_capital = self.initial_capital

        for candidate in candidates:
            entry_time = candidate["Entry_DateTime"]

            still_unsettled = []
            for trade in unsettled:
                if trade["Exit_DateTime"] < entry_time:
                    realised_capital += trade["Net_PnL"]
                else:
                    still_unsettled.append(trade)
            unsettled = still_unsettled
            active = [trade for trade in active if trade["Exit_DateTime"] >= entry_time]

            symbol = candidate["Symbol"]
            rejection_base = {
                "Symbol": symbol,
                "Signal_DateTime": candidate["Signal_DateTime"],
                "Planned_Entry_DateTime": entry_time,
            }
            if any(trade["Symbol"] == symbol for trade in active):
                rejections.append({**rejection_base, "Reason": "SYMBOL_ALREADY_OPEN"})
                continue
            if len(active) >= self.maximum_open_positions:
                rejections.append({**rejection_base, "Reason": "PORTFOLIO_LIMIT"})
                continue

            df = prepared_data[symbol]
            entry_row = df.iloc[candidate["Entry_Position"]]
            raw_open = float(entry_row["Open"])
            stop_distance = candidate["Stop_Distance"]
            if not np.isfinite(raw_open) or raw_open <= 0 or not np.isfinite(stop_distance) or stop_distance <= 0:
                rejections.append({**rejection_base, "Reason": "INVALID_ENTRY_OR_STOP"})
                continue

            direction = candidate["Direction"]
            entry = self._apply_entry_slippage(raw_open, direction)
            if direction == "BUY":
                stop = self._round_down(entry - stop_distance)
                target = self._round_down(entry + self.reward_to_risk_ratio * stop_distance)
            else:
                stop = self._round_up(entry + stop_distance)
                target = self._round_up(entry - self.reward_to_risk_ratio * stop_distance)

            risk_per_share = abs(entry - stop)
            if risk_per_share <= 0 or target <= 0:
                rejections.append({**rejection_base, "Reason": "INVALID_ROUNDED_LEVELS"})
                continue

            risk_budget = realised_capital * self.risk_per_trade_fraction
            position_cap = realised_capital * self.maximum_position_fraction
            risk_quantity = int(np.floor(risk_budget / risk_per_share))
            capital_quantity = int(np.floor(position_cap / entry))
            quantity = min(risk_quantity, capital_quantity)
            if quantity < 1:
                rejections.append({**rejection_base, "Reason": "QUANTITY_ZERO"})
                continue

            exit_position, raw_exit, exit_reason = self._simulate_exit(
                df, candidate, entry, stop, target
            )
            exit_row = df.iloc[exit_position]
            # Profit targets are limit orders and are conservatively capped at
            # the target. Stops and close-based exits receive adverse slippage.
            if exit_reason in {
                "TARGET", "TARGET_GAP", "STOP_AND_TARGET_TARGET_FIRST"
            }:
                exit_price = raw_exit
            else:
                exit_price = self._apply_exit_slippage(raw_exit, direction)

            gross_pnl = (
                (exit_price - entry) * quantity
                if direction == "BUY"
                else (entry - exit_price) * quantity
            )
            cost = self._trading_cost(entry, exit_price, quantity, direction)
            net_pnl = gross_pnl - cost["Trading_Cost"]
            initial_risk = risk_per_share * quantity
            trade = {
                "Trade_ID": len(trades) + 1,
                "Symbol": symbol,
                "Regime": candidate["Regime"],
                "Direction": direction,
                "Signal_DateTime": candidate["Signal_DateTime"],
                "Entry_DateTime": entry_time,
                "Exit_DateTime": exit_row["DateTime"],
                "Raw_Entry_Open": raw_open,
                "Entry_Price": entry,
                "Raw_Exit_Price": raw_exit,
                "Exit_Price": exit_price,
                "Stop_Loss": stop,
                "Target_Price": target,
                "Quantity": quantity,
                "Position_Value": entry * quantity,
                "Risk_Per_Share": risk_per_share,
                "Initial_Risk": initial_risk,
                "Gross_PnL": gross_pnl,
                **cost,
                "Net_PnL": net_pnl,
                "Return_Percent": 100.0 * net_pnl / (entry * quantity),
                "R_Multiple": net_pnl / initial_risk,
                "Exit_Reason": exit_reason,
                "Holding_Candles": exit_position - candidate["Entry_Position"] + 1,
                "Signal_Strength": candidate["Signal_Strength"],
                "Cross_Sectional_Rank": candidate["Cross_Sectional_Rank"],
                "Capital_Before": realised_capital,
                "Capital_After": np.nan,
            }
            trades.append(trade)
            active.append(trade)
            unsettled.append(trade)

        # Capital_After is assigned in realised exit order, which is the valid
        # equity sequence when trades overlap.
        capital = self.initial_capital
        for trade in sorted(trades, key=lambda row: (row["Exit_DateTime"], row["Trade_ID"])):
            capital += trade["Net_PnL"]
            trade["Capital_After"] = capital
        return trades, rejections

    def _analytics(self, ledger):
        if ledger.empty:
            equity = pd.DataFrame({
                "DateTime": [pd.NaT], "Trade_ID": [0],
                "Net_PnL": [0.0], "Equity": [self.initial_capital],
                "Peak_Equity": [self.initial_capital], "Drawdown": [0.0],
                "Drawdown_Percent": [0.0],
            })
            empty = pd.DataFrame()
            return equity, self._empty_summary(), empty, empty, empty, empty, empty, empty

        exits = ledger.sort_values(["Exit_DateTime", "Trade_ID"]).reset_index(drop=True)
        equity = exits[["Exit_DateTime", "Trade_ID", "Net_PnL"]].rename(
            columns={"Exit_DateTime": "DateTime"}
        ).copy()
        equity["Equity"] = self.initial_capital + equity["Net_PnL"].cumsum()
        equity["Peak_Equity"] = equity["Equity"].cummax().clip(lower=self.initial_capital)
        equity["Drawdown"] = equity["Equity"] - equity["Peak_Equity"]
        equity["Drawdown_Percent"] = 100.0 * equity["Drawdown"] / equity["Peak_Equity"]

        wins = ledger["Net_PnL"] > 0
        gross_profit = ledger.loc[wins, "Net_PnL"].sum()
        gross_loss = -ledger.loc[~wins, "Net_PnL"].sum()
        profit_factor = gross_profit / gross_loss if gross_loss > 0 else np.inf

        daily = equity.copy()
        daily["Date"] = daily["DateTime"].dt.normalize()
        daily_equity = daily.groupby("Date", sort=True)["Equity"].last()
        daily_returns = daily_equity.pct_change().dropna()
        sharpe = np.nan
        sortino = np.nan
        if len(daily_returns) >= 2 and daily_returns.std(ddof=1) > 0:
            sharpe = np.sqrt(252.0) * daily_returns.mean() / daily_returns.std(ddof=1)
        downside = daily_returns[daily_returns < 0]
        if len(downside) >= 2 and downside.std(ddof=1) > 0:
            sortino = np.sqrt(252.0) * daily_returns.mean() / downside.std(ddof=1)

        summary = pd.DataFrame({
            "Metric": [
                "Initial Capital", "Ending Capital", "Net Profit",
                "Return Percent", "Total Trades", "Winning Trades",
                "Losing Trades", "Win Rate Percent", "Average Net PnL",
                "Expectancy (R)", "Profit Factor", "Maximum Drawdown",
                "Maximum Drawdown Percent", "Annualized Daily Sharpe",
                "Annualized Daily Sortino", "Total Trading Costs",
                "Average Holding Candles",
            ],
            "Value": [
                self.initial_capital, equity["Equity"].iloc[-1], ledger["Net_PnL"].sum(),
                100.0 * ledger["Net_PnL"].sum() / self.initial_capital,
                len(ledger), int(wins.sum()), int((~wins).sum()),
                100.0 * wins.mean(), ledger["Net_PnL"].mean(),
                ledger["R_Multiple"].mean(), profit_factor,
                equity["Drawdown"].min(), equity["Drawdown_Percent"].min(),
                sharpe, sortino, ledger["Trading_Cost"].sum(),
                ledger["Holding_Candles"].mean(),
            ],
        })

        symbol_perf = self._group_performance(ledger, "Symbol")
        regime_perf = self._group_performance(ledger, "Regime")
        monthly_source = ledger.copy()
        monthly_source["Month"] = monthly_source["Exit_DateTime"].dt.to_period("M").astype(str)
        monthly_perf = self._group_performance(monthly_source, "Month")
        direction_perf = self._group_performance(ledger, "Direction")
        exit_perf = self._group_performance(ledger, "Exit_Reason")
        daily_source = ledger.copy()
        daily_source["Date"] = daily_source["Exit_DateTime"].dt.strftime("%Y-%m-%d")
        daily_perf = self._group_performance(daily_source, "Date")
        return (
            equity, summary, symbol_perf, regime_perf, monthly_perf,
            direction_perf, exit_perf, daily_perf,
        )

    def _empty_summary(self):
        return pd.DataFrame({
            "Metric": ["Initial Capital", "Ending Capital", "Net Profit", "Total Trades"],
            "Value": [self.initial_capital, self.initial_capital, 0.0, 0],
        })

    @staticmethod
    def _group_performance(df, group_column):
        grouped = df.groupby(group_column, observed=True, sort=True)
        result = grouped.agg(
            Trades=("Trade_ID", "count"),
            Net_PnL=("Net_PnL", "sum"),
            Average_PnL=("Net_PnL", "mean"),
            Average_R=("R_Multiple", "mean"),
            Trading_Cost=("Trading_Cost", "sum"),
        ).reset_index()
        win_rate = grouped["Net_PnL"].apply(lambda values: 100.0 * (values > 0).mean())
        result["Win_Rate_Percent"] = result[group_column].map(win_rate)
        return result

    def _maximum_concurrent_positions(self, ledger):
        if ledger.empty:
            return 0
        events = []
        for row in ledger.itertuples(index=False):
            # Entries are processed before exits at an identical timestamp;
            # a trade exiting inside that candle still occupied a slot at open.
            events.append((row.Entry_DateTime, 0, 1))
            events.append((row.Exit_DateTime, 1, -1))
        running = maximum = 0
        for _, _, change in sorted(events):
            running += change
            maximum = max(maximum, running)
        return maximum

    def _build_checks(self, ledger, equity):
        ending_equity = (
            float(equity["Equity"].iloc[-1])
            if not equity.empty else self.initial_capital
        )
        net_pnl = float(ledger["Net_PnL"].sum()) if not ledger.empty else 0.0
        expected_equity = self.initial_capital + net_pnl
        equity_difference = ending_equity - expected_equity
        maximum_concurrent = self._maximum_concurrent_positions(ledger)

        if ledger.empty:
            positive_quantities = True
            nonnegative_costs = True
            intraday_only = True
            valid_geometry = True
        else:
            positive_quantities = bool(ledger["Quantity"].gt(0).all())
            nonnegative_costs = bool(ledger["Trading_Cost"].ge(0).all())
            intraday_only = bool(
                ledger["Entry_DateTime"].dt.date.eq(
                    ledger["Exit_DateTime"].dt.date
                ).all()
            )
            buy = ledger["Direction"].eq("BUY")
            sell = ledger["Direction"].eq("SELL")
            valid_geometry = bool((
                (
                    ~buy
                    | (
                        ledger["Stop_Loss"].lt(ledger["Entry_Price"])
                        & ledger["Target_Price"].gt(ledger["Entry_Price"])
                    )
                )
                & (
                    ~sell
                    | (
                        ledger["Stop_Loss"].gt(ledger["Entry_Price"])
                        & ledger["Target_Price"].lt(ledger["Entry_Price"])
                    )
                )
            ).all())

        rows = [
            (
                "Equity reconciliation", ending_equity, expected_equity,
                equity_difference, 0.01, abs(equity_difference) <= 0.01,
                "Ending equity must equal initial capital plus net P&L.",
            ),
            (
                "Maximum concurrent positions", maximum_concurrent,
                self.maximum_open_positions,
                maximum_concurrent - self.maximum_open_positions, 0,
                maximum_concurrent <= self.maximum_open_positions,
                "Actual concurrency cannot exceed the configured limit.",
            ),
            (
                "Positive quantities", int(positive_quantities), 1,
                int(positive_quantities) - 1, 0, positive_quantities,
                "Every executed trade must have quantity greater than zero.",
            ),
            (
                "Non-negative costs", int(nonnegative_costs), 1,
                int(nonnegative_costs) - 1, 0, nonnegative_costs,
                "Every trading-cost total must be non-negative.",
            ),
            (
                "Intraday exits", int(intraday_only), 1,
                int(intraday_only) - 1, 0, intraday_only,
                "Entry and exit dates must match.",
            ),
            (
                "Stop/target geometry", int(valid_geometry), 1,
                int(valid_geometry) - 1, 0, valid_geometry,
                "BUY and SELL stop/target levels must be directionally valid.",
            ),
        ]
        checks = pd.DataFrame(
            rows,
            columns=[
                "Check", "Actual", "Expected", "Difference", "Tolerance",
                "Passed", "Notes",
            ],
        )
        checks["Status"] = np.where(checks["Passed"], "OK", "FAILED")
        return checks[
            [
                "Check", "Actual", "Expected", "Difference", "Tolerance",
                "Status", "Notes",
            ]
        ]

    @staticmethod
    def _style_workbook(writer):
        from openpyxl.chart import LineChart, Reference
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        header_fill = PatternFill("solid", fgColor="17365D")
        green_fill = PatternFill("solid", fgColor="C6EFCE")
        red_fill = PatternFill("solid", fgColor="FFC7CE")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.fill = header_fill
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for cells in sheet.columns:
                values = [str(cell.value) for cell in cells if cell.value is not None]
                width = min(max([len(value) for value in values] + [10]) + 2, 28)
                sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
            headers = {cell.value: cell.column for cell in sheet[1]}
            for name in ["Signal_DateTime", "Entry_DateTime", "Exit_DateTime", "DateTime"]:
                if name in headers:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, headers[name]).number_format = "yyyy-mm-dd hh:mm:ss"
            for name in ["Entry_Price", "Exit_Price", "Stop_Loss", "Target_Price", "Gross_PnL", "Trading_Cost", "Net_PnL", "Equity", "Drawdown"]:
                if name in headers:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, headers[name]).number_format = "₹#,##0.00;[Red](₹#,##0.00);-"
            pnl_column = headers.get("Net_PnL")
            if pnl_column and sheet.max_row >= 2:
                letter = get_column_letter(pnl_column)
                target = f"{letter}2:{letter}{sheet.max_row}"
                sheet.conditional_formatting.add(target, CellIsRule(operator="greaterThan", formula=["0"], fill=green_fill))
                sheet.conditional_formatting.add(target, CellIsRule(operator="lessThan", formula=["0"], fill=red_fill))

            status_column = headers.get("Status")
            if status_column and sheet.max_row >= 2:
                letter = get_column_letter(status_column)
                target = f"{letter}2:{letter}{sheet.max_row}"
                sheet.conditional_formatting.add(
                    target,
                    CellIsRule(operator="equal", formula=['"OK"'], fill=green_fill),
                )
                sheet.conditional_formatting.add(
                    target,
                    CellIsRule(operator="equal", formula=['"FAILED"'], fill=red_fill),
                )

            if sheet.title == "Equity_Curve" and sheet.max_row >= 2:
                equity_column = headers.get("Equity")
                date_column = headers.get("DateTime")
                if equity_column and date_column:
                    chart = LineChart()
                    chart.title = "FMRSS Equity Curve"
                    chart.y_axis.title = "Equity (INR)"
                    chart.x_axis.title = "Exit time"
                    chart.height = 8
                    chart.width = 16
                    chart.add_data(
                        Reference(
                            sheet,
                            min_col=equity_column,
                            min_row=1,
                            max_row=sheet.max_row,
                        ),
                        titles_from_data=True,
                    )
                    chart.set_categories(
                        Reference(
                            sheet,
                            min_col=date_column,
                            min_row=2,
                            max_row=sheet.max_row,
                        )
                    )
                    chart.legend = None
                    sheet.add_chart(chart, "I2")

    def _save_excel(
        self, output_path, summary, ledger, equity, symbol_perf, regime_perf,
        monthly_perf, direction_perf, exit_perf, daily_perf, rejections, checks,
    ):
        output_path.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(
            prefix="fmrss_backtest_", suffix=".xlsx", dir=output_path.parent
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(temporary_name, engine="openpyxl") as writer:
                summary.to_excel(writer, sheet_name="Summary", index=False)
                ledger.to_excel(writer, sheet_name="Trade_Ledger", index=False)
                equity.to_excel(writer, sheet_name="Equity_Curve", index=False)
                symbol_perf.to_excel(writer, sheet_name="Symbol_Performance", index=False)
                regime_perf.to_excel(writer, sheet_name="Regime_Performance", index=False)
                direction_perf.to_excel(writer, sheet_name="Direction_Performance", index=False)
                exit_perf.to_excel(writer, sheet_name="Exit_Performance", index=False)
                daily_perf.to_excel(writer, sheet_name="Daily_Performance", index=False)
                monthly_perf.to_excel(writer, sheet_name="Monthly_Performance", index=False)
                rejections.to_excel(writer, sheet_name="Rejections", index=False)
                checks.to_excel(writer, sheet_name="Checks", index=False)
                assumptions = pd.DataFrame({
                    "Parameter": [
                        "Initial Capital", "Risk Per Trade Fraction",
                        "Maximum Position Fraction", "Maximum Open Positions",
                        "Reward/Risk Ratio", "Slippage Rate", "Brokerage Rate",
                        "Brokerage Cap Per Order", "STT Sell Rate",
                        "Exchange Transaction Rate", "SEBI Turnover Rate",
                        "GST Rate", "Stamp Duty Buy Rate", "Same-Candle Rule",
                        "Stop Gap Rule", "Target Gap Rule", "Last Entry Time",
                        "Forced Exit Time", "Maximum Holding Candles",
                    ],
                    "Value": [
                        self.initial_capital, self.risk_per_trade_fraction,
                        self.maximum_position_fraction, self.maximum_open_positions,
                        self.reward_to_risk_ratio, self.slippage_rate,
                        self.brokerage_rate, self.brokerage_cap_per_order,
                        self.stt_sell_rate, self.exchange_transaction_rate,
                        self.sebi_turnover_rate, self.gst_rate,
                        self.stamp_duty_buy_rate,
                        "STOP_FIRST" if self.conservative_same_candle else "TARGET_FIRST",
                        "ADVERSE OPEN PLUS SLIPPAGE", "TARGET LIMIT",
                        self.last_entry_time.strftime("%H:%M"),
                        self.forced_exit_time.strftime("%H:%M"),
                        self.maximum_holding_candles,
                    ],
                })
                assumptions.to_excel(writer, sheet_name="Assumptions", index=False)
                self._style_workbook(writer)
            os.replace(temporary_name, output_path)
        except Exception:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)
            raise

    def run(self, risk_data, output_filename="FMRSS_Backtest_Report.xlsx"):
        if not isinstance(risk_data, dict) or not risk_data:
            raise ValueError("risk_data must be a non-empty {symbol: DataFrame} dictionary")

        print("=" * 120)
        print("FMRSS : FIVE-MINUTE PORTFOLIO BACKTEST")
        print(f"Backend : {self.calculation_backend}")
        print("=" * 120)

        prepared, preparation_failures = {}, []
        for raw_symbol, df in risk_data.items():
            symbol = str(raw_symbol).strip().upper()
            try:
                prepared[symbol] = self._validate_and_prepare(symbol, df)
            except Exception as error:
                preparation_failures.append({
                    "Symbol": symbol, "Signal_DateTime": pd.NaT,
                    "Planned_Entry_DateTime": pd.NaT,
                    "Reason": f"PREPARATION_FAILED: {error}",
                })
        if not prepared:
            raise RuntimeError("No symbols could be prepared for backtesting")

        candidates, rejections = self._build_candidates(prepared)
        rejections.extend(preparation_failures)
        trades, rejections = self._run_portfolio(prepared, candidates, rejections)

        ledger = pd.DataFrame(trades, columns=self.LEDGER_COLUMNS)
        rejection_df = pd.DataFrame(rejections, columns=self.REJECTION_COLUMNS)
        (
            equity, summary, symbol_perf, regime_perf, monthly_perf,
            direction_perf, exit_perf, daily_perf,
        ) = self._analytics(ledger)
        checks = self._build_checks(ledger, equity)

        additional_summary = pd.DataFrame({
            "Metric": [
                "Candidate Signals", "Rejected Candidates",
                "Maximum Concurrent Positions", "Model Check Status",
            ],
            "Value": [
                len(candidates), len(rejection_df),
                self._maximum_concurrent_positions(ledger),
                "OK" if checks["Status"].eq("OK").all() else "FAILED",
            ],
        })
        summary = pd.concat([summary, additional_summary], ignore_index=True)

        output_path = self.output_folder / output_filename
        self._save_excel(
            output_path, summary, ledger, equity, symbol_perf,
            regime_perf, monthly_perf, direction_perf, exit_perf,
            daily_perf, rejection_df, checks,
        )

        summary.attrs["Output_File"] = str(output_path)
        summary.attrs["Checks_Passed"] = bool(checks["Status"].eq("OK").all())

        print(f"Prepared symbols : {len(prepared)}/{len(risk_data)}")
        print(f"Candidate signals : {len(candidates)}")
        print(f"Executed trades : {len(ledger)}")
        print(f"Rejected candidates : {len(rejection_df)}")
        print(f"Backtest report : {output_path}")
        print("=" * 120)
        return ledger, equity, summary

    def calculate_all(self, risk_data, output_filename="FMRSS_Backtest_Report.xlsx"):
        return self.run(risk_data, output_filename)

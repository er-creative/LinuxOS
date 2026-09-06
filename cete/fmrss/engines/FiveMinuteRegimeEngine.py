import gc
import os
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd


try:
    from engines._regime_rolling import rolling_lag1_autocorrelation
    CYTHON_REGIME_AVAILABLE = True
except ImportError:
    try:
        from _regime_rolling import rolling_lag1_autocorrelation
        CYTHON_REGIME_AVAILABLE = True
    except ImportError:
        rolling_lag1_autocorrelation = None
        CYTHON_REGIME_AVAILABLE = False


class FiveMinuteRegimeEngine:

    # =====================================================
    # Required Input Columns
    # =====================================================

    REQUIRED_COLUMNS = [
        "Symbol",
        "DateTime",
        "Residual_Return",
        "Residual_Features_Ready"
    ]

    VALID_REGIMES = {
        "MOMENTUM",
        "MEAN_REVERSION",
        "RANDOM",
        "NOT_READY"
    }

    # =====================================================
    # Constructor
    # =====================================================

    def __init__(
        self,
        autocorrelation_window=60,
        autocorrelation_min_periods=50,
        momentum_threshold=0.10,
        mean_reversion_threshold=-0.10,
        stability_candles=3,
        strength_reference=0.30,
        numerical_epsilon=1e-12,
        use_cython=True,
        memory_optimized=True,
        retain_results_in_memory=False,
        report_output_file=(
            "/home/devinderjeet/fmrss/report/"
            "RegimeEngine.xlsx"
        ),
        save_excel=True,
    ):

        if autocorrelation_window < 10:
            raise ValueError(
                "autocorrelation_window must be at least 10"
            )

        if not 3 <= autocorrelation_min_periods <= autocorrelation_window:
            raise ValueError(
                "autocorrelation_min_periods must be between "
                "3 and autocorrelation_window"
            )

        if momentum_threshold <= 0:
            raise ValueError(
                "momentum_threshold must be greater than zero"
            )

        if mean_reversion_threshold >= 0:
            raise ValueError(
                "mean_reversion_threshold must be less than zero"
            )

        if stability_candles < 1:
            raise ValueError(
                "stability_candles must be at least 1"
            )

        if not 0 < strength_reference <= 1:
            raise ValueError(
                "strength_reference must be between 0 and 1"
            )

        if numerical_epsilon <= 0:
            raise ValueError(
                "numerical_epsilon must be positive"
            )

        self.autocorrelation_window = int(
            autocorrelation_window
        )

        self.autocorrelation_min_periods = int(
            autocorrelation_min_periods
        )

        self.momentum_threshold = float(
            momentum_threshold
        )

        self.mean_reversion_threshold = float(
            mean_reversion_threshold
        )

        self.stability_candles = int(
            stability_candles
        )

        self.strength_reference = float(
            strength_reference
        )

        self.numerical_epsilon = float(
            numerical_epsilon
        )

        self.use_cython = bool(use_cython)
        self.memory_optimized = bool(memory_optimized)
        self.retain_results_in_memory = bool(
            retain_results_in_memory
        )
        self.report_output_file = (
            str(report_output_file)
            if report_output_file
            else None
        )
        self.save_excel_enabled = bool(save_excel)
        self.regime_backend = (
            "CYTHON"
            if self.use_cython and CYTHON_REGIME_AVAILABLE
            else "PYTHON"
        )

    # =====================================================
    # Validate Input
    # =====================================================

    def _validate_input(self, symbol, df):

        if not isinstance(df, pd.DataFrame):
            raise TypeError(
                f"{symbol} data must be a pandas DataFrame"
            )

        if df.empty:
            raise ValueError(
                f"{symbol} residual data is empty"
            )

        missing_columns = [
            column
            for column in self.REQUIRED_COLUMNS
            if column not in df.columns
        ]

        if missing_columns:
            raise ValueError(
                f"{symbol} is missing required columns: "
                f"{missing_columns}"
            )

    # =====================================================
    # Prepare Input
    # =====================================================

    def _prepare_input(self, symbol, df):

        self._validate_input(
            symbol=symbol,
            df=df
        )

        # A shallow frame shares the immutable upstream feature blocks while
        # new regime columns receive their own arrays.  This substantially
        # reduces peak memory when residual_data and regime_data coexist.
        result = df.copy(deep=not self.memory_optimized)

        result["DateTime"] = pd.to_datetime(
            result["DateTime"],
            errors="coerce"
        )

        result["Residual_Return"] = pd.to_numeric(
            result["Residual_Return"],
            errors="coerce"
        )

        result["Residual_Return"] = (
            result["Residual_Return"]
            .replace(
                [np.inf, -np.inf],
                np.nan
            )
        )

        result["Residual_Features_Ready"] = (
            result["Residual_Features_Ready"]
            .fillna(False)
            .astype(bool)
        )

        result.dropna(
            subset=["DateTime"],
            inplace=True
        )

        result["Symbol"] = (
            str(symbol)
            .strip()
            .upper()
        )

        result.sort_values(
            "DateTime",
            inplace=True
        )

        result.drop_duplicates(
            subset=["DateTime"],
            keep="last",
            inplace=True
        )

        result.reset_index(
            drop=True,
            inplace=True
        )

        if result.empty:
            raise ValueError(
                f"{symbol} has no valid timestamped rows"
            )

        return result

    # =====================================================
    # Calculate Lag-One Residual Autocorrelation
    # =====================================================

    def _calculate_autocorrelation(self, df):

        residual = df["Residual_Return"]

        if self.regime_backend == "CYTHON":
            values = np.require(
                residual.to_numpy(dtype=np.float64, copy=False),
                dtype=np.float64,
                requirements=["C", "A", "W"],
            )
            correlation_values, observation_values = (
                rolling_lag1_autocorrelation(
                    values,
                    self.autocorrelation_window,
                    self.autocorrelation_min_periods,
                    self.numerical_epsilon,
                )
            )
            rolling_autocorrelation = pd.Series(
                correlation_values,
                index=df.index,
                dtype=np.float64,
            )
            observation_count = pd.Series(
                observation_values,
                index=df.index,
                dtype=np.int32,
            )
        else:
            lagged_residual = residual.shift(1)

            # shift(1) excludes the current candle and prevents look-ahead.
            rolling_autocorrelation = (
                residual
                .rolling(
                    window=self.autocorrelation_window,
                    min_periods=self.autocorrelation_min_periods
                )
                .corr(lagged_residual)
                .shift(1)
            )

            valid_pair = (
                residual.notna()
                & lagged_residual.notna()
            ).astype(np.int8)

            observation_count = (
                valid_pair
                .rolling(
                    window=self.autocorrelation_window,
                    min_periods=1
                )
                .sum()
                .shift(1)
                .fillna(0)
                .astype(np.int32)
            )

        autocorrelation = rolling_autocorrelation.clip(
            lower=-1.0,
            upper=1.0,
        )
        if self.memory_optimized:
            autocorrelation = autocorrelation.astype(
                np.float32,
                copy=False,
            )
        df["Residual_Autocorrelation"] = autocorrelation

        df["Autocorrelation_Observations"] = (
            observation_count
        )

        return df

    # =====================================================
    # Calculate Diagnostic T-Statistic
    # =====================================================

    def _calculate_autocorrelation_diagnostics(self, df):

        rho = df["Residual_Autocorrelation"]

        observations = (
            df["Autocorrelation_Observations"]
            .astype(np.float64)
        )

        denominator = np.maximum(
            1.0 - np.square(rho),
            self.numerical_epsilon
        )

        valid = (
            rho.notna()
            & (
                observations
                >= self.autocorrelation_min_periods
            )
            & (observations > 2)
        )

        t_statistic = np.where(
            valid,
            rho * np.sqrt(
                (observations - 2.0)
                / denominator
            ),
            np.nan
        )

        df["Autocorrelation_TStatistic"] = (
            t_statistic
        )
        if self.memory_optimized:
            df["Autocorrelation_TStatistic"] = df[
                "Autocorrelation_TStatistic"
            ].astype(np.float32, copy=False)

        # This is diagnostic information. It does not alter the initial
        # regime rules, which use the configured rho thresholds.
        df["Autocorrelation_Significant_5Pct"] = (
            pd.Series(
                t_statistic,
                index=df.index
            )
            .abs()
            .ge(1.96)
            & valid
        )

        return df

    # =====================================================
    # Classify Raw Regime
    # =====================================================

    def _classify_raw_regime(self, df):

        rho = df["Residual_Autocorrelation"]

        regime_ready = (
            df["Residual_Features_Ready"]
            & rho.notna()
            & (
                df["Autocorrelation_Observations"]
                >= self.autocorrelation_min_periods
            )
        )

        df["Regime_Features_Ready"] = (
            regime_ready
        )

        df["Raw_Regime"] = np.select(
            condlist=[
                ~regime_ready,
                rho >= self.momentum_threshold,
                rho <= self.mean_reversion_threshold
            ],
            choicelist=[
                "NOT_READY",
                "MOMENTUM",
                "MEAN_REVERSION"
            ],
            default="RANDOM"
        )

        return df

    # =====================================================
    # Calculate Regime Stability and Duration
    # =====================================================

    def _calculate_regime_stability(self, df):

        regime_changed = (
            df["Raw_Regime"]
            .ne(df["Raw_Regime"].shift(1))
        )

        regime_group = regime_changed.cumsum()

        duration = (
            df.groupby(
                regime_group,
                sort=False
            )
            .cumcount()
            .add(1)
        )

        duration = duration.where(
            df["Regime_Features_Ready"],
            0
        )

        df["Regime_Duration"] = (
            duration.astype(np.int32)
        )

        df["Regime_Stable"] = (
            df["Regime_Features_Ready"]
            & (
                df["Regime_Duration"]
                >= self.stability_candles
            )
        )

        df["Stable_Regime"] = np.select(
            condlist=[
                ~df["Regime_Features_Ready"],
                df["Regime_Stable"]
            ],
            choicelist=[
                "NOT_READY",
                df["Raw_Regime"]
            ],
            default="UNSTABLE"
        )

        if self.memory_optimized:
            df["Raw_Regime"] = df["Raw_Regime"].astype("category")
            df["Stable_Regime"] = df["Stable_Regime"].astype("category")

        return df

    # =====================================================
    # Calculate Regime Strength
    # =====================================================

    def _calculate_regime_strength(self, df):

        absolute_rho = (
            df["Residual_Autocorrelation"]
            .abs()
        )

        strength = (
            100.0
            * absolute_rho
            / self.strength_reference
        ).clip(
            lower=0.0,
            upper=100.0
        )

        df["Regime_Strength"] = strength.where(
            df["Regime_Features_Ready"],
            np.nan
        )
        if self.memory_optimized:
            df["Regime_Strength"] = df["Regime_Strength"].astype(
                np.float32,
                copy=False,
            )

        df["Regime_Trade_Eligible"] = (
            df["Regime_Stable"]
            & df["Stable_Regime"].isin(
                [
                    "MOMENTUM",
                    "MEAN_REVERSION"
                ]
            )
        )

        return df

    # =====================================================
    # Calculate One Symbol
    # =====================================================

    def calculate_symbol(
        self,
        symbol,
        residual_df
    ):

        symbol = (
            str(symbol)
            .strip()
            .upper()
        )

        df = self._prepare_input(
            symbol=symbol,
            df=residual_df
        )

        df = self._calculate_autocorrelation(df)

        df = (
            self._calculate_autocorrelation_diagnostics(df)
        )

        df = self._classify_raw_regime(df)

        df = self._calculate_regime_stability(df)

        df = self._calculate_regime_strength(df)

        ready_rows = int(
            df["Regime_Features_Ready"].sum()
        )

        latest = df.iloc[-1]

        report = {
            "Symbol": symbol,
            "Status": (
                "CALCULATED"
                if ready_rows > 0
                else "NOT_READY"
            ),
            "Input_Rows": len(df),
            "Regime_Ready_Rows": ready_rows,
            "Latest_DateTime": latest["DateTime"],
            "Latest_Autocorrelation": (
                latest["Residual_Autocorrelation"]
            ),
            "Latest_TStatistic": (
                latest["Autocorrelation_TStatistic"]
            ),
            "Latest_Raw_Regime": latest["Raw_Regime"],
            "Latest_Stable_Regime": latest["Stable_Regime"],
            "Latest_Regime_Stable": bool(
                latest["Regime_Stable"]
            ),
            "Latest_Regime_Duration": int(
                latest["Regime_Duration"]
            ),
            "Latest_Regime_Strength": (
                latest["Regime_Strength"]
            ),
            "Latest_Trade_Eligible": bool(
                latest["Regime_Trade_Eligible"]
            ),
            "Regime_Backend": self.regime_backend,
            "Error": "",
        }

        return df, report

    # =====================================================
    # Excel Report
    # =====================================================

    @staticmethod
    def _excel_safe(frame):
        safe = frame.copy(deep=False).replace([np.inf, -np.inf], np.nan)
        for column in safe.columns:
            if isinstance(safe[column].dtype, pd.CategoricalDtype):
                safe[column] = safe[column].astype("string")
            if isinstance(safe[column].dtype, pd.DatetimeTZDtype):
                safe[column] = safe[column].dt.tz_localize(None)
        return safe

    def _build_summary(self, report_df):
        status = report_df["Status"].astype("string")
        regime = report_df["Latest_Stable_Regime"].astype("string")
        eligible = report_df["Latest_Trade_Eligible"].fillna(False)
        return pd.DataFrame(
            [
                ["RUN", "Report generated", pd.Timestamp.now()],
                ["RUN", "Regime backend", self.regime_backend],
                ["RUN", "Memory optimized", self.memory_optimized],
                [
                    "RUN",
                    "Results retained in memory",
                    self.retain_results_in_memory,
                ],
                ["COUNTS", "Total symbols", int(len(report_df))],
                ["COUNTS", "Calculated symbols", int(status.eq("CALCULATED").sum())],
                ["COUNTS", "Not-ready symbols", int(status.eq("NOT_READY").sum())],
                ["COUNTS", "Failed symbols", int(status.eq("FAILED").sum())],
                ["COUNTS", "Trade-eligible symbols", int(eligible.sum())],
                ["REGIME", "MOMENTUM symbols", int(regime.eq("MOMENTUM").sum())],
                ["REGIME", "MEAN_REVERSION symbols", int(regime.eq("MEAN_REVERSION").sum())],
                ["REGIME", "RANDOM symbols", int(regime.eq("RANDOM").sum())],
                ["REGIME", "UNSTABLE symbols", int(regime.eq("UNSTABLE").sum())],
                ["REGIME", "NOT_READY symbols", int(regime.eq("NOT_READY").sum())],
                ["CONFIGURATION", "Autocorrelation window", self.autocorrelation_window],
                ["CONFIGURATION", "Autocorrelation minimum periods", self.autocorrelation_min_periods],
                ["CONFIGURATION", "Momentum threshold", self.momentum_threshold],
                ["CONFIGURATION", "Mean-reversion threshold", self.mean_reversion_threshold],
                ["CONFIGURATION", "Stability candles", self.stability_candles],
                ["CONFIGURATION", "Strength reference", self.strength_reference],
            ],
            columns=["Section", "Metric", "Value"],
        )

    @staticmethod
    def _style_workbook(writer):
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(color="FFFFFF", bold=True)
        green_fill = PatternFill("solid", fgColor="C6EFCE")
        blue_fill = PatternFill("solid", fgColor="DDEBF7")
        red_fill = PatternFill("solid", fgColor="FFC7CE")
        amber_fill = PatternFill("solid", fgColor="FFEB9C")

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Do not materialize every cell in the large history worksheet
            # merely to estimate column widths.
            if worksheet.title == "Regime_History":
                from openpyxl.utils import get_column_letter

                for index in range(1, worksheet.max_column + 1):
                    worksheet.column_dimensions[
                        get_column_letter(index)
                    ].width = 21
                continue

            for cells in worksheet.columns:
                values = [
                    "" if cell.value is None else str(cell.value)
                    for cell in cells[:200]
                ]
                width = min(
                    max(max(map(len, values), default=0) + 2, 11),
                    34,
                )
                worksheet.column_dimensions[cells[0].column_letter].width = width

            headers = {
                cell.value: cell.column_letter
                for cell in worksheet[1]
                if cell.value is not None
            }
            datetime_column = headers.get("Latest_DateTime")
            if datetime_column:
                for row in range(2, worksheet.max_row + 1):
                    worksheet[f"{datetime_column}{row}"].number_format = (
                        "yyyy-mm-dd hh:mm:ss"
                    )

            # Header-only sheets must not receive invalid ranges like G2:G1.
            if worksheet.max_row < 2:
                continue

            regime_column = headers.get("Latest_Stable_Regime")
            if regime_column:
                target = (
                    f"{regime_column}2:"
                    f"{regime_column}{worksheet.max_row}"
                )
                for formula, fill in [
                    (f'{regime_column}2="MOMENTUM"', green_fill),
                    (f'{regime_column}2="MEAN_REVERSION"', blue_fill),
                    (f'{regime_column}2="UNSTABLE"', amber_fill),
                    (f'{regime_column}2="RANDOM"', red_fill),
                ]:
                    worksheet.conditional_formatting.add(
                        target,
                        FormulaRule(formula=[formula], fill=fill),
                    )

    def _write_regime_history(self, writer, regime_data):
        """Append complete regime results one symbol at a time."""
        maximum_data_rows = 1_048_575
        total_rows = sum(
            len(frame)
            for frame in regime_data.values()
            if isinstance(frame, pd.DataFrame) and not frame.empty
        )
        if total_rows > maximum_data_rows:
            raise ValueError(
                "Regime_History exceeds Excel's worksheet limit: "
                f"{total_rows:,} rows > {maximum_data_rows:,} rows"
            )

        if total_rows == 0:
            pd.DataFrame(columns=self.REQUIRED_COLUMNS).to_excel(
                writer,
                sheet_name="Regime_History",
                index=False,
            )
            return 0

        start_row = 0
        write_header = True
        written_rows = 0
        expected_columns = None

        for symbol, frame in regime_data.items():
            if not isinstance(frame, pd.DataFrame) or frame.empty:
                continue

            if expected_columns is None:
                expected_columns = list(frame.columns)
            elif list(frame.columns) != expected_columns:
                frame = frame.reindex(columns=expected_columns)

            chunk = self._excel_safe(frame)
            chunk.to_excel(
                writer,
                sheet_name="Regime_History",
                index=False,
                header=write_header,
                startrow=start_row,
            )
            count = len(chunk)
            written_rows += count
            start_row += count + (1 if write_header else 0)
            write_header = False
            del chunk

        return written_rows

    def _save_excel_report(self, report_df, output_file, regime_data=None):
        output_path = Path(output_file).expanduser().resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        tradable = report_df.loc[
            report_df["Latest_Trade_Eligible"].fillna(False)
            & report_df["Latest_Stable_Regime"].isin(
                ["MOMENTUM", "MEAN_REVERSION"]
            )
        ].copy()
        unstable_random = report_df.loc[
            report_df["Latest_Stable_Regime"].isin(
                ["UNSTABLE", "RANDOM"]
            )
        ].copy()
        not_ready_failed = report_df.loc[
            report_df["Status"].isin(["NOT_READY", "FAILED"])
            | report_df["Latest_Stable_Regime"].eq("NOT_READY")
        ].copy()
        summary = self._build_summary(report_df)

        descriptor, temporary_name = tempfile.mkstemp(
            prefix="RegimeEngine_",
            suffix=".xlsx",
            dir=output_path.parent,
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(
                temporary_name,
                engine="openpyxl",
                datetime_format="yyyy-mm-dd hh:mm:ss",
            ) as writer:
                self._excel_safe(summary).to_excel(
                    writer, sheet_name="Summary", index=False
                )
                self._excel_safe(tradable).to_excel(
                    writer, sheet_name="Tradable_Regimes", index=False
                )
                self._excel_safe(unstable_random).to_excel(
                    writer, sheet_name="Unstable_Random", index=False
                )
                self._excel_safe(not_ready_failed).to_excel(
                    writer, sheet_name="Not_Ready_Failed", index=False
                )
                self._excel_safe(report_df).to_excel(
                    writer, sheet_name="All_Shares", index=False
                )
                history_rows = self._write_regime_history(
                    writer,
                    regime_data or {},
                )
                self._style_workbook(writer)

            # Replaces an existing workbook only after the new file is valid.
            os.replace(temporary_name, output_path)
        except Exception:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)
            raise

        report_df.attrs["Excel_File"] = str(output_path)
        report_df.attrs["Regime_History_Rows"] = history_rows
        return str(output_path)

    # =====================================================
    # Calculate from ResidualEngine.xlsx
    # =====================================================

    def calculate_all_from_excel(
        self,
        residual_file,
        sheet_name="Residual_History",
        output_file=None,
    ):
        """Stream residual history and calculate one symbol at a time."""
        from openpyxl import load_workbook

        input_path = Path(residual_file).expanduser().resolve()
        if not input_path.is_file():
            raise FileNotFoundError(
                f"Residual workbook not found: {input_path}"
            )

        workbook = load_workbook(
            filename=input_path,
            read_only=True,
            data_only=True,
        )
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise ValueError(
                f"Worksheet {sheet_name!r} was not found in "
                f"{input_path}. Available sheets: {workbook.sheetnames}"
            )

        regime_data = {}
        reports = []
        processed_symbols = set()

        def failed_report(symbol, row_count, error):
            return {
                "Symbol": symbol,
                "Status": "FAILED",
                "Input_Rows": int(row_count),
                "Regime_Ready_Rows": 0,
                "Latest_DateTime": pd.NaT,
                "Latest_Autocorrelation": np.nan,
                "Latest_TStatistic": np.nan,
                "Latest_Raw_Regime": "NOT_READY",
                "Latest_Stable_Regime": "NOT_READY",
                "Latest_Regime_Stable": False,
                "Latest_Regime_Duration": 0,
                "Latest_Regime_Strength": np.nan,
                "Latest_Trade_Eligible": False,
                "Regime_Backend": self.regime_backend,
                "Error": str(error),
            }

        def process_symbol(symbol, rows, columns):
            if not symbol or not rows:
                return
            try:
                residual_df = pd.DataFrame.from_records(
                    rows,
                    columns=columns,
                )
                result_df, report = self.calculate_symbol(
                    symbol=symbol,
                    residual_df=residual_df,
                )
                reports.append(report)
                if report["Status"] != "NOT_READY":
                    regime_data[symbol] = result_df
                else:
                    del result_df
                del residual_df
            except Exception as error:
                reports.append(
                    failed_report(symbol, len(rows), error)
                )

        try:
            worksheet = workbook[sheet_name]
            row_iterator = worksheet.iter_rows(values_only=True)
            try:
                raw_header = next(row_iterator)
            except StopIteration as error:
                raise ValueError(
                    f"Worksheet {sheet_name!r} is empty"
                ) from error

            columns = [
                str(value).strip() if value is not None else ""
                for value in raw_header
            ]
            if not columns or "Symbol" not in columns:
                raise ValueError(
                    f"Worksheet {sheet_name!r} has no Symbol column"
                )

            missing_columns = [
                column
                for column in self.REQUIRED_COLUMNS
                if column not in columns
            ]
            if missing_columns:
                raise ValueError(
                    f"Worksheet {sheet_name!r} is missing columns: "
                    f"{missing_columns}"
                )

            symbol_index = columns.index("Symbol")
            current_symbol = None
            current_rows = []

            for values in row_iterator:
                if not values or all(value is None for value in values):
                    continue

                raw_symbol = values[symbol_index]
                symbol = str(raw_symbol).strip().upper()
                if not symbol or symbol == "NONE":
                    continue

                if current_symbol is None:
                    current_symbol = symbol

                if symbol != current_symbol:
                    process_symbol(
                        current_symbol,
                        current_rows,
                        columns,
                    )
                    processed_symbols.add(current_symbol)
                    current_rows.clear()
                    gc.collect()

                    if symbol in processed_symbols:
                        raise ValueError(
                            "Residual_History must be grouped by Symbol; "
                            f"{symbol} appears in multiple blocks"
                        )
                    current_symbol = symbol

                current_rows.append(tuple(values))

            process_symbol(current_symbol, current_rows, columns)
            current_rows.clear()
        finally:
            workbook.close()

        report_df = pd.DataFrame(reports)
        if report_df.empty:
            raise RuntimeError(
                "Residual_History contained no usable symbol rows"
            )

        report_df.sort_values(
            by=["Status", "Symbol"],
            inplace=True,
        )
        report_df.reset_index(drop=True, inplace=True)

        selected_output_file = (
            output_file
            if output_file is not None
            else self.report_output_file
        )
        saved_file = None
        if self.save_excel_enabled and selected_output_file:
            saved_file = self._save_excel_report(
                report_df=report_df,
                output_file=selected_output_file,
                regime_data=regime_data,
            )
            print("RegimeEngine.xlsx file is created")

        report_df.attrs["Excel_File"] = saved_file
        report_df.attrs["Regime_Backend"] = self.regime_backend
        report_df.attrs["Residual_Source_File"] = str(input_path)
        report_df.attrs["Residual_Source_Sheet"] = sheet_name
        report_df.attrs["Results_Retained_In_Memory"] = (
            self.retain_results_in_memory
        )

        if not regime_data:
            raise RuntimeError(
                "Regimes could not be calculated for any symbol. "
                f"Diagnostic workbook: {saved_file}"
            )

        if not self.retain_results_in_memory:
            regime_data.clear()
            gc.collect()

        return regime_data, report_df

    # =====================================================
    # Calculate All Symbols
    # =====================================================

    def calculate_all(self, residual_data, output_file=None):

        if not isinstance(residual_data, dict):
            raise TypeError(
                "residual_data must be a dictionary "
                "containing {symbol: DataFrame}"
            )

        if not residual_data:
            raise ValueError(
                "residual_data is empty"
            )

        regime_data = {}
        reports = []

        for symbol, residual_df in residual_data.items():

            symbol = (
                str(symbol)
                .strip()
                .upper()
            )

            try:

                result_df, report = self.calculate_symbol(
                    symbol=symbol,
                    residual_df=residual_df
                )

                reports.append(report)

                if report["Status"] == "NOT_READY":
                    continue

                regime_data[symbol] = result_df

            except Exception as error:

                reports.append({
                    "Symbol": symbol,
                    "Status": "FAILED",
                    "Input_Rows": (
                        len(residual_df)
                        if isinstance(
                            residual_df,
                            pd.DataFrame
                        )
                        else 0
                    ),
                    "Regime_Ready_Rows": 0,
                    "Latest_DateTime": pd.NaT,
                    "Latest_Autocorrelation": np.nan,
                    "Latest_TStatistic": np.nan,
                    "Latest_Raw_Regime": "NOT_READY",
                    "Latest_Stable_Regime": "NOT_READY",
                    "Latest_Regime_Stable": False,
                    "Latest_Regime_Duration": 0,
                    "Latest_Regime_Strength": np.nan,
                    "Latest_Trade_Eligible": False,
                    "Regime_Backend": self.regime_backend,
                    "Error": str(error)
                })

        report_df = pd.DataFrame(reports)

        if not report_df.empty:

            report_df.sort_values(
                by=["Status", "Symbol"],
                inplace=True
            )

            report_df.reset_index(
                drop=True,
                inplace=True
            )

        selected_output_file = (
            output_file
            if output_file is not None
            else self.report_output_file
        )
        saved_file = None
        if self.save_excel_enabled and selected_output_file:
            saved_file = self._save_excel_report(
                report_df=report_df,
                output_file=selected_output_file,
                regime_data=regime_data,
            )
            print("RegimeEngine.xlsx file is created")

        report_df.attrs["Excel_File"] = saved_file
        report_df.attrs["Regime_Backend"] = self.regime_backend

        if not regime_data:
            raise RuntimeError(
                "Regimes could not be calculated for any symbol. "
                f"Diagnostic workbook: {saved_file}"
            )

        report_df.attrs["Results_Retained_In_Memory"] = (
            self.retain_results_in_memory
        )
        if not self.retain_results_in_memory:
            regime_data.clear()
            gc.collect()

        return regime_data, report_df

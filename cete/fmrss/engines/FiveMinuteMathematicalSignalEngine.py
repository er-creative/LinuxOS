import os
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd


class FiveMinuteMathematicalSignalEngine:
    """Calculate signals only for shares selected in RankingShares.xlsx."""

    RANKING_COLUMNS = [
        "Symbol", "Decision_DateTime", "Latest_DateTime", "Latest_Regime",
        "Mathematical_Setup_Accepted", "Residual_ZScore",
        "Volume_Factor", "Latest_Mathematical_Score",
        "Latest_Cross_Sectional_Score", "Latest_Rank",
        "Latest_Eligible", "Latest_Selected",
        "Latest_Selection_Reason", "Score_Rejection_Reason",
    ]
    MATHEMATICAL_COLUMNS = [
        "Symbol", "Latest_DateTime", "Latest_Stable_Regime",
        "Latest_Residual_Impulse", "Latest_Residual_ZScore",
        "Latest_Residual_Return", "Latest_Close", "Latest_Previous_Close",
        "Latest_Residual_Fair_Value", "Latest_Expected_Move_Percent",
        "Latest_Edge_To_Cost_Ratio", "Latest_Residual_Half_Life",
        "Latest_Residual_Quality_Score",
        "Latest_Volume_Factor", "Latest_Mathematical_Score",
        "Latest_Setup_Ready", "Latest_Setup_Accepted",
        "Latest_Volatility_Eligible", "Latest_Volume_Confirmed",
        "Latest_Rejection_Reason",
    ]
    OUTPUT_COLUMNS = [
        "Symbol", "Decision_DateTime", "Ranking_Source_DateTime",
        "Mathematical_DateTime", "Rank", "Regime",
        "Residual_ZScore", "Previous_Residual_ZScore",
        "Signal_Armed", "Armed_Direction", "Armed_DateTime",
        "Armed_Age_Candles", "Maximum_Armed_Abs_ZScore",
        "ZScore_Turn_Confirmed", "Residual_Return_Confirmed",
        "Reversion_Confirmation_Count",
        "Reversion_Turn_Confirmed", "Price_Turn_Confirmed",
        "Residual_Impulse", "Residual_Return", "Volume_Factor",
        "Latest_Close", "Residual_Fair_Value", "Expected_Move_Percent",
        "Edge_To_Cost_Ratio", "Residual_Half_Life",
        "Residual_Quality_Score",
        "Mathematical_Score", "Cross_Sectional_Score",
        "Timestamp_Matched", "Mathematical_Validated",
        "Ranking_Fallback_Used", "Signal_Eligible",
        "Mathematical_Signal", "Signal_Direction", "Signal_Strength",
        "Signal_Reason", "Calculation_Backend",
    ]
    STATE_COLUMNS = [
        "Symbol", "Last_DateTime", "Previous_ZScore",
        "Maximum_Abs_ZScore", "Armed_Direction", "Armed_DateTime",
        "Armed_Cycles", "Last_Rank", "Last_Regime",
        "Last_Selected_DateTime", "Last_Signal_DateTime",
    ]
    DIRECTIONAL_REGIMES = {"MOMENTUM", "MEAN_REVERSION"}

    def __init__(
        self,
        ranking_file="/home/devinderjeet/fmrss/report/RankingShares.xlsx",
        mathematical_score_file=(
            "/home/devinderjeet/fmrss/report/MathematicalScoreEngine.xlsx"
        ),
        output_file=(
            "/home/devinderjeet/fmrss/report/MathematicalSignals.xlsx"
        ),
        signal_state_file=None,
        maximum_selected_symbols=5,
        momentum_minimum_abs_zscore=1.00,
        mean_reversion_minimum_abs_zscore=2.50,
        minimum_abs_impulse=0.0,
        allow_momentum_entries=False,
        allow_mean_reversion_entries=True,
        require_mean_reversion_turn=True,
        require_previous_zscore_turn=True,
        mean_reversion_arm_abs_zscore=None,
        mean_reversion_entry_abs_zscore=2.20,
        mean_reversion_minimum_remaining_abs_zscore=1.50,
        minimum_reversion_confirmations=2,
        maximum_armed_age_candles=12,
        mathematical_strength_weight=0.70,
        cross_sectional_strength_weight=0.30,
        score_match_tolerance=0.05,
        allow_ranking_snapshot_fallback=True,
        memory_optimized=True,
        save_excel=True,
        print_signal_details=False,
        weight_tolerance=1e-9,
    ):
        if maximum_selected_symbols < 1:
            raise ValueError("maximum_selected_symbols must be at least 1")
        for name, value in {
            "momentum_minimum_abs_zscore": momentum_minimum_abs_zscore,
            "mean_reversion_minimum_abs_zscore": mean_reversion_minimum_abs_zscore,
        }.items():
            if not np.isfinite(value) or value <= 0:
                raise ValueError(f"{name} must be finite and positive")
        if not np.isfinite(minimum_abs_impulse) or minimum_abs_impulse < 0:
            raise ValueError("minimum_abs_impulse must be finite and non-negative")
        if not np.isfinite(score_match_tolerance) or score_match_tolerance < 0:
            raise ValueError("score_match_tolerance must be finite and non-negative")
        arm_threshold = (
            mean_reversion_minimum_abs_zscore
            if mean_reversion_arm_abs_zscore is None
            else mean_reversion_arm_abs_zscore
        )
        for name, value in {
            "mean_reversion_arm_abs_zscore": arm_threshold,
            "mean_reversion_entry_abs_zscore": mean_reversion_entry_abs_zscore,
            "mean_reversion_minimum_remaining_abs_zscore": (
                mean_reversion_minimum_remaining_abs_zscore
            ),
        }.items():
            if not np.isfinite(value) or value <= 0:
                raise ValueError(f"{name} must be finite and positive")
        if not (
            mean_reversion_minimum_remaining_abs_zscore
            <= mean_reversion_entry_abs_zscore
            < arm_threshold
        ):
            raise ValueError(
                "Mean-reversion thresholds must satisfy minimum remaining "
                "<= entry < arm"
            )
        if minimum_reversion_confirmations not in {1, 2, 3}:
            raise ValueError("minimum_reversion_confirmations must be 1, 2 or 3")
        if maximum_armed_age_candles < 1:
            raise ValueError("maximum_armed_age_candles must be positive")

        weights = np.asarray(
            [mathematical_strength_weight, cross_sectional_strength_weight],
            dtype=np.float64,
        )
        if not np.isfinite(weights).all() or (weights < 0).any():
            raise ValueError("Signal weights must be finite and non-negative")
        if not np.isclose(weights.sum(), 1.0, rtol=0.0, atol=weight_tolerance):
            raise ValueError("Signal weights must sum to 1.0")

        self.ranking_file = str(ranking_file)
        self.mathematical_score_file = str(mathematical_score_file)
        self.output_file = str(output_file) if output_file else None
        self.signal_state_file = (
            str(signal_state_file) if signal_state_file else None
        )
        self.maximum_selected_symbols = int(maximum_selected_symbols)
        self.momentum_minimum_abs_zscore = float(momentum_minimum_abs_zscore)
        self.mean_reversion_minimum_abs_zscore = float(
            mean_reversion_minimum_abs_zscore
        )
        self.minimum_abs_impulse = float(minimum_abs_impulse)
        self.allow_momentum_entries = bool(allow_momentum_entries)
        self.allow_mean_reversion_entries = bool(
            allow_mean_reversion_entries
        )
        self.require_mean_reversion_turn = bool(
            require_mean_reversion_turn
        )
        self.require_previous_zscore_turn = bool(
            require_previous_zscore_turn
        )
        self.mean_reversion_arm_abs_zscore = float(arm_threshold)
        self.mean_reversion_entry_abs_zscore = float(
            mean_reversion_entry_abs_zscore
        )
        self.mean_reversion_minimum_remaining_abs_zscore = float(
            mean_reversion_minimum_remaining_abs_zscore
        )
        self.minimum_reversion_confirmations = int(
            minimum_reversion_confirmations
        )
        self.maximum_armed_age_candles = int(maximum_armed_age_candles)
        self.mathematical_strength_weight = float(mathematical_strength_weight)
        self.cross_sectional_strength_weight = float(
            cross_sectional_strength_weight
        )
        self.score_match_tolerance = float(score_match_tolerance)
        self.allow_ranking_snapshot_fallback = bool(
            allow_ranking_snapshot_fallback
        )
        self.memory_optimized = bool(memory_optimized)
        self.save_excel_enabled = bool(save_excel)
        self.print_signal_details = bool(print_signal_details)
        self.calculation_backend = (
            "NUMPY_PANDAS_SELECTED_ONLY_MEMORY_OPTIMIZED"
            if self.memory_optimized else "NUMPY_PANDAS_SELECTED_ONLY"
        )

    @staticmethod
    def _path(file_path, description):
        path = Path(file_path).expanduser().resolve()
        if not path.exists():
            raise FileNotFoundError(f"{description} not found:\n{path}")
        return path

    @staticmethod
    def _check_columns(frame, required, description):
        missing = [column for column in required if column not in frame]
        if missing:
            raise ValueError(f"{description} missing columns: {missing}")

    @staticmethod
    def _symbols(series):
        return series.astype("string").str.strip().str.upper()

    @staticmethod
    def _bool(series):
        if pd.api.types.is_bool_dtype(series.dtype):
            return series.fillna(False).astype(bool)
        return series.astype("string").str.strip().str.upper().isin(
            ["TRUE", "1", "YES"]
        )

    @staticmethod
    def _number(series):
        return pd.to_numeric(series, errors="coerce").replace(
            [np.inf, -np.inf], np.nan
        )

    def _load_ranking(self, file_path):
        path = self._path(file_path, "Ranking workbook")
        # Selected_Shares is the sole candidate universe.  Do not fall back to
        # All_Shares because that could accidentally generate signals for
        # symbols outside the top-N selection.
        frame = pd.read_excel(
            path, sheet_name="Selected_Shares", usecols=self.RANKING_COLUMNS
        )
        self._check_columns(frame, self.RANKING_COLUMNS, "Ranking worksheet")
        frame["Symbol"] = self._symbols(frame["Symbol"])
        for column in ["Decision_DateTime", "Latest_DateTime"]:
            frame[column] = pd.to_datetime(frame[column], errors="coerce")
        decision_times = frame["Decision_DateTime"].dropna().unique()
        ranking_decision_timestamp = (
            pd.Timestamp(decision_times[0])
            if len(decision_times) == 1
            else pd.NaT
        )
        for column in [
            "Mathematical_Setup_Accepted", "Latest_Eligible", "Latest_Selected"
        ]:
            frame[column] = self._bool(frame[column])
        frame["Latest_Selection_Reason"] = (
            frame["Latest_Selection_Reason"]
            .astype("string").str.strip().str.upper()
        )
        frame = frame.loc[
            frame["Latest_Selected"]
            & frame["Latest_Selection_Reason"].eq("SELECTED")
        ].copy()
        # Zero selected shares is a valid no-trade decision.  It commonly
        # occurs at the beginning of a session while intraday residual and
        # impulse features are warming up.  Preserve the expected schema and
        # let calculate_all create an empty signal report instead of stopping
        # a continuous forward test.
        if frame.empty:
            frame.reset_index(drop=True, inplace=True)
            frame.attrs["Decision_Timestamp"] = ranking_decision_timestamp
            return frame
        if len(frame) > self.maximum_selected_symbols:
            raise RuntimeError(
                f"Selected shares {len(frame)} exceed maximum "
                f"{self.maximum_selected_symbols}"
            )
        if frame["Symbol"].duplicated().any():
            raise RuntimeError("Ranking workbook contains duplicate selected symbols")
        times = frame["Decision_DateTime"].dropna().unique()
        if len(times) != 1 or frame["Decision_DateTime"].isna().any():
            raise RuntimeError("Selected shares do not share one decision timestamp")
        if frame["Latest_DateTime"].isna().any():
            raise RuntimeError("Selected shares contain an invalid source timestamp")
        frame.attrs["Decision_Timestamp"] = ranking_decision_timestamp
        return frame

    def _load_mathematical(self, file_path, symbols):
        path = self._path(file_path, "Mathematical score workbook")
        try:
            frame = pd.read_excel(
                path,
                sheet_name="Accepted_Setups",
                usecols=self.MATHEMATICAL_COLUMNS,
            )
        except ValueError:
            # Compatibility with older MathematicalScoreEngine workbooks.
            frame = pd.read_excel(
                path, sheet_name="All_Shares", usecols=self.MATHEMATICAL_COLUMNS
            )
        self._check_columns(
            frame, self.MATHEMATICAL_COLUMNS, "Mathematical worksheet"
        )
        frame["Symbol"] = self._symbols(frame["Symbol"])
        frame = frame.loc[frame["Symbol"].isin(symbols)].copy()
        if frame["Symbol"].duplicated().any():
            raise RuntimeError("Mathematical workbook contains duplicate symbols")
        frame["Latest_DateTime"] = pd.to_datetime(
            frame["Latest_DateTime"], errors="coerce"
        )
        for column in [
            "Latest_Setup_Ready", "Latest_Setup_Accepted",
            "Latest_Volatility_Eligible", "Latest_Volume_Confirmed",
        ]:
            frame[column] = self._bool(frame[column])
        return frame

    def _combine(self, ranking, mathematical):
        ranking = ranking.rename(columns={
            "Latest_DateTime": "Ranking_Source_DateTime",
            "Latest_Regime": "Regime",
            "Latest_Mathematical_Score": "Mathematical_Score",
            "Latest_Cross_Sectional_Score": "Cross_Sectional_Score",
            "Latest_Rank": "Rank",
        })
        mathematical = mathematical.rename(columns={
            "Latest_DateTime": "Mathematical_DateTime",
            "Latest_Stable_Regime": "Mathematical_Regime",
            "Latest_Residual_Impulse": "Residual_Impulse",
            "Latest_Residual_Return": "Residual_Return",
            "Latest_Residual_Fair_Value": "Residual_Fair_Value",
            "Latest_Expected_Move_Percent": "Expected_Move_Percent",
            "Latest_Edge_To_Cost_Ratio": "Edge_To_Cost_Ratio",
            "Latest_Residual_Half_Life": "Residual_Half_Life",
            "Latest_Residual_Quality_Score": "Residual_Quality_Score",
            "Latest_Residual_ZScore": "Mathematical_Residual_ZScore",
            "Latest_Volume_Factor": "Mathematical_Volume_Factor",
            "Latest_Mathematical_Score": "Mathematical_Report_Score",
        })
        result = ranking.merge(
            mathematical, on="Symbol", how="left", validate="one_to_one", copy=False
        )
        for column in ["Regime", "Mathematical_Regime"]:
            result[column] = (
                result[column].astype("string").str.strip().str.upper()
            )
        for column in [
            "Residual_ZScore", "Residual_Impulse", "Volume_Factor",
            "Mathematical_Score", "Cross_Sectional_Score", "Rank",
            "Mathematical_Residual_ZScore", "Mathematical_Volume_Factor",
            "Mathematical_Report_Score",
            "Residual_Return", "Latest_Close", "Latest_Previous_Close",
            "Residual_Fair_Value", "Expected_Move_Percent",
            "Edge_To_Cost_Ratio", "Residual_Half_Life",
            "Residual_Quality_Score",
        ]:
            result[column] = self._number(result[column])

        result["Timestamp_Matched"] = result["Ranking_Source_DateTime"].eq(
            result["Mathematical_DateTime"]
        )
        # Never expose or consume mathematical-only values from a later or
        # earlier candle. Mean-reversion fallback uses only synchronized
        # fields already embedded in RankingShares.xlsx.
        result["Residual_Impulse"] = result["Residual_Impulse"].where(
            result["Timestamp_Matched"], np.nan
        )
        for column in [
            "Residual_Return", "Latest_Close", "Latest_Previous_Close",
            "Residual_Fair_Value", "Expected_Move_Percent",
            "Edge_To_Cost_Ratio", "Residual_Half_Life",
            "Residual_Quality_Score",
        ]:
            result[column] = result[column].where(
                result["Timestamp_Matched"], np.nan
            )
        regime_match = result["Regime"].eq(result["Mathematical_Regime"])
        score_match = result["Mathematical_Score"].sub(
            result["Mathematical_Report_Score"]
        ).abs().le(self.score_match_tolerance)
        result["Mathematical_Validated"] = (
            result["Timestamp_Matched"] & regime_match & score_match
            & result["Latest_Setup_Ready"].fillna(False)
            & result["Latest_Setup_Accepted"].fillna(False)
            & result["Latest_Volatility_Eligible"].fillna(False)
            & result["Latest_Volume_Confirmed"].fillna(False)
            & result["Latest_Rejection_Reason"].eq("ACCEPTED")
        )
        result["Ranking_Validated"] = (
            result["Mathematical_Setup_Accepted"]
            & result["Latest_Eligible"] & result["Latest_Selected"]
            & result["Latest_Selection_Reason"].eq("SELECTED")
            & result["Score_Rejection_Reason"].eq("ACCEPTED")
            & result["Regime"].isin(self.DIRECTIONAL_REGIMES)
        )
        result["Ranking_Fallback_Used"] = (
            self.allow_ranking_snapshot_fallback
            & result["Ranking_Validated"] & ~result["Timestamp_Matched"]
        )
        return result

    def _empty_state(self):
        return pd.DataFrame(columns=self.STATE_COLUMNS)

    def _load_signal_state(self):
        if not self.signal_state_file \
                or not os.path.isfile(self.signal_state_file):
            return self._empty_state()
        try:
            state = pd.read_excel(
                self.signal_state_file,
                sheet_name="Signal_State",
                engine="openpyxl",
            ).reindex(columns=self.STATE_COLUMNS)
        except (ValueError, OSError):
            return self._empty_state()
        if state.empty:
            return state
        state["Symbol"] = self._symbols(state["Symbol"])
        for column in [
            "Last_DateTime", "Armed_DateTime", "Last_Selected_DateTime",
            "Last_Signal_DateTime",
        ]:
            state[column] = pd.to_datetime(state[column], errors="coerce")
        for column in [
            "Previous_ZScore", "Maximum_Abs_ZScore", "Last_Rank"
        ]:
            state[column] = self._number(state[column])
        state["Armed_Cycles"] = pd.to_numeric(
            state["Armed_Cycles"], errors="coerce"
        ).fillna(0).astype(np.int16)
        for column in ["Armed_Direction", "Last_Regime"]:
            state[column] = (
                state[column].astype("string").fillna("").str.upper()
            )
        state.dropna(subset=["Symbol"], inplace=True)
        state.drop_duplicates("Symbol", keep="last", inplace=True)
        state.reset_index(drop=True, inplace=True)
        return state

    def _save_signal_state(self, state):
        if not self.signal_state_file:
            return None
        target = Path(self.signal_state_file).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        metadata = pd.DataFrame({
            "Metric": [
                "State Version", "Updated UTC", "Tracked Symbols",
                "Armed Symbols", "Arm Z-Score", "Entry Z-Score",
                "Minimum Remaining Z-Score", "Confirmations Required",
                "Maximum Armed Age Candles",
            ],
            "Value": [
                1, pd.Timestamp.now(tz="UTC").isoformat(), len(state),
                int(state["Armed_Direction"].isin(["BUY", "SELL"]).sum()),
                self.mean_reversion_arm_abs_zscore,
                self.mean_reversion_entry_abs_zscore,
                self.mean_reversion_minimum_remaining_abs_zscore,
                self.minimum_reversion_confirmations,
                self.maximum_armed_age_candles,
            ],
        })
        descriptor, temporary = tempfile.mkstemp(
            prefix=".fmrss_signal_state_", suffix=".xlsx", dir=target.parent
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                self._excel_safe(metadata).to_excel(
                    writer, sheet_name="Metadata", index=False
                )
                self._excel_safe(state).to_excel(
                    writer, sheet_name="Signal_State", index=False
                )
                self._style(writer)
            os.replace(temporary, target)
        except Exception:
            if os.path.exists(temporary):
                os.unlink(temporary)
            raise
        return str(target)

    def _signals(self, result, state):
        validation = result["Ranking_Validated"] & (
            result["Mathematical_Validated"] | result["Ranking_Fallback_Used"]
        )
        state_by_symbol = {
            str(row.Symbol): row._asdict()
            for row in state.itertuples(index=False)
        }
        diagnostics = []

        for index, row in result.iterrows():
            symbol = str(row["Symbol"])
            now = pd.to_datetime(row["Decision_DateTime"], errors="coerce")
            zscore = float(row["Residual_ZScore"])
            regime = str(row["Regime"])
            valid = bool(validation.loc[index])
            prior = state_by_symbol.get(symbol, {})
            previous_zscore = pd.to_numeric(
                prior.get("Previous_ZScore"), errors="coerce"
            )
            armed_direction = str(
                prior.get("Armed_Direction", "") or ""
            ).upper()
            armed_time = pd.to_datetime(
                prior.get("Armed_DateTime"), errors="coerce"
            )
            armed_cycles_value = pd.to_numeric(
                prior.get("Armed_Cycles", 0), errors="coerce"
            )
            armed_cycles = (
                int(armed_cycles_value)
                if pd.notna(armed_cycles_value)
                else 0
            )
            maximum_abs_zscore = pd.to_numeric(
                prior.get("Maximum_Abs_ZScore"), errors="coerce"
            )
            armed_age = (
                (now - armed_time).total_seconds() / 300.0
                if pd.notna(now) and pd.notna(armed_time)
                else np.nan
            )
            expired = bool(
                np.isfinite(armed_age)
                and armed_age > self.maximum_armed_age_candles
            )
            if expired or regime != "MEAN_REVERSION":
                armed_direction = ""
                armed_time = pd.NaT
                armed_cycles = 0
                maximum_abs_zscore = np.nan

            expected_direction = "BUY" if zscore < 0 else "SELL"
            was_armed = (
                armed_direction == expected_direction
                and pd.notna(armed_time)
                and np.isfinite(maximum_abs_zscore)
                and maximum_abs_zscore >= self.mean_reversion_arm_abs_zscore
            )
            zscore_turn = bool(
                np.isfinite(previous_zscore)
                and (
                    (zscore < 0 and zscore > previous_zscore)
                    or (zscore > 0 and zscore < previous_zscore)
                )
            )
            residual_return = pd.to_numeric(
                row["Residual_Return"], errors="coerce"
            )
            return_turn = bool(
                np.isfinite(residual_return)
                and (
                    (zscore < 0 and residual_return > 0)
                    or (zscore > 0 and residual_return < 0)
                )
            )
            close = pd.to_numeric(row["Latest_Close"], errors="coerce")
            previous_close = pd.to_numeric(
                row["Latest_Previous_Close"], errors="coerce"
            )
            price_turn = bool(
                np.isfinite(close) and np.isfinite(previous_close)
                and (
                    (zscore < 0 and close > previous_close)
                    or (zscore > 0 and close < previous_close)
                )
            )
            confirmation_count = int(
                zscore_turn + return_turn + price_turn
            )
            crossback = bool(
                was_armed
                and np.isfinite(previous_zscore)
                and self.mean_reversion_minimum_remaining_abs_zscore
                <= abs(zscore)
                <= self.mean_reversion_entry_abs_zscore
                and zscore_turn
            )
            reversion_signal = bool(
                self.allow_mean_reversion_entries
                and valid
                and regime == "MEAN_REVERSION"
                and crossback
                and (
                    not self.require_mean_reversion_turn
                    or confirmation_count >= self.minimum_reversion_confirmations
                )
            )
            impulse = pd.to_numeric(row["Residual_Impulse"], errors="coerce")
            momentum_signal = bool(
                self.allow_momentum_entries and valid
                and bool(row["Timestamp_Matched"])
                and regime == "MOMENTUM"
                and abs(zscore) >= self.momentum_minimum_abs_zscore
                and np.isfinite(impulse)
                and abs(impulse) > self.minimum_abs_impulse
                and np.sign(zscore) == np.sign(impulse)
            )
            signal = "NO_TRADE"
            if reversion_signal:
                signal = expected_direction
            elif momentum_signal:
                signal = "BUY" if impulse > 0 else "SELL"

            newly_armed = False
            if signal in {"BUY", "SELL"}:
                armed_direction = ""
                armed_time = pd.NaT
                armed_cycles = 0
                maximum_abs_zscore = np.nan
            elif valid and regime == "MEAN_REVERSION" \
                    and abs(zscore) >= self.mean_reversion_arm_abs_zscore:
                if armed_direction != expected_direction or pd.isna(armed_time):
                    armed_direction = expected_direction
                    armed_time = now
                    armed_cycles = 1
                    maximum_abs_zscore = abs(zscore)
                    newly_armed = True
                else:
                    armed_cycles += 1
                    maximum_abs_zscore = max(
                        float(maximum_abs_zscore), abs(zscore)
                    )
            elif abs(zscore) < self.mean_reversion_minimum_remaining_abs_zscore:
                armed_direction = ""
                armed_time = pd.NaT
                armed_cycles = 0
                maximum_abs_zscore = np.nan

            if signal == "BUY":
                reason = "BUY_CROSSBACK_CONFIRMED"
            elif signal == "SELL":
                reason = "SELL_CROSSBACK_CONFIRMED"
            elif not bool(row["Ranking_Validated"]):
                reason = "RANKING_VALIDATION_FAILED"
            elif not bool(row["Timestamp_Matched"]) \
                    and not bool(row["Ranking_Fallback_Used"]):
                reason = "TIMESTAMP_MISMATCH"
            elif bool(row["Timestamp_Matched"]) \
                    and not bool(row["Mathematical_Validated"]):
                reason = "MATHEMATICAL_VALIDATION_FAILED"
            elif regime == "MOMENTUM" and not self.allow_momentum_entries:
                reason = "MOMENTUM_DISABLED"
            elif regime == "MEAN_REVERSION" \
                    and not self.allow_mean_reversion_entries:
                reason = "MEAN_REVERSION_DISABLED"
            elif expired:
                reason = "ARMED_SIGNAL_EXPIRED"
            elif newly_armed:
                reason = "RESIDUAL_EXTREME_ARMED"
            elif not was_armed:
                reason = "WAITING_FOR_RESIDUAL_EXTREME"
            elif not crossback:
                reason = "WAITING_FOR_INWARD_CROSSBACK"
            elif confirmation_count < self.minimum_reversion_confirmations:
                reason = "INSUFFICIENT_REVERSAL_CONFIRMATIONS"
            else:
                reason = "NO_VALID_SIGNAL"

            state_by_symbol[symbol] = {
                "Symbol": symbol,
                "Last_DateTime": now,
                "Previous_ZScore": zscore,
                "Maximum_Abs_ZScore": maximum_abs_zscore,
                "Armed_Direction": armed_direction,
                "Armed_DateTime": armed_time,
                "Armed_Cycles": armed_cycles,
                "Last_Rank": row["Rank"],
                "Last_Regime": regime,
                "Last_Selected_DateTime": now,
                "Last_Signal_DateTime": (
                    now if signal in {"BUY", "SELL"}
                    else prior.get("Last_Signal_DateTime", pd.NaT)
                ),
            }
            diagnostics.append({
                "Previous_Residual_ZScore": previous_zscore,
                "Signal_Armed": armed_direction in {"BUY", "SELL"},
                "Armed_Direction": armed_direction,
                "Armed_DateTime": armed_time,
                "Armed_Age_Candles": armed_age,
                "Maximum_Armed_Abs_ZScore": maximum_abs_zscore,
                "ZScore_Turn_Confirmed": zscore_turn,
                "Residual_Return_Confirmed": return_turn,
                "Reversion_Confirmation_Count": confirmation_count,
                "Reversion_Turn_Confirmed": (
                    confirmation_count >= self.minimum_reversion_confirmations
                ),
                "Price_Turn_Confirmed": price_turn,
                "Mathematical_Signal": signal,
                "Signal_Direction": 1 if signal == "BUY" else -1 if signal == "SELL" else 0,
                "Signal_Reason": reason,
            })

        diagnostic_frame = pd.DataFrame(diagnostics, index=result.index)
        for column in diagnostic_frame:
            result[column] = diagnostic_frame[column]
        strength = (
            self.mathematical_strength_weight * result["Mathematical_Score"]
            + self.cross_sectional_strength_weight * result["Cross_Sectional_Score"]
        ).clip(0.0, 100.0)
        dtype = np.float32 if self.memory_optimized else np.float64
        result["Signal_Eligible"] = validation.to_numpy(dtype=bool, copy=False)
        result["Mathematical_Signal"] = pd.Categorical(
            result["Mathematical_Signal"], categories=["NO_TRADE", "BUY", "SELL"]
        )
        result["Signal_Direction"] = result["Signal_Direction"].astype(np.int8)
        result["Signal_Strength"] = strength.where(
            result["Mathematical_Signal"].astype("string") != "NO_TRADE",
            np.nan,
        ).astype(dtype, copy=False)
        result["Signal_Reason"] = pd.Categorical(result["Signal_Reason"])
        result["Calculation_Backend"] = self.calculation_backend
        updated_state = pd.DataFrame(
            state_by_symbol.values(), columns=self.STATE_COLUMNS
        )
        updated_state.sort_values("Symbol", inplace=True)
        updated_state.reset_index(drop=True, inplace=True)
        return result, updated_state

    @staticmethod
    def _excel_safe(frame):
        safe = frame.copy(deep=False)
        for column in safe.columns:
            if isinstance(safe[column].dtype, pd.CategoricalDtype):
                safe[column] = safe[column].astype("string")
            if isinstance(safe[column].dtype, pd.DatetimeTZDtype):
                safe[column] = safe[column].dt.tz_localize(None)
        return safe

    @staticmethod
    def _style(writer):
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        dark = PatternFill("solid", fgColor="17365D")
        green = PatternFill("solid", fgColor="C6EFCE")
        red = PatternFill("solid", fgColor="FFC7CE")
        amber = PatternFill("solid", fgColor="FFEB9C")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.fill = dark
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(horizontal="center")
            for cells in sheet.columns:
                values = ["" if c.value is None else str(c.value) for c in cells[:100]]
                width = min(max(max(map(len, values), default=0) + 2, 11), 32)
                sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
            headers = {c.value: c.column for c in sheet[1]}
            signal_col = headers.get("Mathematical_Signal")
            if signal_col and sheet.max_row >= 2:
                letter = get_column_letter(signal_col)
                target = f"{letter}2:{letter}{sheet.max_row}"
                sheet.conditional_formatting.add(
                    target, FormulaRule(formula=[f'${letter}2="BUY"'], fill=green)
                )
                sheet.conditional_formatting.add(
                    target, FormulaRule(formula=[f'${letter}2="SELL"'], fill=red)
                )
            fallback_col = headers.get("Ranking_Fallback_Used")
            if fallback_col and sheet.max_row >= 2:
                letter = get_column_letter(fallback_col)
                sheet.conditional_formatting.add(
                    f"{letter}2:{letter}{sheet.max_row}",
                    FormulaRule(formula=[f'${letter}2=TRUE'], fill=amber),
                )
            for header in [
                "Decision_DateTime", "Ranking_Source_DateTime",
                "Mathematical_DateTime",
            ]:
                column = headers.get(header)
                if column:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, column).number_format = "yyyy-mm-dd hh:mm:ss"
            for header in [
                "Residual_ZScore", "Residual_Impulse", "Volume_Factor",
                "Mathematical_Score", "Cross_Sectional_Score", "Signal_Strength",
            ]:
                column = headers.get(header)
                if column:
                    for row in range(2, sheet.max_row + 1):
                        sheet.cell(row, column).number_format = "0.0000"

    def _save(self, report, state, output_file):
        output_path = Path(output_file).expanduser().resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        actionable = report.loc[
            report["Mathematical_Signal"].astype("string").isin(["BUY", "SELL"])
        ].copy()
        issues = report.loc[
            ~report["Timestamp_Matched"] | ~report["Mathematical_Validated"]
        ].copy()
        decision_timestamp = (
            report["Decision_DateTime"].iloc[0]
            if not report.empty
            else report.attrs.get("Decision_Timestamp", pd.NaT)
        )
        summary = pd.DataFrame({
            "Metric": [
                "Decision Timestamp", "Selected Shares Processed", "BUY Signals",
                "SELL Signals", "NO_TRADE Signals", "Exact Validations",
                "Ranking Snapshot Fallbacks", "Armed Symbols",
                "Calculation Backend",
            ],
            "Value": [
                decision_timestamp, len(report),
                int(report["Mathematical_Signal"].eq("BUY").sum()),
                int(report["Mathematical_Signal"].eq("SELL").sum()),
                int(report["Mathematical_Signal"].eq("NO_TRADE").sum()),
                int(report["Mathematical_Validated"].sum()),
                int(report["Ranking_Fallback_Used"].sum()),
                int(state["Armed_Direction"].isin(["BUY", "SELL"]).sum()),
                self.calculation_backend,
            ],
        })
        descriptor, temporary = tempfile.mkstemp(
            prefix="fmrss_selected_signals_", suffix=".xlsx", dir=output_path.parent
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                self._excel_safe(summary).to_excel(writer, sheet_name="Summary", index=False)
                self._excel_safe(report).to_excel(
                    writer, sheet_name="Selected_Signals", index=False
                )
                self._excel_safe(actionable).to_excel(
                    writer, sheet_name="Actionable_Signals", index=False
                )
                self._excel_safe(issues).to_excel(
                    writer, sheet_name="Validation_Issues", index=False
                )
                self._excel_safe(state).to_excel(
                    writer, sheet_name="Signal_State_Snapshot", index=False
                )
                self._style(writer)
            os.replace(temporary, output_path)
        except Exception:
            if os.path.exists(temporary):
                os.unlink(temporary)
            raise
        return str(output_path)

    def calculate_all(
        self, ranking_file=None, mathematical_score_file=None, output_file=None
    ):
        ranking = self._load_ranking(ranking_file or self.ranking_file)
        selected_symbols = ranking["Symbol"].tolist()
        state = self._load_signal_state()

        if selected_symbols:
            mathematical = self._load_mathematical(
                mathematical_score_file or self.mathematical_score_file,
                selected_symbols,
            )
            calculated, state = self._signals(
                self._combine(ranking, mathematical), state
            )
            report = calculated.reindex(columns=self.OUTPUT_COLUMNS).copy()
            report.sort_values(["Rank", "Symbol"], inplace=True)
            report.reset_index(drop=True, inplace=True)
            signal_data = {
                symbol: report.loc[
                    report["Symbol"].eq(symbol)
                ].copy(deep=False)
                for symbol in report["Symbol"]
            }
        else:
            # An explicitly typed, schema-complete empty report keeps the
            # caller and Excel export stable without allocating symbol frames.
            report = pd.DataFrame(columns=self.OUTPUT_COLUMNS)
            report.attrs["Decision_Timestamp"] = ranking.attrs.get(
                "Decision_Timestamp", pd.NaT
            )
            signal_data = {}

        selected_output = output_file or self.output_file
        saved = None
        state_saved = self._save_signal_state(state)
        if self.save_excel_enabled and selected_output:
            saved = self._save(report, state, selected_output)
        actionable = report["Mathematical_Signal"].astype("string").isin(
            ["BUY", "SELL"]
        )
        if self.print_signal_details:
            for row in report.itertuples(index=False):
                print(
                    f"{row.Symbol:<15}CALCULATED : Signal={row.Mathematical_Signal} | "
                    f"Regime={row.Regime} | Rank={int(row.Rank)} | "
                    f"Reason={row.Signal_Reason}"
                )
        report.attrs["Decision_Timestamp"] = (
            report["Decision_DateTime"].iloc[0]
            if not report.empty
            else ranking.attrs.get("Decision_Timestamp", pd.NaT)
        )
        report.attrs["Selected_Symbols"] = selected_symbols
        report.attrs["Excel_File"] = saved
        report.attrs["Signal_State_File"] = state_saved
        print("=" * 120)
        print("FMRSS : SELECTED-SHARE MATHEMATICAL SIGNALS")
        print("=" * 120)
        print(f"Decision timestamp : {report.attrs['Decision_Timestamp']}")
        print(f"Selected shares processed : {len(report)}")
        print(f"Actionable signals : {int(actionable.sum())}")
        if saved:
            print(f"Signal workbook : {saved}")
        print("=" * 120)
        return signal_data, report

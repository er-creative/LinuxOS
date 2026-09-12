import io
import os
import tempfile
from collections import deque
from datetime import time as clock_time
from pathlib import Path

import numpy as np
import pandas as pd


class FiveMinuteResidualStatisticalStrategy:
    """Paper-execute and manage FMRSS trades approved by the Risk Engine.

    State is restored from and saved to an Excel workbook. The class never
    sends broker orders. One call represents one strategy cycle.
    """

    PLAN_COLUMNS = [
        "Decision_ID", "Symbol", "Decision_DateTime", "Market_DateTime",
        "Rank", "Regime", "Trade_Action", "Signal_Strength", "Entry_Price",
        "ATR_14", "Stop_Distance", "Stop_Loss", "Target_Price",
        "Reward_Risk_Ratio", "Quantity", "Planned_Notional", "Maximum_Loss",
        "Potential_Profit", "Trade_Approved", "Risk_Decision_Reason",
    ]

    MARKET_COLUMNS = [
        "Symbol", "Date", "Time", "Open", "High", "Low", "Close", "Volume"
    ]

    MATHEMATICAL_COLUMNS = [
        "Symbol", "Latest_DateTime", "Latest_Stable_Regime",
        "Latest_Residual_ZScore", "Latest_Residual_Return",
        "Latest_Residual_Stationary", "Latest_Residual_Quality_Score",
        "Latest_Close", "Latest_Previous_Close",
        "Latest_Setup_Ready", "Latest_Setup_Accepted",
    ]

    POSITION_COLUMNS = [
        "Position_ID", "Decision_ID", "Symbol", "Direction", "Entry_Action",
        "Regime", "Rank", "Signal_Strength", "Entry_Time", "Entry_Price",
        "Entry_Market_Price", "Entry_Residual_ZScore",
        "Stop_Loss", "Target_Price", "Reward_Risk_Ratio", "Quantity",
        "Initial_Risk", "Initial_Stop_Loss", "Highest_Price_Since_Entry",
        "Lowest_Price_Since_Entry", "Maximum_Favorable_Excursion",
        "Estimated_Exit_Charges", "Estimated_Net_PnL",
        "BreakEven_Activated", "Trailing_Stop_Activated",
        "Stop_Adjustment_Reason", "Last_Checked_DateTime", "Current_Price",
        "Unrealized_PnL", "Entry_Confirmation_Count",
        "Reversal_Confirmation_Count", "Candles_Held", "Status",
    ]

    TRADE_COLUMNS = POSITION_COLUMNS + [
        "Exit_Time", "Exit_Price", "Exit_Market_Price", "Exit_Reason",
        "Gross_PnL", "Filled_Gross_PnL", "Entry_Brokerage",
        "Exit_Brokerage", "STT", "Exchange_Transaction_Charges",
        "SEBI_Charges", "GST", "Stamp_Duty", "Total_Statutory_Charges",
        "Slippage_Cost", "Total_Trading_Cost", "Transaction_Cost",
        "Net_PnL", "Return_Percent", "Holding_Candles",
        "Equity_After",
    ]

    EVENT_COLUMNS = [
        "Event_Time", "Cycle", "Symbol", "Position_ID", "Event",
        "Price", "Quantity", "Reason",
    ]

    MANUAL_LEDGER_COLUMNS = [
        "Trade_ID", "Symbol", "Entry_Action", "Regime", "Rank",
        "Entry_Time", "Entry_Price", "Stop_Loss", "Target_Price",
        "Quantity", "Status", "Current_Price", "Exit_Time", "Exit_Price",
        "Exit_Reason", "Gross_PnL", "Total_Trading_Cost", "Net_PnL",
        "Portfolio_Value", "Manual_Exit_Price",
    ]

    def __init__(
        self,
        trading_ledger_file=(
            "/home/devinderjeet/fmrss/report/TradingLedger.xlsx"
        ),
        mathematical_score_file=None,
        data_folder="/home/hadoop/shareMarket_Data",
        output_file=(
            "/home/devinderjeet/fmrss/report/"
            "ResidualStatisticalStrategy.xlsx"
        ),
        manual_ledger_file=None,
        position_journal_file=None,
        approved_sheet="Approved_Trades",
        initial_capital=100000.0,
        maximum_open_positions=2,
        transaction_cost_rate=None,
        slippage_rate=0.0001,
        brokerage_rate=0.0003,
        brokerage_cap_per_order=20.0,
        stt_sell_rate=0.00025,
        nse_transaction_charge_rate=0.0000307,
        gst_rate=0.18,
        sebi_turnover_charge_rate=0.000001,
        stamp_duty_buy_rate=0.00003,
        enable_profit_protection=True,
        break_even_trigger_r=0.50,
        break_even_cost_buffer_multiplier=1.20,
        trailing_trigger_r=1.00,
        trailing_lock_r=0.50,
        trailing_distance_r=0.50,
        force_exit_at_session_end=True,
        session_exit_time="15:14",
        exit_when_risk_approval_removed=False,
        entry_confirmation_cycles=1,
        maximum_entry_rank=2,
        maximum_entry_age_minutes=5.0,
        minimum_entry_abs_zscore=2.50,
        maximum_entry_abs_zscore=None,
        minimum_residual_quality_score=50.0,
        maximum_entry_deviation_atr=0.20,
        reversal_confirmation_cycles=2,
        reentry_cooldown_candles=6,
        stop_loss_cooldown_candles=12,
        mean_reversion_normalization_zscore=0.50,
        require_profitable_residual_normalization=True,
        minimum_normalization_net_pnl=0.0,
        residual_adverse_expansion=0.75,
        exit_on_residual_divergence=False,
        maximum_consecutive_losses=3,
        daily_loss_limit_fraction=0.01,
        mean_reversion_maximum_holding_candles=12,
        momentum_maximum_holding_candles=18,
        conservative_intrabar_order=True,
        market_tail_rows=5,
        maximum_closed_trades=10000,
        maximum_events=20000,
        memory_optimized=True,
        retain_results_in_memory=False,
        save_excel=True,
        print_details=False,
    ):
        if not np.isfinite(initial_capital) or initial_capital <= 0:
            raise ValueError("initial_capital must be finite and positive")
        if not isinstance(maximum_open_positions, (int, np.integer)) \
                or maximum_open_positions < 1:
            raise ValueError("maximum_open_positions must be a positive integer")
        for name, value in {"slippage_rate": slippage_rate}.items():
            if not np.isfinite(value) or not 0 <= value <= 0.05:
                raise ValueError(f"{name} must be in [0, 0.05]")
        for name, value in {
            "brokerage_rate": brokerage_rate,
            "stt_sell_rate": stt_sell_rate,
            "nse_transaction_charge_rate": nse_transaction_charge_rate,
            "gst_rate": gst_rate,
            "sebi_turnover_charge_rate": sebi_turnover_charge_rate,
            "stamp_duty_buy_rate": stamp_duty_buy_rate,
        }.items():
            if not np.isfinite(value) or not 0 <= value <= 0.25:
                raise ValueError(f"{name} must be in [0, 0.25]")
        if not np.isfinite(brokerage_cap_per_order) \
                or brokerage_cap_per_order < 0:
            raise ValueError("brokerage_cap_per_order must be non-negative")
        for name, value in {
            "break_even_trigger_r": break_even_trigger_r,
            "break_even_cost_buffer_multiplier": (
                break_even_cost_buffer_multiplier
            ),
            "trailing_trigger_r": trailing_trigger_r,
            "trailing_lock_r": trailing_lock_r,
            "trailing_distance_r": trailing_distance_r,
        }.items():
            if not np.isfinite(value) or value < 0:
                raise ValueError(f"{name} must be finite and non-negative")
        if trailing_trigger_r < break_even_trigger_r:
            raise ValueError(
                "trailing_trigger_r must be at least break_even_trigger_r"
            )
        if trailing_lock_r > trailing_trigger_r:
            raise ValueError(
                "trailing_lock_r cannot exceed trailing_trigger_r"
            )
        if market_tail_rows < 1:
            raise ValueError("market_tail_rows must be positive")
        for name, value in {
            "entry_confirmation_cycles": entry_confirmation_cycles,
            "maximum_entry_rank": maximum_entry_rank,
            "reversal_confirmation_cycles": reversal_confirmation_cycles,
            "reentry_cooldown_candles": reentry_cooldown_candles,
            "stop_loss_cooldown_candles": stop_loss_cooldown_candles,
            "maximum_consecutive_losses": maximum_consecutive_losses,
            "mean_reversion_maximum_holding_candles": (
                mean_reversion_maximum_holding_candles
            ),
            "momentum_maximum_holding_candles": momentum_maximum_holding_candles,
        }.items():
            if not isinstance(value, (int, np.integer)) or value < 0:
                raise ValueError(f"{name} must be a non-negative integer")
        if entry_confirmation_cycles < 1 or reversal_confirmation_cycles < 1:
            raise ValueError("confirmation cycles must be at least 1")
        for name, value in {
            "maximum_entry_age_minutes": maximum_entry_age_minutes,
            "minimum_entry_abs_zscore": minimum_entry_abs_zscore,
            "minimum_residual_quality_score": minimum_residual_quality_score,
            "maximum_entry_deviation_atr": maximum_entry_deviation_atr,
        }.items():
            if not np.isfinite(value) or value < 0:
                raise ValueError(f"{name} must be finite and non-negative")
        if maximum_entry_abs_zscore is not None and (
            not np.isfinite(maximum_entry_abs_zscore)
            or maximum_entry_abs_zscore < 0
        ):
            raise ValueError(
                "maximum_entry_abs_zscore must be finite and non-negative"
            )
        if not np.isfinite(mean_reversion_normalization_zscore) \
                or mean_reversion_normalization_zscore < 0:
            raise ValueError(
                "mean_reversion_normalization_zscore must be non-negative"
            )

        self.trading_ledger_file = str(trading_ledger_file)
        self.mathematical_score_file = (
            str(mathematical_score_file) if mathematical_score_file else None
        )
        self.data_folder = str(data_folder)
        self.output_file = str(output_file)
        self.manual_ledger_file = (
            str(manual_ledger_file) if manual_ledger_file else None
        )
        self.position_journal_file = (
            str(position_journal_file) if position_journal_file else None
        )
        self.approved_sheet = str(approved_sheet)
        self.initial_capital = float(initial_capital)
        self.maximum_open_positions = int(maximum_open_positions)
        # Accepted only so older notebooks do not fail at construction. The
        # former blanket transaction-cost percentage is intentionally unused;
        # actual equity-intraday components below replace it.
        self.transaction_cost_rate = (
            None if transaction_cost_rate is None
            else float(transaction_cost_rate)
        )
        self.slippage_rate = float(slippage_rate)
        self.brokerage_rate = float(brokerage_rate)
        self.brokerage_cap_per_order = float(brokerage_cap_per_order)
        self.stt_sell_rate = float(stt_sell_rate)
        self.nse_transaction_charge_rate = float(
            nse_transaction_charge_rate
        )
        self.gst_rate = float(gst_rate)
        self.sebi_turnover_charge_rate = float(
            sebi_turnover_charge_rate
        )
        self.stamp_duty_buy_rate = float(stamp_duty_buy_rate)
        self.enable_profit_protection = bool(enable_profit_protection)
        self.break_even_trigger_r = float(break_even_trigger_r)
        self.break_even_cost_buffer_multiplier = float(
            break_even_cost_buffer_multiplier
        )
        self.trailing_trigger_r = float(trailing_trigger_r)
        self.trailing_lock_r = float(trailing_lock_r)
        self.trailing_distance_r = float(trailing_distance_r)
        self.force_exit_at_session_end = bool(force_exit_at_session_end)
        self.session_exit_time = self._parse_time(session_exit_time)
        self.exit_when_risk_approval_removed = bool(
            exit_when_risk_approval_removed
        )
        self.entry_confirmation_cycles = int(entry_confirmation_cycles)
        self.maximum_entry_rank = int(maximum_entry_rank)
        self.maximum_entry_age_minutes = float(maximum_entry_age_minutes)
        self.minimum_entry_abs_zscore = float(minimum_entry_abs_zscore)
        self.maximum_entry_abs_zscore = (
            float(maximum_entry_abs_zscore)
            if maximum_entry_abs_zscore is not None
            else float("inf")
        )
        if self.maximum_entry_abs_zscore < self.minimum_entry_abs_zscore:
            raise ValueError(
                "maximum_entry_abs_zscore must be greater than or equal to "
                "minimum_entry_abs_zscore"
            )
        self.minimum_residual_quality_score = float(
            minimum_residual_quality_score
        )
        self.maximum_entry_deviation_atr = float(maximum_entry_deviation_atr)
        self.reversal_confirmation_cycles = int(reversal_confirmation_cycles)
        self.reentry_cooldown_candles = int(reentry_cooldown_candles)
        self.stop_loss_cooldown_candles = int(stop_loss_cooldown_candles)
        self.mean_reversion_normalization_zscore = float(
            mean_reversion_normalization_zscore
        )
        if not np.isfinite(minimum_normalization_net_pnl):
            raise ValueError("minimum_normalization_net_pnl must be finite")
        self.require_profitable_residual_normalization = bool(
            require_profitable_residual_normalization
        )
        self.minimum_normalization_net_pnl = float(
            minimum_normalization_net_pnl
        )
        if not np.isfinite(residual_adverse_expansion) \
                or residual_adverse_expansion <= 0:
            raise ValueError("residual_adverse_expansion must be positive")
        if not np.isfinite(daily_loss_limit_fraction) \
                or not 0 < daily_loss_limit_fraction <= 0.10:
            raise ValueError("daily_loss_limit_fraction must be in (0, 0.10]")
        self.residual_adverse_expansion = float(residual_adverse_expansion)
        self.exit_on_residual_divergence = bool(
            exit_on_residual_divergence
        )
        self.maximum_consecutive_losses = int(maximum_consecutive_losses)
        self.daily_loss_limit_fraction = float(daily_loss_limit_fraction)
        self.mean_reversion_maximum_holding_candles = int(
            mean_reversion_maximum_holding_candles
        )
        self.momentum_maximum_holding_candles = int(
            momentum_maximum_holding_candles
        )
        self.conservative_intrabar_order = bool(conservative_intrabar_order)
        self.market_tail_rows = int(market_tail_rows)
        self.maximum_closed_trades = int(maximum_closed_trades)
        self.maximum_events = int(maximum_events)
        self.memory_optimized = bool(memory_optimized)
        self.retain_results_in_memory = bool(retain_results_in_memory)
        self.save_excel = bool(save_excel)
        self.print_details = bool(print_details)
        self.calculation_backend = "PANDAS_TAIL_ONLY_PAPER_EXECUTION"

        self.open_positions = pd.DataFrame(columns=self.POSITION_COLUMNS)
        self.closed_trades = pd.DataFrame(columns=self.TRADE_COLUMNS)
        self.events = pd.DataFrame(columns=self.EVENT_COLUMNS)
        self.processed_decision_ids = set()
        self.realized_pnl = 0.0
        self.equity = self.initial_capital
        self.cycle_number = 0
        restored = self._restore_state()
        if not restored:
            self._restore_position_journal()

    @staticmethod
    def _parse_time(value):
        if isinstance(value, clock_time):
            return value
        parsed = pd.to_datetime(str(value), format="%H:%M", errors="coerce")
        if pd.isna(parsed):
            raise ValueError("session_exit_time must use HH:MM format")
        return parsed.time()

    @staticmethod
    def _boolean(series):
        if pd.api.types.is_bool_dtype(series):
            return series.fillna(False)
        return series.astype("string").str.strip().str.upper().isin(
            {"TRUE", "1", "YES", "Y"}
        )

    @staticmethod
    def _normalise_symbol(series):
        return series.astype("string").str.strip().str.upper()

    def _load_approved_plans(self):
        if not os.path.isfile(self.trading_ledger_file):
            raise FileNotFoundError(
                f"Trading ledger not found: {self.trading_ledger_file}"
            )
        frame = pd.read_excel(
            self.trading_ledger_file,
            sheet_name=self.approved_sheet,
            usecols=self.PLAN_COLUMNS,
            engine="openpyxl",
        )
        if frame.empty:
            return frame
        frame["Symbol"] = self._normalise_symbol(frame["Symbol"])
        frame["Trade_Action"] = (
            frame["Trade_Action"].astype("string").str.strip().str.upper()
        )
        frame["Regime"] = (
            frame["Regime"].astype("string").str.strip().str.upper()
        )
        frame["Trade_Approved"] = self._boolean(frame["Trade_Approved"])
        for column in ["Decision_DateTime", "Market_DateTime"]:
            frame[column] = pd.to_datetime(frame[column], errors="coerce")
        frame = frame.loc[
            frame["Trade_Approved"]
            & frame["Trade_Action"].isin({"BUY", "SELL_SHORT"})
        ].copy()
        frame.sort_values(["Rank", "Symbol"], inplace=True)
        frame.drop_duplicates("Symbol", keep="first", inplace=True)
        frame.reset_index(drop=True, inplace=True)
        return frame

    def _load_plan_history(self):
        try:
            frame = pd.read_excel(
                self.trading_ledger_file,
                sheet_name="Trade_Plans",
                usecols=self.PLAN_COLUMNS,
                engine="openpyxl",
            )
        except (ValueError, FileNotFoundError):
            return pd.DataFrame(columns=self.PLAN_COLUMNS)
        if frame.empty:
            return frame
        frame["Symbol"] = self._normalise_symbol(frame["Symbol"])
        frame["Trade_Action"] = (
            frame["Trade_Action"].astype("string").str.strip().str.upper()
        )
        frame["Trade_Approved"] = self._boolean(frame["Trade_Approved"])
        frame["Decision_DateTime"] = pd.to_datetime(
            frame["Decision_DateTime"], errors="coerce"
        )
        frame.sort_values(["Decision_DateTime", "Decision_ID"], inplace=True)
        frame.drop_duplicates("Decision_ID", keep="last", inplace=True)
        frame.reset_index(drop=True, inplace=True)
        return frame

    def _load_mathematical_snapshot(self):
        if not self.mathematical_score_file \
                or not os.path.isfile(self.mathematical_score_file):
            return {}
        try:
            frame = pd.read_excel(
                self.mathematical_score_file,
                sheet_name="All_Shares",
                usecols=self.MATHEMATICAL_COLUMNS,
                engine="openpyxl",
            )
        except ValueError:
            return {}
        if frame.empty:
            return {}
        frame["Symbol"] = self._normalise_symbol(frame["Symbol"])
        frame["Latest_DateTime"] = pd.to_datetime(
            frame["Latest_DateTime"], errors="coerce"
        )
        frame["Latest_Stable_Regime"] = (
            frame["Latest_Stable_Regime"]
            .astype("string").str.strip().str.upper()
        )
        frame["Latest_Residual_ZScore"] = pd.to_numeric(
            frame["Latest_Residual_ZScore"], errors="coerce"
        )
        for column in [
            "Latest_Residual_Return", "Latest_Residual_Quality_Score",
            "Latest_Close", "Latest_Previous_Close",
        ]:
            frame[column] = pd.to_numeric(frame[column], errors="coerce")
        for column in [
            "Latest_Residual_Stationary", "Latest_Setup_Ready",
            "Latest_Setup_Accepted",
        ]:
            frame[column] = self._boolean(frame[column])
        frame.drop_duplicates("Symbol", keep="last", inplace=True)
        return {
            row.Symbol: row for row in frame.itertuples(index=False)
        }

    @staticmethod
    def _consecutive_action_count(history, symbol, action):
        if history.empty:
            return 0
        rows = history.loc[history["Symbol"].eq(symbol)].sort_values(
            "Decision_DateTime", ascending=False
        )
        count = 0
        for row in rows.itertuples(index=False):
            if bool(row.Trade_Approved) and row.Trade_Action == action:
                count += 1
            else:
                break
        return count

    def _cooldown_complete(self, symbol, current_time):
        if self.reentry_cooldown_candles == 0 or self.closed_trades.empty:
            return True
        rows = self.closed_trades.loc[
            self.closed_trades["Symbol"].astype(str).eq(symbol)
        ]
        if rows.empty:
            return True
        rows = rows.copy()
        rows["_ExitTime"] = pd.to_datetime(rows["Exit_Time"], errors="coerce")
        rows.sort_values("_ExitTime", inplace=True)
        last = rows.iloc[-1]
        last_exit = last["_ExitTime"]
        if pd.isna(last_exit):
            return True
        elapsed = (pd.Timestamp(current_time) - last_exit).total_seconds() / 300.0
        cooldown = (
            self.stop_loss_cooldown_candles
            if str(last.get("Exit_Reason", "")).startswith("STOP")
            else self.reentry_cooldown_candles
        )
        return elapsed >= cooldown

    def _entry_risk_halt_reason(self, current_time):
        if self.closed_trades.empty:
            return None
        pnl = pd.to_numeric(self.closed_trades["Net_PnL"], errors="coerce")
        exits = pd.to_datetime(self.closed_trades["Exit_Time"], errors="coerce")
        order = pd.DataFrame({"pnl": pnl, "exit": exits}).dropna(
            subset=["exit"]
        ).sort_values("exit")
        if order.empty:
            return None
        today = pd.Timestamp(current_time).date()
        today_rows = order.loc[order["exit"].dt.date.eq(today)]
        consecutive_losses = 0
        for value in reversed(today_rows["pnl"].fillna(0.0).tolist()):
            if value <= 0:
                consecutive_losses += 1
            else:
                break
        if consecutive_losses >= self.maximum_consecutive_losses:
            return "MAXIMUM_CONSECUTIVE_LOSSES"
        daily_pnl = today_rows["pnl"].sum()
        if daily_pnl <= -self.initial_capital * self.daily_loss_limit_fraction:
            return "DAILY_LOSS_LIMIT"
        return None

    def _read_market_tail(self, symbol):
        path = os.path.join(self.data_folder, f"{symbol}_5mins.txt")
        if not os.path.isfile(path):
            raise FileNotFoundError(f"Market data missing for {symbol}: {path}")
        with open(path, "r", encoding="utf-8", errors="replace") as handle:
            lines = deque(handle, maxlen=self.market_tail_rows + 2)
        frame = pd.read_csv(
            io.StringIO("".join(lines)), names=self.MARKET_COLUMNS, header=None
        )
        frame["DateTime"] = pd.to_datetime(
            frame["Date"].astype(str).str.strip()
            + " " + frame["Time"].astype(str).str.strip(),
            errors="coerce", format="mixed",
        )
        for column in ["Open", "High", "Low", "Close", "Volume"]:
            frame[column] = pd.to_numeric(frame[column], errors="coerce")
        frame.dropna(
            subset=["DateTime", "Open", "High", "Low", "Close"], inplace=True
        )
        frame.sort_values("DateTime", inplace=True)
        frame.drop_duplicates("DateTime", keep="last", inplace=True)
        if frame.empty:
            raise ValueError(f"No valid completed candle for {symbol}")
        return frame.tail(self.market_tail_rows).copy()

    def _restore_state(self):
        if not os.path.isfile(self.output_file):
            return False
        try:
            open_frame = pd.read_excel(
                self.output_file, sheet_name="Open_Positions", engine="openpyxl"
            ).reindex(columns=self.POSITION_COLUMNS)
            closed_frame = pd.read_excel(
                self.output_file, sheet_name="Closed_Trades", engine="openpyxl"
            ).reindex(columns=self.TRADE_COLUMNS)
            event_frame = pd.read_excel(
                self.output_file, sheet_name="Event_Log", engine="openpyxl"
            ).reindex(columns=self.EVENT_COLUMNS)
            self.open_positions = open_frame
            numeric_defaults = {
                "Initial_Stop_Loss": self.open_positions.get("Stop_Loss"),
                "Highest_Price_Since_Entry": self.open_positions.get(
                    "Current_Price"
                ),
                "Lowest_Price_Since_Entry": self.open_positions.get(
                    "Current_Price"
                ),
                "Maximum_Favorable_Excursion": 0.0,
                "Estimated_Exit_Charges": 0.0,
                "Estimated_Net_PnL": 0.0,
            }
            for column, default in numeric_defaults.items():
                values = pd.to_numeric(
                    self.open_positions[column], errors="coerce"
                )
                if isinstance(default, pd.Series):
                    values = values.fillna(
                        pd.to_numeric(default, errors="coerce")
                    )
                else:
                    values = values.fillna(default)
                self.open_positions[column] = values.astype(float)
            for column in [
                "BreakEven_Activated", "Trailing_Stop_Activated"
            ]:
                self.open_positions[column] = self._boolean(
                    self.open_positions[column]
                )
            self.open_positions["Stop_Adjustment_Reason"] = (
                self.open_positions["Stop_Adjustment_Reason"]
                .astype("string")
                .fillna("INITIAL_STOP")
            )
            for column in [
                "Entry_Confirmation_Count", "Reversal_Confirmation_Count",
                "Candles_Held",
            ]:
                self.open_positions[column] = pd.to_numeric(
                    self.open_positions[column], errors="coerce"
                ).fillna(0).astype(np.int32)
            self.closed_trades = closed_frame.tail(
                self.maximum_closed_trades
            ).copy()
            self.events = event_frame.tail(self.maximum_events).copy()
            for frame in [self.open_positions, self.closed_trades]:
                if "Decision_ID" in frame:
                    self.processed_decision_ids.update(
                        frame["Decision_ID"].dropna().astype(str)
                    )
            if not self.closed_trades.empty:
                pnl = pd.to_numeric(
                    self.closed_trades["Net_PnL"], errors="coerce"
                ).fillna(0.0)
                self.realized_pnl = float(pnl.sum())
                self.equity = self.initial_capital + self.realized_pnl
            if not self.events.empty:
                cycles = pd.to_numeric(self.events["Cycle"], errors="coerce")
                self.cycle_number = int(cycles.max()) if cycles.notna().any() else 0
            return True
        except (ValueError, OSError):
            # An unrelated/legacy workbook is not valid strategy state.
            self.open_positions = pd.DataFrame(columns=self.POSITION_COLUMNS)
            self.closed_trades = pd.DataFrame(columns=self.TRADE_COLUMNS)
            self.events = pd.DataFrame(columns=self.EVENT_COLUMNS)
            return False

    def _restore_position_journal(self):
        """Recover execution state when disposable report files were removed."""
        if not self.position_journal_file \
                or not os.path.isfile(self.position_journal_file):
            return False
        try:
            open_frame = pd.read_excel(
                self.position_journal_file,
                sheet_name="Open_Positions",
                engine="openpyxl",
            ).reindex(columns=self.POSITION_COLUMNS)
            closed_frame = pd.read_excel(
                self.position_journal_file,
                sheet_name="Closed_Trades",
                engine="openpyxl",
            ).reindex(columns=self.TRADE_COLUMNS)
            event_frame = pd.read_excel(
                self.position_journal_file,
                sheet_name="Event_Log",
                engine="openpyxl",
            ).reindex(columns=self.EVENT_COLUMNS)
            metadata = pd.read_excel(
                self.position_journal_file,
                sheet_name="Recovery_Metadata",
                engine="openpyxl",
            )

            self.open_positions = open_frame
            for column in [
                "Entry_Time", "Last_Checked_DateTime"
            ]:
                self.open_positions[column] = pd.to_datetime(
                    self.open_positions[column], errors="coerce"
                )
            for column in [
                "Direction", "Quantity", "Entry_Confirmation_Count",
                "Reversal_Confirmation_Count", "Candles_Held",
            ]:
                self.open_positions[column] = pd.to_numeric(
                    self.open_positions[column], errors="coerce"
                ).fillna(0).astype(np.int32)
            for column in [
                "Entry_Price", "Entry_Market_Price",
                "Entry_Residual_ZScore", "Stop_Loss", "Target_Price",
                "Reward_Risk_Ratio", "Initial_Risk", "Initial_Stop_Loss",
                "Highest_Price_Since_Entry", "Lowest_Price_Since_Entry",
                "Maximum_Favorable_Excursion", "Estimated_Exit_Charges",
                "Estimated_Net_PnL", "Current_Price", "Unrealized_PnL",
            ]:
                self.open_positions[column] = pd.to_numeric(
                    self.open_positions[column], errors="coerce"
                )
            for column in [
                "BreakEven_Activated", "Trailing_Stop_Activated"
            ]:
                self.open_positions[column] = self._boolean(
                    self.open_positions[column]
                )
            self.open_positions = self.open_positions.loc[
                self.open_positions["Status"]
                .astype("string").str.upper().eq("OPEN")
                & self.open_positions["Position_ID"].notna()
                & self.open_positions["Symbol"].notna()
            ].copy()
            self.open_positions.drop_duplicates(
                "Position_ID", keep="last", inplace=True
            )
            self.open_positions.reset_index(drop=True, inplace=True)

            self.closed_trades = closed_frame.tail(
                self.maximum_closed_trades
            ).copy()
            self.events = event_frame.tail(self.maximum_events).copy()
            for frame in [self.open_positions, self.closed_trades]:
                self.processed_decision_ids.update(
                    frame["Decision_ID"].dropna().astype(str)
                )

            values = {}
            if {"Metric", "Value"}.issubset(metadata.columns):
                values = dict(zip(metadata["Metric"], metadata["Value"]))
            realized = pd.to_numeric(
                values.get("Realized PnL"), errors="coerce"
            )
            if pd.isna(realized):
                realized = pd.to_numeric(
                    self.closed_trades.get(
                        "Net_PnL", pd.Series(dtype=float)
                    ), errors="coerce"
                ).fillna(0.0).sum()
            self.realized_pnl = float(realized)
            self.equity = self.initial_capital + self.realized_pnl
            cycle = pd.to_numeric(values.get("Cycle"), errors="coerce")
            if pd.notna(cycle):
                self.cycle_number = int(cycle)
            elif not self.events.empty:
                cycles = pd.to_numeric(self.events["Cycle"], errors="coerce")
                self.cycle_number = (
                    int(cycles.max()) if cycles.notna().any() else 0
                )
            return True
        except (ValueError, OSError, KeyError, TypeError):
            self.open_positions = pd.DataFrame(columns=self.POSITION_COLUMNS)
            self.closed_trades = pd.DataFrame(columns=self.TRADE_COLUMNS)
            self.events = pd.DataFrame(columns=self.EVENT_COLUMNS)
            self.processed_decision_ids.clear()
            self.realized_pnl = 0.0
            self.equity = self.initial_capital
            self.cycle_number = 0
            return False

    def _event(self, timestamp, symbol, position_id, event, price, quantity, reason):
        row = pd.DataFrame(
            [[timestamp, self.cycle_number, symbol, position_id, event,
              price, quantity, reason]],
            columns=self.EVENT_COLUMNS,
        )
        if self.events.empty:
            self.events = row
        else:
            self.events = pd.concat(
                [self.events, row], ignore_index=True
            )
        self.events = self.events.tail(self.maximum_events)

    def _equity_intraday_charges(
        self, direction, entry_fill, exit_fill, quantity
    ):
        entry_turnover = abs(float(entry_fill) * int(quantity))
        exit_turnover = abs(float(exit_fill) * int(quantity))
        total_turnover = entry_turnover + exit_turnover

        entry_brokerage = min(
            entry_turnover * self.brokerage_rate,
            self.brokerage_cap_per_order,
        )
        exit_brokerage = min(
            exit_turnover * self.brokerage_rate,
            self.brokerage_cap_per_order,
        )

        # Long: entry is BUY and exit is SELL. Short: entry is SELL and
        # cover is BUY. STT is sell-side only; stamp duty is buy-side only.
        if int(direction) == 1:
            buy_turnover = entry_turnover
            sell_turnover = exit_turnover
        else:
            buy_turnover = exit_turnover
            sell_turnover = entry_turnover

        stt = sell_turnover * self.stt_sell_rate
        exchange = total_turnover * self.nse_transaction_charge_rate
        sebi = total_turnover * self.sebi_turnover_charge_rate
        gst = (
            entry_brokerage + exit_brokerage + exchange + sebi
        ) * self.gst_rate
        stamp = buy_turnover * self.stamp_duty_buy_rate
        statutory = (
            entry_brokerage + exit_brokerage + stt
            + exchange + sebi + gst + stamp
        )
        return {
            "Entry_Brokerage": entry_brokerage,
            "Exit_Brokerage": exit_brokerage,
            "STT": stt,
            "Exchange_Transaction_Charges": exchange,
            "SEBI_Charges": sebi,
            "GST": gst,
            "Stamp_Duty": stamp,
            "Total_Statutory_Charges": statutory,
        }

    def _estimate_open_trade(self, position, market_price):
        direction = int(position["Direction"])
        quantity = int(position["Quantity"])
        entry_fill = float(position["Entry_Price"])
        entry_market = float(pd.to_numeric(
            position.get("Entry_Market_Price", entry_fill), errors="coerce"
        ))
        if not np.isfinite(entry_market):
            entry_market = entry_fill
        market_price = float(market_price)
        estimated_exit_fill = market_price * (
            1.0 - direction * self.slippage_rate
        )
        gross = (
            (market_price - entry_market) * direction * quantity
        )
        filled_gross = (
            (estimated_exit_fill - entry_fill) * direction * quantity
        )
        slippage = max(0.0, gross - filled_gross)
        statutory = self._equity_intraday_charges(
            direction, entry_fill, estimated_exit_fill, quantity
        )["Total_Statutory_Charges"]
        estimated_cost = statutory + slippage
        return gross, estimated_cost, gross - estimated_cost

    def _update_profit_protection(self, index, candle):
        if not self.enable_profit_protection \
                or index not in self.open_positions.index:
            return
        position = self.open_positions.loc[index]
        direction = int(position["Direction"])
        quantity = int(position["Quantity"])
        entry_market = float(pd.to_numeric(
            position.get("Entry_Market_Price", position["Entry_Price"]),
            errors="coerce",
        ))
        initial_risk = float(position["Initial_Risk"])
        if quantity < 1 or initial_risk <= 0 or not np.isfinite(entry_market):
            return

        previous_high = float(pd.to_numeric(
            position.get("Highest_Price_Since_Entry", entry_market),
            errors="coerce",
        ))
        previous_low = float(pd.to_numeric(
            position.get("Lowest_Price_Since_Entry", entry_market),
            errors="coerce",
        ))
        if not np.isfinite(previous_high):
            previous_high = entry_market
        if not np.isfinite(previous_low):
            previous_low = entry_market
        highest = max(previous_high, float(candle.High))
        lowest = min(previous_low, float(candle.Low))
        favorable_price = highest if direction == 1 else lowest
        favorable_move = (
            (favorable_price - entry_market) * direction
        )
        mfe = max(0.0, favorable_move * quantity)
        risk_per_share = initial_risk / quantity

        _, estimated_cost, estimated_net = self._estimate_open_trade(
            position, float(candle.Close)
        )
        # Estimate the round-trip cost at entry price for a stable break-even
        # threshold that does not move adversely as the market fluctuates.
        _, entry_cost, _ = self._estimate_open_trade(position, entry_market)
        cost_buffer_per_share = (
            entry_cost * self.break_even_cost_buffer_multiplier / quantity
        )

        old_stop = float(position["Stop_Loss"])
        new_stop = old_stop
        break_even_active = bool(position.get("BreakEven_Activated", False))
        trailing_active = bool(position.get("Trailing_Stop_Activated", False))
        reason = str(position.get("Stop_Adjustment_Reason", "INITIAL_STOP"))

        if mfe >= self.break_even_trigger_r * initial_risk:
            break_even_stop = (
                entry_market + direction * cost_buffer_per_share
            )
            new_stop = (
                max(new_stop, break_even_stop)
                if direction == 1
                else min(new_stop, break_even_stop)
            )
            break_even_active = True
            reason = "COST_ADJUSTED_BREAK_EVEN"

        if mfe >= self.trailing_trigger_r * initial_risk:
            locked_profit_stop = entry_market + direction * (
                self.trailing_lock_r * risk_per_share
                + cost_buffer_per_share
            )
            dynamic_stop = favorable_price - direction * (
                self.trailing_distance_r * risk_per_share
            )
            candidate = (
                max(locked_profit_stop, dynamic_stop)
                if direction == 1
                else min(locked_profit_stop, dynamic_stop)
            )
            new_stop = (
                max(new_stop, candidate)
                if direction == 1
                else min(new_stop, candidate)
            )
            trailing_active = True
            reason = "TRAILING_STOP"

        self.open_positions.loc[index, "Stop_Loss"] = new_stop
        self.open_positions.loc[index, "Highest_Price_Since_Entry"] = highest
        self.open_positions.loc[index, "Lowest_Price_Since_Entry"] = lowest
        self.open_positions.loc[
            index, "Maximum_Favorable_Excursion"
        ] = mfe
        self.open_positions.loc[
            index, "Estimated_Exit_Charges"
        ] = estimated_cost
        self.open_positions.loc[index, "Estimated_Net_PnL"] = estimated_net
        self.open_positions.loc[
            index, "BreakEven_Activated"
        ] = break_even_active
        self.open_positions.loc[
            index, "Trailing_Stop_Activated"
        ] = trailing_active
        self.open_positions.loc[index, "Stop_Adjustment_Reason"] = reason

    def _close(
        self,
        index,
        exit_time,
        exit_price,
        reason,
        holding_candles=0,
        apply_exit_slippage=True,
    ):
        position = self.open_positions.loc[index].copy()
        direction = int(position["Direction"])
        quantity = int(position["Quantity"])
        entry = float(position["Entry_Price"])
        entry_market = float(pd.to_numeric(
            position.get("Entry_Market_Price", entry), errors="coerce"
        ))
        if not np.isfinite(entry_market):
            entry_market = entry
        raw_exit = float(exit_price)
        # A manually entered exit is treated as the actual executed fill.
        # Strategy-generated exits retain adverse paper slippage.
        filled_exit = (
            raw_exit * (1.0 - direction * self.slippage_rate)
            if apply_exit_slippage else raw_exit
        )
        gross = (raw_exit - entry_market) * direction * quantity
        filled_gross = (filled_exit - entry) * direction * quantity
        slippage_cost = max(0.0, gross - filled_gross)
        charges = self._equity_intraday_charges(
            direction, entry, filled_exit, quantity
        )
        total_trading_cost = (
            charges["Total_Statutory_Charges"] + slippage_cost
        )
        net = gross - total_trading_cost
        self.realized_pnl += net
        self.equity = self.initial_capital + self.realized_pnl
        closed = position.to_dict()
        closed.update(
            {
                "Current_Price": filled_exit,
                "Unrealized_PnL": 0.0,
                "Status": "CLOSED",
                "Exit_Time": exit_time,
                "Exit_Price": filled_exit,
                "Exit_Market_Price": raw_exit,
                "Exit_Reason": reason,
                "Gross_PnL": gross,
                "Filled_Gross_PnL": filled_gross,
                **charges,
                "Slippage_Cost": slippage_cost,
                "Total_Trading_Cost": total_trading_cost,
                # Backward-compatible alias used by older analysis notebooks.
                "Transaction_Cost": total_trading_cost,
                "Net_PnL": net,
                "Return_Percent": (
                    net / max(abs(entry * quantity), 1e-12) * 100.0
                ),
                "Holding_Candles": int(holding_candles),
                "Equity_After": self.equity,
            }
        )
        closed_row = pd.DataFrame([closed], columns=self.TRADE_COLUMNS)
        if self.closed_trades.empty:
            self.closed_trades = closed_row
        else:
            self.closed_trades = pd.concat(
                [self.closed_trades, closed_row], ignore_index=True
            )
        self.closed_trades = self.closed_trades.reindex(
            columns=self.TRADE_COLUMNS
        ).tail(self.maximum_closed_trades)
        self._event(
            exit_time, position["Symbol"], position["Position_ID"],
            "EXIT", filled_exit, quantity, reason,
        )
        self.open_positions.drop(index=index, inplace=True)

    def _process_manual_exits(self):
        """Close positions whose manual fill was entered in TradingLedge.xlsx."""
        if not self.manual_ledger_file or self.open_positions.empty \
                or not os.path.isfile(self.manual_ledger_file):
            return
        try:
            manual = pd.read_excel(
                self.manual_ledger_file,
                sheet_name="Open_Trades",
                usecols=["Trade_ID", "Manual_Exit_Price"],
                engine="openpyxl",
            )
        except (ValueError, OSError, PermissionError):
            return
        if manual.empty:
            return
        manual["Trade_ID"] = manual["Trade_ID"].astype("string").str.strip()
        manual["Manual_Exit_Price"] = pd.to_numeric(
            manual["Manual_Exit_Price"], errors="coerce"
        )
        requested = manual.loc[
            manual["Trade_ID"].notna()
            & manual["Manual_Exit_Price"].gt(0)
        ]
        for request in requested.itertuples(index=False):
            matches = self.open_positions.index[
                self.open_positions["Position_ID"].astype(str).eq(
                    str(request.Trade_ID)
                )
            ]
            if len(matches) != 1:
                continue
            index = matches[0]
            symbol = str(self.open_positions.loc[index, "Symbol"])
            try:
                latest = self._read_market_tail(symbol).iloc[-1]
                exit_time = latest["DateTime"]
            except Exception:
                exit_time = pd.Timestamp.now()
            candles_held = int(pd.to_numeric(
                self.open_positions.loc[index, "Candles_Held"],
                errors="coerce",
            ))
            self._close(
                index=index,
                exit_time=exit_time,
                exit_price=float(request.Manual_Exit_Price),
                reason="Sold Manually",
                holding_candles=candles_held,
                apply_exit_slippage=False,
            )
        del manual, requested

    def _monitor_open_positions(self, plans, plan_history, mathematical):
        approved_by_symbol = {
            row.Symbol: row for row in plans.itertuples(index=False)
        }
        for index in list(self.open_positions.index):
            if index not in self.open_positions.index:
                continue
            position = self.open_positions.loc[index]
            symbol = str(position["Symbol"])
            try:
                candles = self._read_market_tail(symbol)
                last_checked = pd.to_datetime(
                    position["Last_Checked_DateTime"], errors="coerce"
                )
                new_candles = candles.loc[
                    candles["DateTime"].gt(last_checked)
                    if pd.notna(last_checked)
                    else pd.Series(True, index=candles.index)
                ]
                plan = approved_by_symbol.get(symbol)
                if plan is None and self.exit_when_risk_approval_removed:
                    latest = candles.iloc[-1]
                    self._close(
                        index, latest["DateTime"], latest["Close"],
                        "RISK_APPROVAL_REMOVED", len(new_candles),
                    )
                    continue
                exited = False
                candles_held = int(pd.to_numeric(
                    position.get("Candles_Held", 0), errors="coerce"
                ) if pd.notna(position.get("Candles_Held", 0)) else 0)
                for sequence, candle in enumerate(
                    new_candles.itertuples(index=False), start=1
                ):
                    candles_held += 1
                    position = self.open_positions.loc[index]
                    direction = int(position["Direction"])
                    stop = float(position["Stop_Loss"])
                    target = float(position["Target_Price"])
                    stop_hit = candle.Low <= stop if direction == 1 else candle.High >= stop
                    target_hit = candle.High >= target if direction == 1 else candle.Low <= target
                    if stop_hit and target_hit:
                        if self.conservative_intrabar_order:
                            self._close(index, candle.DateTime, stop,
                                        "STOP_AND_TARGET_HIT_STOP_ASSUMED",
                                        candles_held)
                        else:
                            self._close(index, candle.DateTime, target,
                                        "STOP_AND_TARGET_HIT_TARGET_ASSUMED",
                                        candles_held)
                        exited = True
                        break
                    if stop_hit:
                        stop_reason = str(position.get(
                            "Stop_Adjustment_Reason", "INITIAL_STOP"
                        ))
                        exit_reason = (
                            "TRAILING_STOP"
                            if stop_reason == "TRAILING_STOP"
                            else (
                                "COST_ADJUSTED_BREAK_EVEN"
                                if stop_reason == "COST_ADJUSTED_BREAK_EVEN"
                                else "STOP_LOSS"
                            )
                        )
                        self._close(
                            index, candle.DateTime, stop, exit_reason,
                            candles_held,
                        )
                        exited = True
                        break
                    # The candle's high/low can activate protection, but the
                    # adjusted stop becomes executable from the next candle.
                    # This prevents same-candle high/low ordering look-ahead.
                    self._update_profit_protection(index, candle)
                    if target_hit:
                        self._close(index, candle.DateTime, target,
                                    "TARGET_REACHED", candles_held)
                        exited = True
                        break
                    if self.force_exit_at_session_end \
                            and candle.DateTime.time() >= self.session_exit_time:
                        self._close(
                            index, candle.DateTime, candle.Close,
                            "SESSION_END", candles_held,
                        )
                        exited = True
                        break
                if not exited and index in self.open_positions.index:
                    latest = candles.iloc[-1]
                    direction = int(position["Direction"])
                    quantity = int(position["Quantity"])
                    self.open_positions.loc[index, "Candles_Held"] = candles_held
                    self.open_positions.loc[index, "Last_Checked_DateTime"] = latest["DateTime"]
                    self.open_positions.loc[index, "Current_Price"] = latest["Close"]
                    self.open_positions.loc[index, "Unrealized_PnL"] = (
                        (float(latest["Close"]) - float(position["Entry_Price"]))
                        * direction * quantity
                    )
                    _, estimated_cost, estimated_net = (
                        self._estimate_open_trade(position, latest["Close"])
                    )
                    self.open_positions.loc[
                        index, "Estimated_Exit_Charges"
                    ] = estimated_cost
                    self.open_positions.loc[
                        index, "Estimated_Net_PnL"
                    ] = estimated_net

                    # Mathematical exits are evaluated only after price-based
                    # stop, target and session exits have been checked.
                    snapshot = mathematical.get(symbol)
                    if snapshot is not None and bool(snapshot.Latest_Setup_Ready):
                        zscore = float(snapshot.Latest_Residual_ZScore)
                        entry_zscore = float(pd.to_numeric(
                            position.get("Entry_Residual_ZScore", np.nan),
                            errors="coerce",
                        ))
                        if (
                            self.exit_on_residual_divergence
                            and
                            position["Regime"] == "MEAN_REVERSION"
                            and np.isfinite(zscore)
                            and np.isfinite(entry_zscore)
                            and abs(zscore) >= abs(entry_zscore) + self.residual_adverse_expansion
                        ):
                            self._close(
                                index, latest["DateTime"], latest["Close"],
                                "RESIDUAL_DIVERGENCE_EXPANDED", candles_held,
                            )
                            exited = True
                        if (not exited and
                            position["Regime"] == "MEAN_REVERSION"
                            and np.isfinite(zscore)
                            and abs(zscore)
                            <= self.mean_reversion_normalization_zscore
                            and candles_held > 0
                            and (
                                not self.require_profitable_residual_normalization
                                or estimated_net
                                >= self.minimum_normalization_net_pnl
                            )
                        ):
                            self._close(
                                index, latest["DateTime"], latest["Close"],
                                "PROFITABLE_RESIDUAL_NORMALIZATION",
                                candles_held,
                            )
                            exited = True

                    if not exited and index in self.open_positions.index:
                        reversal_count = 0
                        if plan is not None:
                            desired = 1 if plan.Trade_Action == "BUY" else -1
                            if desired != direction:
                                reversal_count = self._consecutive_action_count(
                                    plan_history, symbol, plan.Trade_Action
                                )
                        self.open_positions.loc[
                            index, "Reversal_Confirmation_Count"
                        ] = reversal_count
                        if reversal_count >= self.reversal_confirmation_cycles:
                            self._close(
                                index, latest["DateTime"], latest["Close"],
                                "CONFIRMED_SIGNAL_REVERSAL", candles_held,
                            )
                            exited = True

                    if not exited and index in self.open_positions.index:
                        maximum_holding = (
                            self.mean_reversion_maximum_holding_candles
                            if position["Regime"] == "MEAN_REVERSION"
                            else self.momentum_maximum_holding_candles
                        )
                        if maximum_holding and candles_held >= maximum_holding:
                            self._close(
                                index, latest["DateTime"], latest["Close"],
                                "MAXIMUM_HOLDING_PERIOD", candles_held,
                            )
                del candles, new_candles
            except Exception as error:
                self._event(
                    pd.Timestamp.now(), symbol, position["Position_ID"],
                    "MONITOR_ERROR", np.nan, position["Quantity"],
                    f"{type(error).__name__}: {error}",
                )

    def _open_new_positions(self, plans, plan_history, mathematical):
        open_symbols = set(self.open_positions["Symbol"].astype(str))
        available = self.maximum_open_positions - len(self.open_positions)
        if available <= 0:
            return
        for plan in plans.itertuples(index=False):
            decision_id = str(plan.Decision_ID)
            pending_position_id = f"{decision_id}_{plan.Trade_Action}"
            if plan.Symbol in open_symbols or decision_id in self.processed_decision_ids:
                continue
            if available <= 0:
                break
            try:
                candles = self._read_market_tail(plan.Symbol)
                latest = candles.iloc[-1]
                snapshot = mathematical.get(plan.Symbol)
                rejection = None
                direction = 1 if plan.Trade_Action == "BUY" else -1
                entry_age = (
                    (latest["DateTime"] - plan.Decision_DateTime).total_seconds()
                    / 60.0
                    if pd.notna(latest["DateTime"])
                    and pd.notna(plan.Decision_DateTime)
                    else np.nan
                )
                if not np.isfinite(entry_age) or entry_age < 0:
                    rejection = "INVALID_ENTRY_TIMESTAMP"
                elif entry_age > self.maximum_entry_age_minutes:
                    rejection = "STALE_SETUP_AT_ENTRY"
                elif not np.isfinite(plan.Rank) or plan.Rank > self.maximum_entry_rank:
                    rejection = "OUTSIDE_ENTRY_RANK_LIMIT"
                elif snapshot is None:
                    rejection = "MATHEMATICAL_SNAPSHOT_MISSING"
                else:
                    zscore = float(snapshot.Latest_Residual_ZScore)
                    residual_return = float(snapshot.Latest_Residual_Return)
                    quality = float(snapshot.Latest_Residual_Quality_Score)
                    expected_regime = str(snapshot.Latest_Stable_Regime).upper()
                    setup_valid = (
                        bool(snapshot.Latest_Setup_Ready)
                        and bool(snapshot.Latest_Setup_Accepted)
                        and bool(snapshot.Latest_Residual_Stationary)
                        and expected_regime == "MEAN_REVERSION"
                    )
                    residual_direction_valid = (
                        (
                            direction == 1
                            and -self.maximum_entry_abs_zscore
                            <= zscore <= -self.minimum_entry_abs_zscore
                            and residual_return > 0
                        )
                        or (
                            direction == -1
                            and self.minimum_entry_abs_zscore
                            <= zscore <= self.maximum_entry_abs_zscore
                            and residual_return < 0
                        )
                    )
                    price_turn_valid = (
                        (direction == 1 and snapshot.Latest_Close > snapshot.Latest_Previous_Close)
                        or (direction == -1 and snapshot.Latest_Close < snapshot.Latest_Previous_Close)
                    )
                    if not setup_valid:
                        rejection = "SETUP_NO_LONGER_ACCEPTED"
                    elif not np.isfinite(quality) or quality < self.minimum_residual_quality_score:
                        rejection = "RESIDUAL_QUALITY_FAILED_AT_ENTRY"
                    elif not residual_direction_valid:
                        rejection = "RESIDUAL_ENTRY_CONFIRMATION_FAILED"
                    elif not price_turn_valid:
                        rejection = "PRICE_TURN_FAILED_AT_ENTRY"
                    else:
                        deviation = abs(float(latest["Close"]) - float(plan.Entry_Price))
                        maximum_deviation = (
                            self.maximum_entry_deviation_atr * float(plan.ATR_14)
                        )
                        if not np.isfinite(maximum_deviation) or deviation > maximum_deviation:
                            rejection = "ENTRY_PRICE_MOVED_TOO_FAR"
                if rejection:
                    self.processed_decision_ids.add(decision_id)
                    self._event(
                        latest["DateTime"], plan.Symbol, pending_position_id,
                        "ENTRY_SKIPPED",
                        latest["Close"], plan.Quantity, rejection,
                    )
                    del candles
                    continue
                if self.force_exit_at_session_end \
                        and latest["DateTime"].time() >= self.session_exit_time:
                    self.processed_decision_ids.add(decision_id)
                    self._event(
                        latest["DateTime"], plan.Symbol, pending_position_id,
                        "ENTRY_SKIPPED",
                        latest["Close"], plan.Quantity, "SESSION_ENDED",
                    )
                    del candles
                    continue
                confirmation_count = self._consecutive_action_count(
                    plan_history, plan.Symbol, plan.Trade_Action
                )
                if confirmation_count < self.entry_confirmation_cycles:
                    self.processed_decision_ids.add(decision_id)
                    self._event(
                        latest["DateTime"], plan.Symbol, pending_position_id,
                        "ENTRY_WAITING",
                        latest["Close"], plan.Quantity,
                        f"CONFIRMATION_{confirmation_count}_OF_"
                        f"{self.entry_confirmation_cycles}",
                    )
                    del candles
                    continue
                if not self._cooldown_complete(plan.Symbol, latest["DateTime"]):
                    self.processed_decision_ids.add(decision_id)
                    self._event(
                        latest["DateTime"], plan.Symbol, pending_position_id,
                        "ENTRY_WAITING",
                        latest["Close"], plan.Quantity, "REENTRY_COOLDOWN",
                    )
                    del candles
                    continue
                halt_reason = self._entry_risk_halt_reason(latest["DateTime"])
                if halt_reason:
                    self.processed_decision_ids.add(decision_id)
                    self._event(
                        latest["DateTime"], plan.Symbol, pending_position_id,
                        "ENTRY_SKIPPED",
                        latest["Close"], plan.Quantity, halt_reason,
                    )
                    del candles
                    continue
                # Adverse paper slippage on entry.
                entry = float(latest["Close"]) * (
                    1.0 + direction * self.slippage_rate
                )
                distance = float(plan.Stop_Distance)
                rr = float(plan.Reward_Risk_Ratio)
                quantity = int(plan.Quantity)
                if not np.isfinite(entry) or not np.isfinite(distance) \
                        or distance <= 0 or quantity < 1:
                    raise ValueError("invalid approved trade plan")
                stop = entry - direction * distance
                target = entry + direction * rr * distance
                position_id = pending_position_id
                position = {
                    "Position_ID": position_id,
                    "Decision_ID": decision_id,
                    "Symbol": plan.Symbol,
                    "Direction": direction,
                    "Entry_Action": plan.Trade_Action,
                    "Regime": plan.Regime,
                    "Rank": plan.Rank,
                    "Signal_Strength": plan.Signal_Strength,
                    "Entry_Time": latest["DateTime"],
                    "Entry_Price": entry,
                    "Entry_Market_Price": float(latest["Close"]),
                    "Entry_Residual_ZScore": (
                        float(mathematical[plan.Symbol].Latest_Residual_ZScore)
                        if plan.Symbol in mathematical
                        else np.nan
                    ),
                    "Stop_Loss": stop,
                    "Target_Price": target,
                    "Reward_Risk_Ratio": rr,
                    "Quantity": quantity,
                    "Initial_Risk": distance * quantity,
                    "Initial_Stop_Loss": stop,
                    "Highest_Price_Since_Entry": float(latest["Close"]),
                    "Lowest_Price_Since_Entry": float(latest["Close"]),
                    "Maximum_Favorable_Excursion": 0.0,
                    "Estimated_Exit_Charges": 0.0,
                    "Estimated_Net_PnL": 0.0,
                    "BreakEven_Activated": False,
                    "Trailing_Stop_Activated": False,
                    "Stop_Adjustment_Reason": "INITIAL_STOP",
                    "Last_Checked_DateTime": latest["DateTime"],
                    "Current_Price": entry,
                    "Unrealized_PnL": 0.0,
                    "Entry_Confirmation_Count": confirmation_count,
                    "Reversal_Confirmation_Count": 0,
                    "Candles_Held": 0,
                    "Status": "OPEN",
                }
                position_row = pd.DataFrame(
                    [position], columns=self.POSITION_COLUMNS
                )
                if self.open_positions.empty:
                    self.open_positions = position_row
                else:
                    self.open_positions = pd.concat(
                        [self.open_positions, position_row],
                        ignore_index=True,
                    )
                self.open_positions = self.open_positions.reindex(
                    columns=self.POSITION_COLUMNS
                )
                new_index = self.open_positions.index[-1]
                _, estimated_cost, estimated_net = self._estimate_open_trade(
                    self.open_positions.loc[new_index], latest["Close"]
                )
                self.open_positions.loc[
                    new_index, "Estimated_Exit_Charges"
                ] = estimated_cost
                self.open_positions.loc[
                    new_index, "Estimated_Net_PnL"
                ] = estimated_net
                self.processed_decision_ids.add(decision_id)
                open_symbols.add(plan.Symbol)
                available -= 1
                self._event(
                    latest["DateTime"], plan.Symbol, position_id,
                    "ENTRY", entry, quantity, "RISK_APPROVED",
                )
                del candles
            except Exception as error:
                self.processed_decision_ids.add(decision_id)
                self._event(
                    pd.Timestamp.now(), plan.Symbol, "", "ENTRY_ERROR",
                    np.nan, getattr(plan, "Quantity", 0),
                    f"{type(error).__name__}: {error}",
                )

    def _summary(self):
        net_series = pd.to_numeric(
            self.closed_trades.get("Net_PnL", pd.Series(dtype=float)),
            errors="coerce",
        ).fillna(0.0)
        unrealized = pd.to_numeric(
            self.open_positions.get("Unrealized_PnL", pd.Series(dtype=float)),
            errors="coerce",
        ).fillna(0.0).sum()
        statutory_costs = pd.to_numeric(
            self.closed_trades.get(
                "Total_Statutory_Charges", pd.Series(dtype=float)
            ),
            errors="coerce",
        ).fillna(0.0).sum()
        slippage_costs = pd.to_numeric(
            self.closed_trades.get(
                "Slippage_Cost", pd.Series(dtype=float)
            ),
            errors="coerce",
        ).fillna(0.0).sum()
        total_costs = pd.to_numeric(
            self.closed_trades.get(
                "Total_Trading_Cost", pd.Series(dtype=float)
            ),
            errors="coerce",
        ).fillna(0.0).sum()
        return pd.DataFrame(
            {
                "Metric": [
                    "Execution Mode", "Cycle", "Initial Capital",
                    "Realized PnL", "Unrealized PnL", "Current Equity",
                    "Open Positions", "Closed Trades", "Winning Trades",
                    "Losing Trades", "Equity Intraday Statutory Charges",
                    "Slippage Cost", "Total Trading Cost",
                    "Calculation Backend",
                ],
                "Value": [
                    "PAPER_ONLY", self.cycle_number, self.initial_capital,
                    float(net_series.sum()), float(unrealized), self.equity,
                    len(self.open_positions), len(self.closed_trades),
                    int((net_series > 0).sum()), int((net_series <= 0).sum()),
                    float(statutory_costs), float(slippage_costs),
                    float(total_costs),
                    self.calculation_backend,
                ],
            }
        )

    def _save_state(self, plans):
        target = Path(self.output_file).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        summary = self._summary()
        descriptor, temporary = tempfile.mkstemp(
            prefix=".residual_strategy_", suffix=".xlsx", dir=target.parent
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                summary.to_excel(writer, sheet_name="Summary", index=False)
                self.open_positions.to_excel(
                    writer, sheet_name="Open_Positions", index=False
                )
                self.closed_trades.to_excel(
                    writer, sheet_name="Closed_Trades", index=False
                )
                self.events.to_excel(writer, sheet_name="Event_Log", index=False)
                plans.to_excel(writer, sheet_name="Current_Risk_Plans", index=False)
                for worksheet in writer.book.worksheets:
                    worksheet.freeze_panes = "A2"
                    worksheet.auto_filter.ref = worksheet.dimensions
                    worksheet.sheet_view.showGridLines = False
            os.replace(temporary, target)
        except Exception:
            if os.path.exists(temporary):
                os.unlink(temporary)
            raise
        del summary
        return str(target)

    def _save_position_journal(self, plans):
        """Atomically persist restart-safe execution state outside reports."""
        if not self.position_journal_file:
            return None
        target = Path(self.position_journal_file).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        metadata = pd.DataFrame({
            "Metric": [
                "Journal Version", "Updated UTC", "Execution Mode",
                "Cycle", "Initial Capital", "Realized PnL",
                "Current Equity", "Open Position Count",
                "Closed Trade Count", "Recovery Rule",
            ],
            "Value": [
                1, pd.Timestamp.now(tz="UTC").isoformat(), "PAPER_ONLY",
                self.cycle_number, self.initial_capital, self.realized_pnl,
                self.equity, len(self.open_positions),
                len(self.closed_trades),
                "Restore existing positions; never create duplicate entries",
            ],
        })
        descriptor, temporary = tempfile.mkstemp(
            prefix=".fmrss_position_journal_",
            suffix=".xlsx",
            dir=target.parent,
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                metadata.to_excel(
                    writer, sheet_name="Recovery_Metadata", index=False
                )
                self.open_positions.to_excel(
                    writer, sheet_name="Open_Positions", index=False
                )
                plans.reindex(columns=self.PLAN_COLUMNS).to_excel(
                    writer, sheet_name="Current_Risk_Plans", index=False
                )
                self.closed_trades.to_excel(
                    writer, sheet_name="Closed_Trades", index=False
                )
                self.events.to_excel(
                    writer, sheet_name="Event_Log", index=False
                )
                for worksheet in writer.book.worksheets:
                    worksheet.freeze_panes = "A2"
                    worksheet.auto_filter.ref = worksheet.dimensions
                    worksheet.sheet_view.showGridLines = False
            os.replace(temporary, target)
        except Exception:
            if os.path.exists(temporary):
                os.unlink(temporary)
            raise
        del metadata
        return str(target)

    def _manual_open_view(self):
        if self.open_positions.empty:
            return pd.DataFrame(columns=self.MANUAL_LEDGER_COLUMNS)
        source = self.open_positions.copy()
        result = pd.DataFrame({
            "Trade_ID": source["Position_ID"],
            "Symbol": source["Symbol"],
            "Entry_Action": source["Entry_Action"],
            "Regime": source["Regime"],
            "Rank": source["Rank"],
            "Entry_Time": source["Entry_Time"],
            "Entry_Price": source["Entry_Price"],
            "Stop_Loss": source["Stop_Loss"],
            "Target_Price": source["Target_Price"],
            "Quantity": source["Quantity"],
            "Status": "OPEN",
            "Current_Price": source["Current_Price"],
            "Exit_Time": pd.NaT,
            "Exit_Price": np.nan,
            "Exit_Reason": "",
            "Gross_PnL": source["Unrealized_PnL"],
            "Total_Trading_Cost": source["Estimated_Exit_Charges"],
            "Net_PnL": source["Estimated_Net_PnL"],
            "Portfolio_Value": self.equity,
            # This is the only user-editable execution field.
            "Manual_Exit_Price": np.nan,
        })
        return result.reindex(columns=self.MANUAL_LEDGER_COLUMNS)

    def _manual_closed_view(self):
        if self.closed_trades.empty:
            return pd.DataFrame(columns=self.MANUAL_LEDGER_COLUMNS)
        source = self.closed_trades.copy()
        result = pd.DataFrame({
            "Trade_ID": source["Position_ID"],
            "Symbol": source["Symbol"],
            "Entry_Action": source["Entry_Action"],
            "Regime": source["Regime"],
            "Rank": source["Rank"],
            "Entry_Time": source["Entry_Time"],
            "Entry_Price": source["Entry_Price"],
            "Stop_Loss": source["Stop_Loss"],
            "Target_Price": source["Target_Price"],
            "Quantity": source["Quantity"],
            "Status": "CLOSED",
            "Current_Price": source["Exit_Price"],
            "Exit_Time": source["Exit_Time"],
            "Exit_Price": source["Exit_Price"],
            "Exit_Reason": source["Exit_Reason"],
            "Gross_PnL": source["Gross_PnL"],
            "Total_Trading_Cost": source["Total_Trading_Cost"],
            "Net_PnL": source["Net_PnL"],
            "Portfolio_Value": source["Equity_After"],
            "Manual_Exit_Price": np.nan,
        })
        return result.reindex(columns=self.MANUAL_LEDGER_COLUMNS)

    def _save_manual_ledger(self):
        if not self.manual_ledger_file:
            return None
        target = Path(self.manual_ledger_file).expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        open_view = self._manual_open_view()
        closed_view = self._manual_closed_view()
        ledger_parts = [
            frame for frame in [closed_view, open_view] if not frame.empty
        ]
        trading_history = (
            pd.concat(ledger_parts, ignore_index=True)
            if ledger_parts
            else pd.DataFrame(columns=self.MANUAL_LEDGER_COLUMNS)
        ).reindex(columns=self.MANUAL_LEDGER_COLUMNS)
        summary = self._summary()
        instructions = pd.DataFrame({
            "Instruction": [
                "To close a trade manually, enter the actual executed price in Manual_Exit_Price on Open_Trades.",
                "Save and close the workbook. The next Phase B poll will move the trade to Closed_Trades.",
                "Do not change Trade_ID. A manual fill is recorded with Exit_Reason Sold Manually.",
            ]
        })
        descriptor, temporary = tempfile.mkstemp(
            prefix=".trading_ledge_", suffix=".xlsx", dir=target.parent
        )
        os.close(descriptor)
        try:
            with pd.ExcelWriter(temporary, engine="openpyxl") as writer:
                summary.to_excel(writer, sheet_name="Summary", index=False)
                open_view.to_excel(writer, sheet_name="Open_Trades", index=False)
                closed_view.to_excel(
                    writer, sheet_name="Closed_Trades", index=False
                )
                trading_history.to_excel(
                    writer, sheet_name="Trading_Ledger", index=False
                )
                instructions.to_excel(
                    writer, sheet_name="Instructions", index=False
                )
                for worksheet in writer.book.worksheets:
                    worksheet.freeze_panes = "A2"
                    worksheet.auto_filter.ref = worksheet.dimensions
                    worksheet.sheet_view.showGridLines = False
                open_sheet = writer.book["Open_Trades"]
                manual_column = self.MANUAL_LEDGER_COLUMNS.index(
                    "Manual_Exit_Price"
                ) + 1
                from openpyxl.styles import PatternFill
                fill = PatternFill(
                    fill_type="solid", fgColor="FFF2CC"
                )
                for cell in open_sheet.iter_cols(
                    min_col=manual_column,
                    max_col=manual_column,
                    min_row=2,
                ):
                    for item in cell:
                        item.fill = fill
            os.replace(temporary, target)
        except Exception:
            if os.path.exists(temporary):
                os.unlink(temporary)
            raise
        del open_view, closed_view, ledger_parts
        del trading_history, summary, instructions
        return str(target)

    def run_cycle(self):
        self.cycle_number += 1
        plans = self._load_approved_plans()
        plan_history = self._load_plan_history()
        mathematical = self._load_mathematical_snapshot()
        # Read manual instructions before any workbook is refreshed.
        self._process_manual_exits()
        self._monitor_open_positions(plans, plan_history, mathematical)
        self.open_positions.reset_index(drop=True, inplace=True)
        self._open_new_positions(plans, plan_history, mathematical)

        saved = None
        journal_saved = None
        if self.save_excel:
            saved = self._save_state(plans)
            if not os.path.isfile(saved) or os.path.getsize(saved) == 0:
                raise RuntimeError("Strategy state workbook verification failed")
            manual_saved = self._save_manual_ledger()
            if manual_saved and (
                not os.path.isfile(manual_saved)
                or os.path.getsize(manual_saved) == 0
            ):
                raise RuntimeError("Manual trading ledger verification failed")
        if self.position_journal_file:
            journal_saved = self._save_position_journal(plans)
            if (
                not journal_saved
                or not os.path.isfile(journal_saved)
                or os.path.getsize(journal_saved) == 0
            ):
                raise RuntimeError("Position recovery journal verification failed")

        report = self._summary()
        report.attrs["Excel_File"] = saved
        report.attrs["Position_Journal_File"] = journal_saved
        report.attrs["Open_Positions"] = len(self.open_positions)
        report.attrs["Closed_Trades"] = len(self.closed_trades)
        if self.print_details:
            print(self.open_positions.to_string(index=False))
        if saved:
            print(f"ResidualStatisticalStrategy.xlsx file is created: {saved}")

        if self.retain_results_in_memory:
            strategy_data = {
                "open_positions": self.open_positions.copy(),
                "closed_trades": self.closed_trades.copy(),
                "events": self.events.copy(),
            }
        else:
            strategy_data = {}
        del plans, plan_history, mathematical
        return strategy_data, report

    def calculate_all(self):
        return self.run_cycle()

import gc
import os
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd


# The Cython extension is optional.  Keeping the fallback makes the engine
# usable before compilation and on machines where a compiler is unavailable.
try:
    from engines._rolling_robust import rolling_median_mad
    CYTHON_ROBUST_STATS_AVAILABLE = True
except ImportError:
    try:
        from _rolling_robust import rolling_median_mad
        CYTHON_ROBUST_STATS_AVAILABLE = True
    except ImportError:
        rolling_median_mad = None
        CYTHON_ROBUST_STATS_AVAILABLE = False


class FiveMinuteResidualEngine:

    RESIDUAL_HISTORY_COLUMNS = [
        "Symbol",
        "DateTime",
        "Open",
        "High",
        "Low",
        "Close",
        "Volume",
        "Rolling_Beta",
        "Residual_Return",
        "Residual_Impulse",
        "Residual_ZScore",
        "Residual_Features_Ready",
    ]

    # =====================================================
    # Required Input Columns
    # =====================================================

    REQUIRED_COLUMNS = [
        "Symbol",
        "DateTime",
        "Open",
        "High",
        "Low",
        "Close",
        "Volume",
        "Nifty_Open",
        "Nifty_High",
        "Nifty_Low",
        "Nifty_Close"
    ]

    # =====================================================
    # Constructor
    # =====================================================

    def __init__(
        self,
        beta_window=375,
        beta_min_periods=350,
        impulse_window=3,
        zscore_window=60,
        zscore_min_periods=50,
        beta_minimum=-1.0,
        beta_maximum=3.0,
        variance_floor=1e-12,
        mad_floor=1e-12,
        zscore_clip=5.0,
        use_cython=True,
        memory_optimized=True,
        retain_results_in_memory=False,
        report_output_file=(
            "/home/devinderjeet/fmrss/report/"
            "ResidualEngine.xlsx"
        ),
        save_excel=True,
        print_symbol_details=False,
    ):

        if beta_window < 20:
            raise ValueError(
                "beta_window must be at least 20"
            )

        if not 2 <= beta_min_periods <= beta_window:
            raise ValueError(
                "beta_min_periods must be between "
                "2 and beta_window"
            )

        if impulse_window < 1:
            raise ValueError(
                "impulse_window must be at least 1"
            )

        if zscore_window < 10:
            raise ValueError(
                "zscore_window must be at least 10"
            )

        if not 2 <= zscore_min_periods <= zscore_window:
            raise ValueError(
                "zscore_min_periods must be between "
                "2 and zscore_window"
            )

        if beta_minimum >= beta_maximum:
            raise ValueError(
                "beta_minimum must be less than "
                "beta_maximum"
            )

        if variance_floor <= 0:
            raise ValueError(
                "variance_floor must be positive"
            )

        if mad_floor <= 0:
            raise ValueError(
                "mad_floor must be positive"
            )

        if zscore_clip <= 0:
            raise ValueError(
                "zscore_clip must be positive"
            )

        self.beta_window = int(beta_window)
        self.beta_min_periods = int(
            beta_min_periods
        )

        self.impulse_window = int(
            impulse_window
        )

        self.zscore_window = int(
            zscore_window
        )

        self.zscore_min_periods = int(
            zscore_min_periods
        )

        self.beta_minimum = float(
            beta_minimum
        )

        self.beta_maximum = float(
            beta_maximum
        )

        self.variance_floor = float(
            variance_floor
        )

        self.mad_floor = float(
            mad_floor
        )

        self.zscore_clip = float(
            zscore_clip
        )

        self.use_cython = bool(use_cython)
        self.memory_optimized = bool(memory_optimized)
        self.retain_results_in_memory = bool(
            retain_results_in_memory
        )
        self.report_output_file = (
            str(report_output_file)
            if report_output_file
            else None
        )
        self.save_excel_enabled = bool(save_excel)
        self.print_symbol_details = bool(print_symbol_details)

        self.robust_stats_backend = (
            "CYTHON"
            if self.use_cython
            and CYTHON_ROBUST_STATS_AVAILABLE
            else "PYTHON"
        )

    # =====================================================
    # Validate Input
    # =====================================================

    def _validate_input(self, symbol, df):

        if not isinstance(df, pd.DataFrame):
            raise TypeError(
                f"{symbol} data must be a pandas DataFrame"
            )

        if df.empty:
            raise ValueError(
                f"{symbol} aligned data is empty"
            )

        missing_columns = [
            column
            for column in self.REQUIRED_COLUMNS
            if column not in df.columns
        ]

        if missing_columns:
            raise ValueError(
                f"{symbol} is missing required columns: "
                f"{missing_columns}"
            )

    # =====================================================
    # Prepare Input
    # =====================================================

    def _prepare_input(self, symbol, df):

        self._validate_input(
            symbol=symbol,
            df=df
        )

        # Keep only fields required here and by downstream residual consumers.
        # This avoids copying unrelated aligner diagnostics for every symbol.
        result = df.loc[:, self.REQUIRED_COLUMNS].copy()

        result["DateTime"] = pd.to_datetime(
            result["DateTime"],
            errors="coerce"
        )

        numeric_columns = [
            "Open",
            "High",
            "Low",
            "Close",
            "Volume",
            "Nifty_Open",
            "Nifty_High",
            "Nifty_Low",
            "Nifty_Close"
        ]

        for column in numeric_columns:

            result[column] = pd.to_numeric(
                result[column],
                errors="coerce"
            )

            if self.memory_optimized:
                result[column] = result[column].astype(
                    np.float32,
                    copy=False,
                )

        result.dropna(
            subset=[
                "DateTime",
                "Close",
                "Nifty_Close"
            ],
            inplace=True
        )

        valid_prices = (
            (result["Close"] > 0)
            & (result["Nifty_Close"] > 0)
        )
        if not valid_prices.all():
            result.drop(
                index=result.index[~valid_prices],
                inplace=True,
            )

        result["Symbol"] = (
            str(symbol)
            .strip()
            .upper()
        )

        result.sort_values(
            "DateTime",
            inplace=True
        )

        result.drop_duplicates(
            subset=["DateTime"],
            keep="last",
            inplace=True
        )

        result.reset_index(
            drop=True,
            inplace=True
        )

        if result.empty:
            raise ValueError(
                f"{symbol} has no valid aligned rows"
            )

        # Used to prevent overnight returns from being
        # treated as ordinary five-minute returns.
        result["Trading_Date"] = (
            result["DateTime"].dt.normalize()
        )

        return result

    # =====================================================
    # Calculate Intraday Log Returns
    # =====================================================

    @staticmethod
    def _calculate_intraday_returns(df):

        # Return is reset at the beginning of every day.
        # The first candle of each day receives NaN.
        df["Stock_Log_Close"] = np.log(
            df["Close"].to_numpy(
                dtype=np.float64,
                copy=False,
            )
        )

        df["Nifty_Log_Close"] = np.log(
            df["Nifty_Close"].to_numpy(
                dtype=np.float64,
                copy=False,
            )
        )

        df["Stock_Return"] = (
            df.groupby(
                "Trading_Date",
                sort=False
            )["Stock_Log_Close"]
            .diff()
        )

        df["Nifty_Return"] = (
            df.groupby(
                "Trading_Date",
                sort=False
            )["Nifty_Log_Close"]
            .diff()
        )

        return df

    # =====================================================
    # Calculate Rolling Beta
    # =====================================================

    def _calculate_rolling_beta(self, df):

        stock_return = df["Stock_Return"]
        nifty_return = df["Nifty_Return"]

        # Rolling NIFTY variance
        rolling_nifty_variance = (
            nifty_return
            .rolling(
                window=self.beta_window,
                min_periods=self.beta_min_periods
            )
            .var(ddof=1)
        )

        # Rolling stock/NIFTY covariance
        rolling_covariance = (
            stock_return
            .rolling(
                window=self.beta_window,
                min_periods=self.beta_min_periods
            )
            .cov(nifty_return)
        )

        # Shift by one candle:
        # beta used for candle t is estimated only through
        # candle t-1.
        df["Nifty_Return_Variance"] = (
            rolling_nifty_variance.shift(1)
        )

        df["Stock_Nifty_Covariance"] = (
            rolling_covariance.shift(1)
        )

        valid_variance = (
            df["Nifty_Return_Variance"]
            > self.variance_floor
        )

        df["Rolling_Beta_Raw"] = np.where(
            valid_variance,
            (
                df["Stock_Nifty_Covariance"]
                / df["Nifty_Return_Variance"]
            ),
            np.nan
        )

        df["Beta_Was_Clipped"] = (
            (
                df["Rolling_Beta_Raw"]
                < self.beta_minimum
            )
            |
            (
                df["Rolling_Beta_Raw"]
                > self.beta_maximum
            )
        )

        df["Rolling_Beta"] = (
            df["Rolling_Beta_Raw"]
            .clip(
                lower=self.beta_minimum,
                upper=self.beta_maximum
            )
        )

        df["Beta_Valid"] = (
            df["Rolling_Beta"].notna()
            & valid_variance
        )

        return df

    # =====================================================
    # Calculate Rolling Correlation
    # =====================================================

    def _calculate_market_relationship(self, df):

        rolling_correlation = (
            df["Stock_Return"]
            .rolling(
                window=self.beta_window,
                min_periods=self.beta_min_periods
            )
            .corr(df["Nifty_Return"])
        )

        # Exclude the current candle
        df["Stock_Nifty_Correlation"] = (
            rolling_correlation.shift(1)
        )

        df["Stock_Nifty_RSquared"] = (
            df["Stock_Nifty_Correlation"] ** 2
        )

        return df

    # =====================================================
    # Calculate Residual Return
    # =====================================================

    @staticmethod
    def _calculate_residual_return(df):

        # Expected stock return explained by NIFTY
        df["Market_Expected_Return"] = (
            df["Rolling_Beta"]
            * df["Nifty_Return"]
        )

        # Stock-specific unexplained movement
        df["Residual_Return"] = (
            df["Stock_Return"]
            - df["Market_Expected_Return"]
        )

        return df

    # =====================================================
    # Calculate Residual Impulse
    # =====================================================

    def _calculate_residual_impulse(self, df):

        # The impulse must not cross an overnight boundary.
        df["Residual_Impulse"] = (
            df.groupby(
                "Trading_Date",
                sort=False
            )["Residual_Return"]
            .transform(
                lambda series:
                series.rolling(
                    window=self.impulse_window,
                    min_periods=self.impulse_window
                ).sum()
            )
        )

        return df

    # =====================================================
    # Median Absolute Deviation
    # =====================================================

    @staticmethod
    def _median_absolute_deviation(values):

        values = np.asarray(
            values,
            dtype=np.float64
        )

        values = values[
            np.isfinite(values)
        ]

        if values.size == 0:
            return np.nan

        median = np.median(values)

        return np.median(
            np.abs(values - median)
        )

    # =====================================================
    # Calculate Robust Residual Z-Score
    # =====================================================

    def _calculate_robust_zscore(self, df):

        impulse = df["Residual_Impulse"]

        if self.robust_stats_backend == "CYTHON":

            # The compiled function directly produces statistics for the
            # previous window.  This is equivalent to pandas rolling(...)
            # followed by shift(1), but avoids a Python callback per row.
            # Guarantee a C-contiguous, aligned and writable buffer.  This is
            # compatible with the new const-memoryview kernel as well as an
            # older compiled build that requests a writable source buffer.
            impulse_values = np.require(
                impulse.to_numpy(dtype=np.float64, copy=False),
                dtype=np.float64,
                requirements=["C", "A", "W"],
            )

            median_values, mad_values = rolling_median_mad(
                impulse_values,
                self.zscore_window,
                self.zscore_min_periods
            )

            rolling_median = pd.Series(
                median_values,
                index=df.index,
                dtype=np.float64
            )

            rolling_mad = pd.Series(
                mad_values,
                index=df.index,
                dtype=np.float64
            )

        else:

            # Safe fallback used when the extension has not been compiled.
            rolling_median = (
                impulse
                .rolling(
                    window=self.zscore_window,
                    min_periods=self.zscore_min_periods
                )
                .median()
                .shift(1)
            )

            rolling_mad = (
                impulse
                .rolling(
                    window=self.zscore_window,
                    min_periods=self.zscore_min_periods
                )
                .apply(
                    self._median_absolute_deviation,
                    raw=True
                )
                .shift(1)
            )

        # 1.4826 makes MAD comparable with standard
        # deviation under an approximately normal
        # distribution.
        robust_scale = 1.4826 * rolling_mad

        df["Residual_Rolling_Median"] = (
            rolling_median
        )

        df["Residual_Rolling_MAD"] = (
            rolling_mad
        )

        df["Residual_Robust_Scale"] = (
            robust_scale
        )

        valid_scale = (
            robust_scale > self.mad_floor
        )

        df["Residual_ZScore"] = np.where(
            valid_scale,
            (
                (
                    df["Residual_Impulse"]
                    - rolling_median
                )
                / robust_scale
            ),
            np.nan
        )

        # Keep the real Z-score and a clipped copy.
        # The clipped copy will later be used for scoring.
        df["Residual_ZScore_Clipped"] = (
            df["Residual_ZScore"]
            .clip(
                lower=-self.zscore_clip,
                upper=self.zscore_clip
            )
        )

        return df

    # =====================================================
    # Calculate Feature State
    # =====================================================

    @staticmethod
    def _calculate_feature_state(df):

        required_features = [
            "Stock_Return",
            "Nifty_Return",
            "Rolling_Beta",
            "Residual_Return",
            "Residual_Impulse",
            "Residual_ZScore"
        ]

        df["Residual_Features_Ready"] = (
            df[required_features]
            .notna()
            .all(axis=1)
        )

        df["Residual_Direction"] = np.select(
            condlist=[
                df["Residual_Return"] > 0,
                df["Residual_Return"] < 0
            ],
            choicelist=[
                "POSITIVE",
                "NEGATIVE"
            ],
            default="NEUTRAL"
        )

        df["Residual_Extreme_State"] = np.select(
            condlist=[
                df["Residual_ZScore"] >= 2.0,
                df["Residual_ZScore"] <= -2.0,
                df["Residual_ZScore"].abs() >= 1.0
            ],
            choicelist=[
                "POSITIVE_EXTREME",
                "NEGATIVE_EXTREME",
                "ELEVATED"
            ],
            default="NORMAL"
        )

        # Rows without sufficient warm-up history should
        # not be described as normal.
        df.loc[
            ~df["Residual_Features_Ready"],
            "Residual_Extreme_State"
        ] = "NOT_READY"

        return df

    # =====================================================
    # Calculate Features for One Symbol
    # =====================================================

    def calculate_symbol(
        self,
        symbol,
        aligned_df
    ):

        symbol = (
            str(symbol)
            .strip()
            .upper()
        )

        df = self._prepare_input(
            symbol=symbol,
            df=aligned_df
        )

        df = self._calculate_intraday_returns(df)

        df = self._calculate_rolling_beta(df)

        df = self._calculate_market_relationship(df)

        df = self._calculate_residual_return(df)

        df = self._calculate_residual_impulse(df)

        df = self._calculate_robust_zscore(df)

        df = self._calculate_feature_state(df)

        # Helper logarithm columns are unnecessary in the
        # final output.
        df.drop(
            columns=[
                "Stock_Log_Close",
                "Nifty_Log_Close"
            ],
            inplace=True,
            errors="ignore"
        )

        ready_rows = int(
            df["Residual_Features_Ready"].sum()
        )

        report = {
            "Symbol": symbol,
            "Status": (
                "CALCULATED"
                if ready_rows > 0
                else "NOT_READY"
            ),
            "Input_Rows": len(df),
            "Feature_Ready_Rows": ready_rows,
            "Latest_DateTime": (
                df["DateTime"].iloc[-1]
            ),
            "Latest_Beta": (
                df["Rolling_Beta"].iloc[-1]
            ),
            "Latest_Residual_Return": (
                df["Residual_Return"].iloc[-1]
            ),
            "Latest_Residual_Impulse": (
                df["Residual_Impulse"].iloc[-1]
            ),
            "Latest_Residual_ZScore": (
                df["Residual_ZScore"].iloc[-1]
            ),
            "Latest_Features_Ready": bool(
                df[
                    "Residual_Features_Ready"
                ].iloc[-1]
            ),
            "Robust_Stats_Backend": (
                self.robust_stats_backend
            ),
            "Error": "",
        }

        return df, report

    # =====================================================
    # Excel Reporting
    # =====================================================

    @staticmethod
    def _excel_safe(frame):
        """Return an Excel-safe shallow copy of a report DataFrame."""
        safe = frame.copy(deep=False)
        safe = safe.replace([np.inf, -np.inf], np.nan)
        for column in safe.columns:
            if isinstance(safe[column].dtype, pd.CategoricalDtype):
                safe[column] = safe[column].astype("string")
            if isinstance(safe[column].dtype, pd.DatetimeTZDtype):
                safe[column] = safe[column].dt.tz_localize(None)
        return safe

    def _build_excel_summary(self, report_df):
        status = report_df["Status"].astype("string")
        latest_ready = report_df["Latest_Features_Ready"].fillna(False)
        return pd.DataFrame(
            [
                ["RUN", "Report generated", pd.Timestamp.now()],
                ["RUN", "Robust statistics backend", self.robust_stats_backend],
                ["RUN", "Memory optimized", self.memory_optimized],
                [
                    "RUN",
                    "Results retained in memory",
                    self.retain_results_in_memory,
                ],
                ["COUNTS", "Total symbols", int(len(report_df))],
                ["COUNTS", "Calculated symbols", int(status.eq("CALCULATED").sum())],
                ["COUNTS", "Not-ready symbols", int(status.eq("NOT_READY").sum())],
                ["COUNTS", "Failed symbols", int(status.eq("FAILED").sum())],
                ["COUNTS", "Latest feature-ready symbols", int(latest_ready.sum())],
                ["CONFIGURATION", "Beta window", self.beta_window],
                ["CONFIGURATION", "Beta minimum periods", self.beta_min_periods],
                ["CONFIGURATION", "Impulse window", self.impulse_window],
                ["CONFIGURATION", "Z-score window", self.zscore_window],
                ["CONFIGURATION", "Z-score minimum periods", self.zscore_min_periods],
                ["CONFIGURATION", "Beta minimum", self.beta_minimum],
                ["CONFIGURATION", "Beta maximum", self.beta_maximum],
                ["CONFIGURATION", "Z-score clip", self.zscore_clip],
            ],
            columns=["Section", "Metric", "Value"],
        )

    def _write_residual_history(self, writer, residual_data):
        """Write one combined history sheet in bounded per-symbol chunks."""
        maximum_data_rows = 1_048_575  # Excel limit minus header row.
        total_rows = sum(
            len(frame)
            for frame in residual_data.values()
            if isinstance(frame, pd.DataFrame) and not frame.empty
        )
        if total_rows > maximum_data_rows:
            raise ValueError(
                "Residual_History exceeds the Excel worksheet limit: "
                f"{total_rows:,} rows > {maximum_data_rows:,} rows"
            )

        if total_rows == 0:
            pd.DataFrame(columns=self.RESIDUAL_HISTORY_COLUMNS).to_excel(
                writer,
                sheet_name="Residual_History",
                index=False,
            )
            return 0

        start_row = 0
        write_header = True
        written_rows = 0

        for symbol, frame in residual_data.items():
            if not isinstance(frame, pd.DataFrame) or frame.empty:
                continue

            missing = [
                column
                for column in self.RESIDUAL_HISTORY_COLUMNS
                if column not in frame.columns
            ]
            if missing:
                raise ValueError(
                    f"{symbol} residual history is missing columns: {missing}"
                )

            # Only one symbol (normally at most 1,500 rows) is staged at a
            # time. No second 120,000-row combined DataFrame is constructed.
            history_chunk = frame.loc[
                :,
                self.RESIDUAL_HISTORY_COLUMNS,
            ].copy(deep=False)
            history_chunk = self._excel_safe(history_chunk)
            history_chunk.to_excel(
                writer,
                sheet_name="Residual_History",
                index=False,
                header=write_header,
                startrow=start_row,
            )

            chunk_rows = len(history_chunk)
            written_rows += chunk_rows
            start_row += chunk_rows + (1 if write_header else 0)
            write_header = False
            del history_chunk

        return written_rows

    @staticmethod
    def _style_workbook(writer):
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(color="FFFFFF", bold=True)
        green_fill = PatternFill("solid", fgColor="C6EFCE")
        red_fill = PatternFill("solid", fgColor="FFC7CE")
        amber_fill = PatternFill("solid", fgColor="FFEB9C")

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            # Avoid generating four 120,000-cell tuples merely to estimate
            # widths on the long history worksheet.
            if worksheet.title == "Residual_History":
                from openpyxl.utils import get_column_letter

                for index in range(1, worksheet.max_column + 1):
                    worksheet.column_dimensions[
                        get_column_letter(index)
                    ].width = 21
                continue

            for column_cells in worksheet.columns:
                values = [
                    "" if cell.value is None else str(cell.value)
                    for cell in column_cells[:200]
                ]
                width = min(
                    max(max(map(len, values), default=0) + 2, 11),
                    34,
                )
                worksheet.column_dimensions[
                    column_cells[0].column_letter
                ].width = width

            headers = {
                cell.value: cell.column_letter
                for cell in worksheet[1]
                if cell.value is not None
            }
            for header in ["Latest_DateTime"]:
                column = headers.get(header)
                if column:
                    for row in range(2, worksheet.max_row + 1):
                        worksheet[f"{column}{row}"].number_format = (
                            "yyyy-mm-dd hh:mm:ss"
                        )

            # Empty DataFrames produce a header-only worksheet.  Do not form
            # invalid conditional ranges such as G2:G1.
            if worksheet.max_row < 2:
                continue

            status_column = headers.get("Status")
            if status_column:
                target = (
                    f"{status_column}2:"
                    f"{status_column}{worksheet.max_row}"
                )
                worksheet.conditional_formatting.add(
                    target,
                    FormulaRule(
                        formula=[f'{status_column}2="CALCULATED"'],
                        fill=green_fill,
                    ),
                )
                worksheet.conditional_formatting.add(
                    target,
                    FormulaRule(
                        formula=[f'{status_column}2="FAILED"'],
                        fill=red_fill,
                    ),
                )
                worksheet.conditional_formatting.add(
                    target,
                    FormulaRule(
                        formula=[f'{status_column}2="NOT_READY"'],
                        fill=amber_fill,
                    ),
                )

    def _save_excel_report(self, report_df, residual_data, output_file):
        """Atomically create or replace the residual report workbook."""
        output_path = Path(output_file).expanduser().resolve()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        calculated = report_df.loc[
            report_df["Status"].eq("CALCULATED")
        ].copy()
        exceptions = report_df.loc[
            ~report_df["Status"].eq("CALCULATED")
        ].copy()
        summary = self._build_excel_summary(report_df)

        descriptor, temporary_name = tempfile.mkstemp(
            prefix="ResidualEngine_",
            suffix=".xlsx",
            dir=output_path.parent,
        )
        os.close(descriptor)

        try:
            with pd.ExcelWriter(
                temporary_name,
                engine="openpyxl",
                datetime_format="yyyy-mm-dd hh:mm:ss",
            ) as writer:
                self._excel_safe(summary).to_excel(
                    writer, sheet_name="Summary", index=False
                )
                self._excel_safe(calculated).to_excel(
                    writer, sheet_name="Calculated_Shares", index=False
                )
                self._excel_safe(exceptions).to_excel(
                    writer, sheet_name="Not_Ready_Failed", index=False
                )
                self._excel_safe(report_df).to_excel(
                    writer, sheet_name="All_Shares", index=False
                )
                history_rows = self._write_residual_history(
                    writer=writer,
                    residual_data=residual_data,
                )
                self._style_workbook(writer)

            # os.replace overwrites an existing workbook atomically, so a
            # reader never observes a half-written Excel file.
            os.replace(temporary_name, output_path)
        except Exception:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)
            raise

        report_df.attrs["Excel_File"] = str(output_path)
        report_df.attrs["Residual_History_Rows"] = history_rows
        return str(output_path)

    # =====================================================
    # Calculate Features for All Symbols
    # =====================================================

    def calculate_all(self, aligned_data, output_file=None):

        if not isinstance(aligned_data, dict):
            raise TypeError(
                "aligned_data must be a dictionary "
                "containing {symbol: DataFrame}"
            )

        if not aligned_data:
            raise ValueError(
                "aligned_data is empty"
            )

        residual_data = {}
        reports = []

        for symbol, aligned_df in aligned_data.items():

            symbol = (
                str(symbol)
                .strip()
                .upper()
            )

            try:

                result_df, report = (
                    self.calculate_symbol(
                        symbol=symbol,
                        aligned_df=aligned_df
                    )
                )

                reports.append(report)

                if report["Status"] == "NOT_READY":

                    if self.print_symbol_details:
                        print(
                            f"{symbol:<15}"
                            f"NOT READY : "
                            f"No fully calculated feature rows"
                        )

                    continue

                residual_data[symbol] = result_df

                latest_beta = (
                    report["Latest_Beta"]
                )

                latest_zscore = (
                    report[
                        "Latest_Residual_ZScore"
                    ]
                )

                beta_text = (
                    f"{latest_beta:.4f}"
                    if pd.notna(latest_beta)
                    else "NaN"
                )

                zscore_text = (
                    f"{latest_zscore:.4f}"
                    if pd.notna(latest_zscore)
                    else "NaN"
                )

                if self.print_symbol_details:
                    print(
                        f"{symbol:<15}"
                        f"CALCULATED : "
                        f"Ready="
                        f"{report['Feature_Ready_Rows']} | "
                        f"Latest Beta={beta_text} | "
                        f"Latest Z={zscore_text}"
                    )

            except Exception as error:

                reports.append({
                    "Symbol": symbol,
                    "Status": "FAILED",
                    "Input_Rows": (
                        len(aligned_df)
                        if isinstance(
                            aligned_df,
                            pd.DataFrame
                        )
                        else 0
                    ),
                    "Feature_Ready_Rows": 0,
                    "Latest_DateTime": pd.NaT,
                    "Latest_Beta": np.nan,
                    "Latest_Residual_Return": np.nan,
                    "Latest_Residual_Impulse": np.nan,
                    "Latest_Residual_ZScore": np.nan,
                    "Latest_Features_Ready": False,
                    "Robust_Stats_Backend": (
                        self.robust_stats_backend
                    ),
                    "Error": str(error)
                })

                if self.print_symbol_details:
                    print(
                        f"{symbol:<15}"
                        f"FAILED : {error}"
                    )

        report_df = pd.DataFrame(reports)

        if not report_df.empty:

            report_df.sort_values(
                by=[
                    "Status",
                    "Symbol"
                ],
                inplace=True
            )

            report_df.reset_index(
                drop=True,
                inplace=True
            )

        selected_output_file = (
            output_file
            if output_file is not None
            else self.report_output_file
        )
        saved_file = None
        if self.save_excel_enabled and selected_output_file:
            saved_file = self._save_excel_report(
                report_df=report_df,
                residual_data=residual_data,
                output_file=selected_output_file,
            )
            print("ResidualEngine.xlsx file is created")

        report_df.attrs["Excel_File"] = saved_file
        report_df.attrs["Robust_Stats_Backend"] = (
            self.robust_stats_backend
        )

        if not residual_data:
            raise RuntimeError(
                "Residual features could not be calculated for any symbol. "
                f"Diagnostic workbook: {saved_file}"
            )

        report_df.attrs["Results_Retained_In_Memory"] = (
            self.retain_results_in_memory
        )

        if not self.retain_results_in_memory:
            residual_data.clear()
            gc.collect()

        return residual_data, report_df

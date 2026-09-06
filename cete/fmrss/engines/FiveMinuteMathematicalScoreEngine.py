import os
import tempfile
from pathlib import Path

import numpy as np
import pandas as pd


class FiveMinuteMathematicalScoreEngine:

    # =====================================================
    # Required Input Columns
    # =====================================================

    REQUIRED_COLUMNS = [
        "Symbol",
        "DateTime",
        "Residual_Return",
        "Residual_Impulse",
        "Residual_ZScore",
        "Residual_Features_Ready",
        "Stable_Regime",
        "Regime_Stable",
        "Regime_Strength",
        "Volume_Factor",
        "Volume_Confirmation_Ready",
        "Volume_Confirmed",
        "Volatility_Features_Ready",
        "Volatility_Eligible"
    ]

    DIRECTIONAL_REGIMES = {
        "MOMENTUM",
        "MEAN_REVERSION"
    }

    # =====================================================
    # Constructor
    # =====================================================

    def __init__(
        self,
        residual_weight=0.35,
        regime_weight=0.25,
        volume_weight=0.20,
        consistency_weight=0.20,
        residual_zscore_reference=3.0,
        volume_factor_reference=2.0,
        momentum_consistency_window=5,
        reversion_consistency_window=3,
        minimum_mathematical_score=65.0,
        weight_tolerance=1e-9,
        memory_optimized=True,
        report_output_file=(
            "/home/devinderjeet/fmrss/report/"
            "MathematicalScoreEngine.xlsx"
        ),
        save_excel=True,
        print_symbol_details=True
    ):

        weights = np.asarray(
            [
                residual_weight,
                regime_weight,
                volume_weight,
                consistency_weight
            ],
            dtype=np.float64
        )

        if not np.isfinite(weights).all():
            raise ValueError(
                "All score weights must be finite"
            )

        if (weights < 0).any():
            raise ValueError(
                "Score weights cannot be negative"
            )

        if not np.isclose(
            weights.sum(),
            1.0,
            rtol=0.0,
            atol=weight_tolerance
        ):
            raise ValueError(
                "Score weights must sum to 1.0"
            )

        if residual_zscore_reference <= 0:
            raise ValueError(
                "residual_zscore_reference must be positive"
            )

        if volume_factor_reference <= 0:
            raise ValueError(
                "volume_factor_reference must be positive"
            )

        if momentum_consistency_window < 2:
            raise ValueError(
                "momentum_consistency_window must be at least 2"
            )

        if reversion_consistency_window < 1:
            raise ValueError(
                "reversion_consistency_window must be at least 1"
            )

        if not 0 <= minimum_mathematical_score <= 100:
            raise ValueError(
                "minimum_mathematical_score must be between 0 and 100"
            )

        self.residual_weight = float(
            residual_weight
        )

        self.regime_weight = float(
            regime_weight
        )

        self.volume_weight = float(
            volume_weight
        )

        self.consistency_weight = float(
            consistency_weight
        )

        self.residual_zscore_reference = float(
            residual_zscore_reference
        )

        self.volume_factor_reference = float(
            volume_factor_reference
        )

        self.momentum_consistency_window = int(
            momentum_consistency_window
        )

        self.reversion_consistency_window = int(
            reversion_consistency_window
        )

        self.minimum_mathematical_score = float(
            minimum_mathematical_score
        )

        self.memory_optimized = bool(memory_optimized)
        self.report_output_file = (
            str(report_output_file)
            if report_output_file is not None
            else None
        )
        self.save_excel_enabled = bool(save_excel)
        self.print_symbol_details = bool(print_symbol_details)

        # Rolling mean and arithmetic operations already execute through
        # compiled pandas/NumPy paths. There is no Python rolling callback.
        self.calculation_backend = (
            "NUMPY_PANDAS_MEMORY_OPTIMIZED"
            if self.memory_optimized
            else "NUMPY_PANDAS"
        )

    # =====================================================
    # Validate Input
    # =====================================================

    def _validate_input(self, symbol, df):

        if not isinstance(df, pd.DataFrame):
            raise TypeError(
                f"{symbol} data must be a pandas DataFrame"
            )

        if df.empty:
            raise ValueError(
                f"{symbol} volatility data is empty"
            )

        missing_columns = [
            column
            for column in self.REQUIRED_COLUMNS
            if column not in df.columns
        ]

        if missing_columns:
            raise ValueError(
                f"{symbol} is missing required columns: "
                f"{missing_columns}"
            )

    # =====================================================
    # Prepare Input
    # =====================================================

    def _prepare_input(self, symbol, df):

        self._validate_input(
            symbol=symbol,
            df=df
        )

        # In memory-optimized mode, keep the immutable input blocks shared.
        # All cleaned/recalculated columns below are assigned as complete
        # columns, so pandas allocates new blocks only for changed data.
        result = df.copy(
            deep=not self.memory_optimized
        )

        result["DateTime"] = pd.to_datetime(
            result["DateTime"],
            errors="coerce"
        )

        numeric_columns = [
            "Residual_Return",
            "Residual_Impulse",
            "Residual_ZScore",
            "Regime_Strength",
            "Volume_Factor"
        ]

        for column in numeric_columns:
            result[column] = pd.to_numeric(
                result[column],
                errors="coerce"
            )

        result[numeric_columns] = (
            result[numeric_columns]
            .replace(
                [np.inf, -np.inf],
                np.nan
            )
        )

        boolean_columns = [
            "Residual_Features_Ready",
            "Regime_Stable",
            "Volume_Confirmation_Ready",
            "Volume_Confirmed",
            "Volatility_Features_Ready",
            "Volatility_Eligible"
        ]

        for column in boolean_columns:
            result[column] = (
                result[column]
                .fillna(False)
                .astype(bool)
            )

        result["Stable_Regime"] = (
            result["Stable_Regime"]
            .astype("string")
            .fillna("NOT_READY")
            .str.strip()
            .str.upper()
            .astype("category")
        )

        invalid_timestamp = result["DateTime"].isna()

        if invalid_timestamp.any():
            result.drop(
                index=result.index[invalid_timestamp],
                inplace=True
            )

        result["Symbol"] = (
            str(symbol)
            .strip()
            .upper()
        )

        result.sort_values(
            "DateTime",
            inplace=True
        )

        result.drop_duplicates(
            subset=["DateTime"],
            keep="last",
            inplace=True
        )

        result.reset_index(
            drop=True,
            inplace=True
        )

        if self.memory_optimized:
            # Category codes are compact and all flags require one byte per
            # row. Float64 is retained for the source residual calculations;
            # only bounded reporting/score features are stored as float32.
            for column in boolean_columns:
                result[column] = result[column].astype(
                    np.bool_,
                    copy=False
                )

        if result.empty:
            raise ValueError(
                f"{symbol} has no valid timestamped rows"
            )

        return result

    # =====================================================
    # Component Scores
    # =====================================================

    def _calculate_component_scores(self, df):

        residual_score = (
            100.0
            * df["Residual_ZScore"].abs()
            / self.residual_zscore_reference
        ).clip(
            lower=0.0,
            upper=100.0
        )

        regime_score = (
            df["Regime_Strength"]
            .clip(
                lower=0.0,
                upper=100.0
            )
        )

        volume_score = (
            100.0
            * df["Volume_Factor"]
            / self.volume_factor_reference
        ).clip(
            lower=0.0,
            upper=100.0
        )

        score_dtype = np.float32 if self.memory_optimized else np.float64

        df["Residual_Score"] = residual_score.astype(
            score_dtype,
            copy=False
        )
        df["Regime_Score"] = regime_score.astype(
            score_dtype,
            copy=False
        )
        df["Volume_Score"] = volume_score.astype(
            score_dtype,
            copy=False
        )

        return df

    # =====================================================
    # Momentum Directional Consistency
    # =====================================================

    def _calculate_momentum_consistency(self, df):

        residual = df["Residual_Return"]

        valid_residual = residual.notna()
        indicator_dtype = (
            np.float32 if self.memory_optimized else np.float64
        )

        positive_indicator = (
            residual.gt(0)
            .astype(indicator_dtype)
            .where(valid_residual)
        )

        negative_indicator = (
            residual.lt(0)
            .astype(indicator_dtype)
            .where(valid_residual)
        )

        positive_consistency = (
            positive_indicator
            .rolling(
                window=self.momentum_consistency_window,
                min_periods=self.momentum_consistency_window
            )
            .mean()
        )

        negative_consistency = (
            negative_indicator
            .rolling(
                window=self.momentum_consistency_window,
                min_periods=self.momentum_consistency_window
            )
            .mean()
        )

        impulse = df["Residual_Impulse"]

        direction = np.select(
            condlist=[
                impulse > 0,
                impulse < 0
            ],
            choicelist=[
                "POSITIVE",
                "NEGATIVE"
            ],
            default="NEUTRAL"
        )

        momentum_consistency = np.select(
            condlist=[
                impulse > 0,
                impulse < 0
            ],
            choicelist=[
                positive_consistency,
                negative_consistency
            ],
            default=np.nan
        )

        df["Momentum_Direction"] = pd.Categorical(
            direction,
            categories=[
                "NEGATIVE",
                "NEUTRAL",
                "POSITIVE"
            ]
        )

        df["Momentum_Positive_Consistency"] = (
            positive_consistency.astype(
                indicator_dtype,
                copy=False
            )
        )

        df["Momentum_Negative_Consistency"] = (
            negative_consistency.astype(
                indicator_dtype,
                copy=False
            )
        )

        df["Momentum_Consistency"] = (
            pd.Series(
                momentum_consistency,
                index=df.index,
                dtype=indicator_dtype
            )
        )

        return df

    # =====================================================
    # Mean-Reversion Consistency
    # =====================================================

    def _calculate_reversion_consistency(self, df):

        absolute_zscore = df["Residual_ZScore"].abs()

        movement_toward_zero = (
            absolute_zscore.diff()
        )

        indicator_dtype = (
            np.float32 if self.memory_optimized else np.float64
        )

        improvement_indicator = (
            movement_toward_zero.lt(0)
            .astype(indicator_dtype)
            .where(movement_toward_zero.notna())
        )

        reversion_consistency = (
            improvement_indicator
            .rolling(
                window=self.reversion_consistency_window,
                min_periods=self.reversion_consistency_window
            )
            .mean()
        )

        df["Reversion_Consistency"] = (
            reversion_consistency.astype(
                indicator_dtype,
                copy=False
            )
        )

        return df

    # =====================================================
    # Select Regime-Specific Consistency
    # =====================================================

    def _select_directional_consistency(self, df):

        regime = df["Stable_Regime"]

        consistency = np.select(
            condlist=[
                regime.eq("MOMENTUM"),
                regime.eq("MEAN_REVERSION")
            ],
            choicelist=[
                df["Momentum_Consistency"],
                df["Reversion_Consistency"]
            ],
            default=np.nan
        )

        consistency = pd.Series(
            consistency,
            index=df.index,
            dtype=(
                np.float32
                if self.memory_optimized
                else np.float64
            )
        ).clip(
            lower=0.0,
            upper=1.0
        )

        df["Directional_Consistency"] = consistency

        df["Consistency_Score"] = (
            100.0 * consistency
        )

        df["Consistency_Features_Ready"] = (
            consistency.notna()
            & df["Regime_Stable"]
            & regime.isin(
                [
                    "MOMENTUM",
                    "MEAN_REVERSION"
                ]
            )
        )

        return df

    # =====================================================
    # Final Mathematical Score
    # =====================================================

    def _calculate_final_score(self, df):

        component_columns = [
            "Residual_Score",
            "Regime_Score",
            "Volume_Score",
            "Consistency_Score"
        ]

        score_features_ready = (
            df[component_columns]
            .notna()
            .all(axis=1)
        )

        mathematical_score = (
            self.residual_weight
            * df["Residual_Score"]
            + self.regime_weight
            * df["Regime_Score"]
            + self.volume_weight
            * df["Volume_Score"]
            + self.consistency_weight
            * df["Consistency_Score"]
        )

        df["Score_Features_Ready"] = score_features_ready

        df["Mathematical_Score"] = (
            mathematical_score
            .where(score_features_ready, np.nan)
            .clip(lower=0.0, upper=100.0)
            .astype(
                np.float32
                if self.memory_optimized
                else np.float64,
                copy=False
            )
        )

        return df

    # =====================================================
    # Setup Gates and Rejection Reason
    # =====================================================

    def _calculate_setup_state(self, df):

        directional_regime = df["Stable_Regime"].isin(
            [
                "MOMENTUM",
                "MEAN_REVERSION"
            ]
        )

        setup_ready = (
            df["Residual_Features_Ready"]
            & df["Regime_Stable"]
            & directional_regime
            & df["Volume_Confirmation_Ready"]
            & df["Volume_Confirmed"]
            & df["Volatility_Features_Ready"]
            & df["Volatility_Eligible"]
            & df["Consistency_Features_Ready"]
            & df["Score_Features_Ready"]
        )

        score_accepted = (
            df["Mathematical_Score"]
            >= self.minimum_mathematical_score
        )

        df["Mathematical_Setup_Ready"] = setup_ready

        df["Mathematical_Setup_Accepted"] = (
            setup_ready
            & score_accepted
        )

        reason = np.select(
            condlist=[
                ~df["Residual_Features_Ready"],
                df["Stable_Regime"].eq("RANDOM"),
                ~df["Regime_Stable"],
                ~directional_regime,
                ~df["Volume_Confirmation_Ready"],
                ~df["Volume_Confirmed"],
                ~df["Volatility_Features_Ready"],
                ~df["Volatility_Eligible"],
                ~df["Consistency_Features_Ready"],
                ~df["Score_Features_Ready"],
                ~score_accepted
            ],
            choicelist=[
                "RESIDUAL_NOT_READY",
                "RANDOM_OR_INVALID_REGIME",
                "REGIME_UNSTABLE",
                "RANDOM_OR_INVALID_REGIME",
                "VOLUME_NOT_READY",
                "VOLUME_NOT_CONFIRMED",
                "VOLATILITY_NOT_READY",
                "VOLATILITY_NOT_ELIGIBLE",
                "CONSISTENCY_NOT_READY",
                "SCORE_NOT_READY",
                "LOW_MATHEMATICAL_SCORE"
            ],
            default="ACCEPTED"
        )

        df["Score_Rejection_Reason"] = pd.Categorical(
            reason
        )

        return df

    # =====================================================
    # Calculate One Symbol
    # =====================================================

    def calculate_symbol(
        self,
        symbol,
        volatility_df
    ):

        symbol = (
            str(symbol)
            .strip()
            .upper()
        )

        df = self._prepare_input(
            symbol=symbol,
            df=volatility_df
        )

        df = self._calculate_component_scores(df)

        df = self._calculate_momentum_consistency(df)

        df = self._calculate_reversion_consistency(df)

        df = self._select_directional_consistency(df)

        df = self._calculate_final_score(df)

        df = self._calculate_setup_state(df)

        score_ready_rows = int(
            df["Score_Features_Ready"].sum()
        )

        accepted_rows = int(
            df["Mathematical_Setup_Accepted"].sum()
        )

        latest = df.iloc[-1]

        report = {
            "Symbol": symbol,
            "Status": (
                "CALCULATED"
                if score_ready_rows > 0
                else "NOT_READY"
            ),
            "Input_Rows": len(df),
            "Score_Ready_Rows": score_ready_rows,
            "Accepted_Rows": accepted_rows,
            "Latest_DateTime": latest["DateTime"],
            "Latest_Stable_Regime": str(
                latest["Stable_Regime"]
            ),
            "Latest_Residual_Return": latest["Residual_Return"],
            "Latest_Residual_Impulse": latest["Residual_Impulse"],
            "Latest_Residual_ZScore": latest["Residual_ZScore"],
            "Latest_Regime_Strength": latest["Regime_Strength"],
            "Latest_Volume_Factor": latest["Volume_Factor"],
            "Latest_Momentum_Direction": str(
                latest["Momentum_Direction"]
            ),
            "Latest_Directional_Consistency": (
                latest["Directional_Consistency"]
            ),
            "Latest_Residual_Score": latest["Residual_Score"],
            "Latest_Regime_Score": latest["Regime_Score"],
            "Latest_Volume_Score": latest["Volume_Score"],
            "Latest_Consistency_Score": latest["Consistency_Score"],
            "Latest_Mathematical_Score": (
                latest["Mathematical_Score"]
            ),
            "Latest_Setup_Ready": bool(
                latest["Mathematical_Setup_Ready"]
            ),
            "Latest_Setup_Accepted": bool(
                latest["Mathematical_Setup_Accepted"]
            ),
            "Latest_Volatility_Eligible": bool(
                latest["Volatility_Eligible"]
            ),
            "Latest_Volume_Confirmed": bool(
                latest["Volume_Confirmed"]
            ),
            "Latest_Rejection_Reason": str(
                latest["Score_Rejection_Reason"]
            ),
            "Calculation_Backend": self.calculation_backend
        }

        return df, report

    # =====================================================
    # Excel Report Helpers
    # =====================================================

    @staticmethod
    def _excel_safe(frame):

        safe = frame.copy(deep=False)

        for column in safe.columns:
            if isinstance(safe[column].dtype, pd.CategoricalDtype):
                safe[column] = safe[column].astype("string")

            if pd.api.types.is_datetime64tz_dtype(safe[column].dtype):
                safe[column] = safe[column].dt.tz_localize(None)

        return safe

    def _build_summary(self, report_df):

        status = report_df["Status"].astype("string")
        regime = report_df["Latest_Stable_Regime"].astype("string")
        accepted = report_df["Latest_Setup_Accepted"].fillna(False)

        rows = [
            ["RUN", "Report generated", pd.Timestamp.now()],
            ["RUN", "Calculation backend", self.calculation_backend],
            ["CONFIGURATION", "Minimum mathematical score", self.minimum_mathematical_score],
            ["CONFIGURATION", "Residual weight", self.residual_weight],
            ["CONFIGURATION", "Regime weight", self.regime_weight],
            ["CONFIGURATION", "Volume weight", self.volume_weight],
            ["CONFIGURATION", "Consistency weight", self.consistency_weight],
            ["COUNTS", "Total shares", int(len(report_df))],
            ["COUNTS", "Calculated shares", int(status.eq("CALCULATED").sum())],
            ["COUNTS", "Not-ready shares", int(status.eq("NOT_READY").sum())],
            ["COUNTS", "Failed shares", int(status.eq("FAILED").sum())],
            ["COUNTS", "Accepted setups", int(accepted.sum())],
            ["REGIME", "MOMENTUM shares", int(regime.eq("MOMENTUM").sum())],
            ["REGIME", "MEAN_REVERSION shares", int(regime.eq("MEAN_REVERSION").sum())],
            ["REGIME", "UNSTABLE shares", int(regime.eq("UNSTABLE").sum())],
            ["REGIME", "RANDOM shares", int(regime.eq("RANDOM").sum())],
        ]

        return pd.DataFrame(
            rows,
            columns=["Section", "Metric", "Value"]
        )

    @staticmethod
    def _style_workbook(writer):

        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(color="FFFFFF", bold=True)
        green_fill = PatternFill("solid", fgColor="C6EFCE")
        red_fill = PatternFill("solid", fgColor="FFC7CE")
        amber_fill = PatternFill("solid", fgColor="FFEB9C")

        def add_formula_rule(
            worksheet,
            column_letter,
            formula,
            fill,
        ):
            """Add a data-row rule only when the sheet has data rows.

            An empty DataFrame written by pandas still creates a header row,
            making ``worksheet.max_row`` equal to 1.  Building a range such as
            ``G2:G1`` raises openpyxl's MultiCellRange TypeError.  Keeping the
            range validation here prevents every report sheet from repeating
            that failure-prone logic.
            """
            if not column_letter or worksheet.max_row < 2:
                return

            target = (
                f"{column_letter}2:"
                f"{column_letter}{worksheet.max_row}"
            )
            worksheet.conditional_formatting.add(
                target,
                FormulaRule(formula=[formula], fill=fill),
            )

        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.sheet_view.showGridLines = False
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            for column_cells in worksheet.columns:
                values = [
                    "" if cell.value is None else str(cell.value)
                    for cell in column_cells[:200]
                ]
                width = min(max(max(map(len, values), default=0) + 2, 11), 34)
                worksheet.column_dimensions[
                    column_cells[0].column_letter
                ].width = width

            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, pd.Timestamp):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
                    elif isinstance(cell.value, float):
                        cell.number_format = "0.0000"

            headers = {
                cell.value: cell.column_letter
                for cell in worksheet[1]
            }

            accepted_column = headers.get("Latest_Setup_Accepted")
            add_formula_rule(
                worksheet=worksheet,
                column_letter=accepted_column,
                formula=f'{accepted_column}2=TRUE',
                fill=green_fill,
            )

            regime_column = headers.get("Latest_Stable_Regime")
            add_formula_rule(
                worksheet=worksheet,
                column_letter=regime_column,
                formula=(
                    f'OR({regime_column}2="UNSTABLE",'
                    f'{regime_column}2="RANDOM")'
                ),
                fill=red_fill,
            )

            reason_column = headers.get("Latest_Rejection_Reason")
            add_formula_rule(
                worksheet=worksheet,
                column_letter=reason_column,
                formula=f'{reason_column}2<>"ACCEPTED"',
                fill=amber_fill,
            )

    def _save_excel_report(self, report_df, output_file):

        output_path = Path(output_file).expanduser()
        output_path.parent.mkdir(parents=True, exist_ok=True)

        directional = report_df.loc[
            report_df["Latest_Stable_Regime"].isin(
                ["MOMENTUM", "MEAN_REVERSION"]
            )
        ].copy()

        unstable_random = report_df.loc[
            report_df["Latest_Stable_Regime"].isin(
                ["UNSTABLE", "RANDOM"]
            )
        ].copy()

        accepted = directional.loc[
            directional["Latest_Setup_Accepted"].fillna(False)
        ].copy()

        rejection_summary = (
            report_df
            .groupby(
                ["Latest_Stable_Regime", "Latest_Rejection_Reason"],
                observed=True,
                dropna=False
            )
            .size()
            .rename("Share_Count")
            .reset_index()
            .sort_values(
                ["Share_Count", "Latest_Stable_Regime"],
                ascending=[False, True]
            )
        )

        summary = self._build_summary(report_df)

        temporary = tempfile.NamedTemporaryFile(
            prefix="MathematicalScoreEngine_",
            suffix=".xlsx",
            dir=output_path.parent,
            delete=False
        )
        temporary_name = temporary.name
        temporary.close()

        try:
            with pd.ExcelWriter(
                temporary_name,
                engine="openpyxl",
                datetime_format="yyyy-mm-dd hh:mm:ss"
            ) as writer:
                self._excel_safe(summary).to_excel(
                    writer, sheet_name="Summary", index=False
                )
                self._excel_safe(directional).to_excel(
                    writer, sheet_name="Directional_Setups", index=False
                )
                self._excel_safe(unstable_random).to_excel(
                    writer, sheet_name="Unstable_Random", index=False
                )
                self._excel_safe(accepted).to_excel(
                    writer, sheet_name="Accepted_Setups", index=False
                )
                self._excel_safe(rejection_summary).to_excel(
                    writer, sheet_name="Rejection_Summary", index=False
                )
                self._excel_safe(report_df).to_excel(
                    writer, sheet_name="All_Shares", index=False
                )
                self._style_workbook(writer)

            os.replace(temporary_name, output_path)
        except Exception:
            if os.path.exists(temporary_name):
                os.unlink(temporary_name)
            raise

        report_df.attrs["Excel_File"] = str(output_path)
        return str(output_path)

    # =====================================================
    # Calculate All Symbols
    # =====================================================

    def calculate_all(self, volatility_data, output_file=None):

        if not isinstance(volatility_data, dict):
            raise TypeError(
                "volatility_data must be a dictionary "
                "containing {symbol: DataFrame}"
            )

        if not volatility_data:
            raise ValueError(
                "volatility_data is empty"
            )

        score_data = {}
        reports = []

        print("=" * 120)
        print("FMRSS : CALCULATING MATHEMATICAL SCORES")
        print(
            "Calculation backend : "
            f"{self.calculation_backend}"
        )
        print("=" * 120)

        for symbol, volatility_df in volatility_data.items():

            symbol = (
                str(symbol)
                .strip()
                .upper()
            )

            try:

                result_df, report = self.calculate_symbol(
                    symbol=symbol,
                    volatility_df=volatility_df
                )

                reports.append(report)

                if report["Status"] == "NOT_READY":

                    if self.print_symbol_details:
                        print(
                            f"{symbol:<15}"
                            "NOT READY : No valid score rows"
                        )

                    continue

                score_data[symbol] = result_df

                score = report["Latest_Mathematical_Score"]

                score_text = (
                    f"{score:.2f}"
                    if pd.notna(score)
                    else "NaN"
                )

                if self.print_symbol_details:
                    print(
                        f"{symbol:<15}"
                        f"CALCULATED : "
                        f"Score={score_text} | "
                        f"Regime={report['Latest_Stable_Regime']} | "
                        f"Accepted={report['Latest_Setup_Accepted']} | "
                        f"Reason={report['Latest_Rejection_Reason']}"
                    )

            except Exception as error:

                reports.append({
                    "Symbol": symbol,
                    "Status": "FAILED",
                    "Input_Rows": (
                        len(volatility_df)
                        if isinstance(
                            volatility_df,
                            pd.DataFrame
                        )
                        else 0
                    ),
                    "Score_Ready_Rows": 0,
                    "Accepted_Rows": 0,
                    "Latest_DateTime": pd.NaT,
                    "Latest_Stable_Regime": "NOT_READY",
                    "Latest_Residual_Return": np.nan,
                    "Latest_Residual_Impulse": np.nan,
                    "Latest_Residual_ZScore": np.nan,
                    "Latest_Regime_Strength": np.nan,
                    "Latest_Volume_Factor": np.nan,
                    "Latest_Momentum_Direction": "NEUTRAL",
                    "Latest_Directional_Consistency": np.nan,
                    "Latest_Residual_Score": np.nan,
                    "Latest_Regime_Score": np.nan,
                    "Latest_Volume_Score": np.nan,
                    "Latest_Consistency_Score": np.nan,
                    "Latest_Mathematical_Score": np.nan,
                    "Latest_Setup_Ready": False,
                    "Latest_Setup_Accepted": False,
                    "Latest_Volatility_Eligible": False,
                    "Latest_Volume_Confirmed": False,
                    "Latest_Rejection_Reason": "CALCULATION_FAILED",
                    "Calculation_Backend": self.calculation_backend,
                    "Error": str(error)
                })

                if self.print_symbol_details:
                    print(
                        f"{symbol:<15}"
                        f"FAILED : {error}"
                    )

        report_df = pd.DataFrame(reports)

        if not report_df.empty:

            report_df.sort_values(
                by=["Status", "Symbol"],
                inplace=True
            )

            report_df.reset_index(
                drop=True,
                inplace=True
            )

        print("=" * 120)
        print(
            f"Successfully scored : "
            f"{len(score_data)}/{len(volatility_data)}"
        )
        print("=" * 120)

        if not score_data:
            raise RuntimeError(
                "Mathematical scores could not be "
                "calculated for any symbol"
            )

        selected_output_file = (
            output_file
            if output_file is not None
            else self.report_output_file
        )

        if self.save_excel_enabled and selected_output_file:
            saved_file = self._save_excel_report(
                report_df=report_df,
                output_file=selected_output_file
            )
            print(f"Mathematical score workbook : {saved_file}")

        return score_data, report_df
